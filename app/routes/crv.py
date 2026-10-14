from datetime import datetime, date
from io import BytesIO
from pathlib import Path
import calendar
import json
import uuid
import re
from xml.sax.saxutils import escape
from sqlalchemy import text


from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    session,
    flash,
    send_file,
    current_app
)
from werkzeug.utils import secure_filename

try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None


from app.extensions import db
from app.models.crv_movimiento import CRVMovimiento
from app.models.crv_tercero import CRVTercero
from app.models.sede import Sede
from app.models.crv_cuenta_pagar import CRVCuentaPagar
from app.models.crv_pago_cuenta_pagar import CRVPagoCuentaPagar
from app.models.crv_cuenta_cobrar import CRVCuentaCobrar
from app.models.crv_cobro_cuenta_cobrar import CRVCobroCuentaCobrar
from app.models.crv_proyecto import CRVProyecto
from app.models.crv_cierre_mensual import CRVCierreMensual
from app.models.crv_conciliacion_bancaria import CRVConciliacionBancaria

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image,
    PageBreak,
    KeepTogether
)


crv_bp = Blueprint(
    "crv",
    __name__,
    url_prefix="/crv"
)


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def usuario_valido():
    return "usuario" in session


def empresa_crv_valida():
    """
    Este módulo es exclusivo de CRV Constructora.
    Evita mezclar datos con Hotel, Ecobike,
    Comercializadora o Financiera.
    """

    empresa_id = session.get("empresa_id")

    empresa_nombre = str(
        session.get("empresa_nombre", "")
    ).strip().upper()

    if not empresa_id:
        return False

    return (
        "CRV" in empresa_nombre
        and "CONSTRUCTORA" in empresa_nombre
    )


def proteger_crv():

    if not usuario_valido():
        return redirect("/login")

    if not session.get("empresa_id"):
        return redirect("/empresas")

    if not empresa_crv_valida():
        return redirect("/dashboard")

    return None


def convertir_fecha(valor):

    if not valor:
        return None

    try:
        return datetime.strptime(
            valor,
            "%Y-%m-%d"
        ).date()

    except ValueError:
        return None


def convertir_numero(valor, default=0):
    """
    Convierte valores monetarios sin multiplicarlos por 100.

    Acepta, entre otros:
      152100.84   -> 152100.84
      152100,84   -> 152100.84
      152.100,84  -> 152100.84
      152,100.84  -> 152100.84
      1.190.000   -> 1190000
      1,190,000   -> 1190000
      181000      -> 181000
    """

    if valor is None:
        return default

    texto = str(valor).strip()

    if not texto:
        return default

    texto = (
        texto
        .replace("$", "")
        .replace("COP", "")
        .replace("cop", "")
        .replace(" ", "")
        .strip()
    )

    if not texto:
        return default

    signo = ""

    if texto.startswith("-"):
        signo = "-"
        texto = texto[1:]
    elif texto.startswith("+"):
        texto = texto[1:]

    if not texto:
        return default

    tiene_punto = "." in texto
    tiene_coma = "," in texto

    try:
        if tiene_punto and tiene_coma:

            ultimo_punto = texto.rfind(".")
            ultima_coma = texto.rfind(",")

            if ultimo_punto > ultima_coma:
                texto = texto.replace(",", "")
            else:
                texto = texto.replace(".", "")
                texto = texto.replace(",", ".")

        elif tiene_coma:

            partes = texto.split(",")

            if (
                len(partes) == 2
                and len(partes[-1]) in (1, 2)
            ):
                texto = texto.replace(",", ".")
            else:
                texto = texto.replace(",", "")

        elif tiene_punto:

            partes = texto.split(".")

            if (
                len(partes) == 2
                and len(partes[-1]) in (1, 2)
            ):
                pass
            else:
                texto = texto.replace(".", "")

        return float(
            signo + texto
        )

    except (TypeError, ValueError):
        return default


def convertir_porcentaje(valor, default=0):

    if valor is None:
        return default

    texto = str(valor).strip().replace(",", ".")

    if not texto:
        return default

    try:
        return float(texto)
    except (TypeError, ValueError):
        return default


def calcular_iva_formulario(formulario):
    """
    Devuelve los valores tributarios del documento sin asumir que
    todo IVA pagado sea descontable.

    tipo_iva:
      - NO_APLICA
      - EXCLUIDO
      - EXENTO
      - GENERADO
      - DESCONTABLE
      - MAYOR_VALOR_COSTO

    EXENTO y EXCLUIDO siempre se guardan con tarifa 0 e IVA 0.
    """

    base_gravable = convertir_numero(
        formulario.get("base_gravable")
    )

    tipo_iva = str(
        formulario.get("tipo_iva", "NO_APLICA")
        or "NO_APLICA"
    ).strip().upper()

    tarifa_iva = convertir_porcentaje(
        formulario.get("tarifa_iva"),
        0
    )

    if tipo_iva in (
        "NO_APLICA",
        "EXCLUIDO",
        "EXENTO",
    ):
        tarifa_iva = 0
        valor_iva = 0
    else:
        tarifa_iva = max(0, tarifa_iva)
        valor_iva = round(
            base_gravable * tarifa_iva / 100,
            2
        )

    total_documento = round(
        base_gravable + valor_iva,
        2
    )

    return {
        "base_gravable": base_gravable,
        "tipo_iva": tipo_iva,
        "tarifa_iva": tarifa_iva,
        "valor_iva": valor_iva,
        "total_documento": total_documento,
    }


def resumen_iva_movimientos(movimientos):
    """
    Resumen tributario de IVA.

    IMPORTANTE:
    Los movimientos NO_APLICA siguen existiendo en contabilidad,
    pero NO se suman como base de IVA. Esto evita que nomina,
    transferencias, anticipos u otros conceptos sin IVA inflen la
    base tributaria del reporte.
    """

    resultado = {
        # Base gravable sujeta a IVA distinta de NO_APLICA,
        # EXENTO y EXCLUIDO.
        "base_gravable": 0.0,

        # Bases separadas por tratamiento.
        "base_generado": 0.0,
        "base_descontable": 0.0,
        "base_mayor_valor": 0.0,
        "base_exenta": 0.0,
        "base_excluida": 0.0,
        "base_no_aplica": 0.0,

        # Valores de impuesto.
        "iva_generado": 0.0,
        "iva_descontable": 0.0,
        "iva_mayor_valor": 0.0,
        "iva_total": 0.0,

        # Referencia gerencial. No sustituye la liquidacion fiscal
        # definitiva del contador.
        "saldo_iva_preliminar": 0.0,
    }

    for movimiento in movimientos:
        base = float(
            getattr(movimiento, "base_gravable", 0)
            or 0
        )

        iva = float(
            getattr(movimiento, "valor_iva", 0)
            or 0
        )

        tipo = str(
            getattr(movimiento, "tipo_iva", "NO_APLICA")
            or "NO_APLICA"
        ).strip().upper()

        if tipo == "GENERADO":
            resultado["base_gravable"] += base
            resultado["base_generado"] += base
            resultado["iva_generado"] += iva
            resultado["iva_total"] += iva

        elif tipo == "DESCONTABLE":
            resultado["base_gravable"] += base
            resultado["base_descontable"] += base
            resultado["iva_descontable"] += iva
            resultado["iva_total"] += iva

        elif tipo == "MAYOR_VALOR_COSTO":
            resultado["base_gravable"] += base
            resultado["base_mayor_valor"] += base
            resultado["iva_mayor_valor"] += iva
            resultado["iva_total"] += iva

        elif tipo == "EXENTO":
            resultado["base_exenta"] += base

        elif tipo == "EXCLUIDO":
            resultado["base_excluida"] += base

        else:
            # NO_APLICA y cualquier registro historico no tributario.
            resultado["base_no_aplica"] += base

    resultado["saldo_iva_preliminar"] = (
        resultado["iva_generado"]
        - resultado["iva_descontable"]
    )

    return resultado


def obtener_tercero_por_identificacion(
    empresa_id,
    identificacion
):

    identificacion = str(
        identificacion or ""
    ).strip()

    if not identificacion:
        return None

    return CRVTercero.query.filter_by(
        empresa_id=empresa_id,
        identificacion=identificacion
    ).first()


def crear_o_actualizar_tercero_desde_formulario(
    empresa_id,
    formulario
):

    identificacion = formulario.get(
        "tercero_identificacion",
        ""
    ).strip()

    razon_social = formulario.get(
        "tercero_nombre",
        ""
    ).strip()

    if not identificacion and not razon_social:
        return None

    tercero = None

    if identificacion:

        tercero = obtener_tercero_por_identificacion(
            empresa_id,
            identificacion
        )

    if tercero is None:

        tercero = CRVTercero(
            empresa_id=empresa_id,

            tipo_tercero=formulario.get(
                "tipo_tercero",
                "TERCERO"
            ).strip().upper() or "TERCERO",

            tipo_documento=formulario.get(
                "tipo_documento_tercero",
                ""
            ).strip(),

            identificacion=(
                identificacion
                or f"SIN-DOC-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
            ),

            razon_social=(
                razon_social
                or "TERCERO SIN NOMBRE"
            ),

            nombre_comercial=formulario.get(
                "tercero_nombre_comercial",
                ""
            ).strip(),

            direccion=formulario.get(
                "tercero_direccion",
                ""
            ).strip(),

            ciudad=formulario.get(
                "tercero_ciudad",
                ""
            ).strip(),

            departamento=formulario.get(
                "tercero_departamento",
                ""
            ).strip(),

            telefono=formulario.get(
                "tercero_telefono",
                ""
            ).strip(),

            celular=formulario.get(
                "tercero_celular",
                ""
            ).strip(),

            email=formulario.get(
                "tercero_email",
                ""
            ).strip(),

            contacto=formulario.get(
                "tercero_contacto",
                ""
            ).strip()
        )

        db.session.add(tercero)
        db.session.flush()

    else:

        if razon_social:
            tercero.razon_social = razon_social

        tipo_tercero = formulario.get(
            "tipo_tercero",
            ""
        ).strip()

        if tipo_tercero:
            tercero.tipo_tercero = tipo_tercero.upper()

        campos = {
            "nombre_comercial": "tercero_nombre_comercial",
            "direccion": "tercero_direccion",
            "ciudad": "tercero_ciudad",
            "departamento": "tercero_departamento",
            "telefono": "tercero_telefono",
            "celular": "tercero_celular",
            "email": "tercero_email",
            "contacto": "tercero_contacto"
        }

        for atributo, campo_form in campos.items():

            valor = formulario.get(
                campo_form,
                ""
            ).strip()

            if valor:
                setattr(
                    tercero,
                    atributo,
                    valor
                )

    return tercero



# ============================================================
# COMPROBANTES - CONSECUTIVOS PERMANENTES CE / RC / CC
# ============================================================

def _tipo_comprobante_movimiento(movimiento):
    """
    Determina el tipo de comprobante interno.

    RC = Recibo de caja:
         entrada real de dinero.

    CE = Comprobante de egreso:
         salida real de dinero.

    CC = Comprobante contable:
         reconocimiento contable sin entrada/salida inmediata,
         saldos iniciales, transferencias internas y operaciones
         a credito.
    """

    tipo = str(
        movimiento.tipo or ""
    ).strip().upper()

    forma_pago = str(
        movimiento.forma_pago or ""
    ).strip().upper()

    cuenta = str(
        movimiento.cuenta or ""
    ).strip().upper()

    if tipo == "COBRO CUENTA POR COBRAR":
        return "RECIBO DE CAJA", "RC"

    if tipo == "INGRESO":
        if (
            forma_pago in ("CREDITO", "CRÉDITO")
            or cuenta == "CUENTAS POR COBRAR"
        ):
            return "COMPROBANTE CONTABLE", "CC"

        return "RECIBO DE CAJA", "RC"

    if tipo == "COMPRA":
        if (
            forma_pago in ("CREDITO", "CRÉDITO")
            or cuenta == "CUENTAS POR PAGAR"
        ):
            return "COMPROBANTE CONTABLE", "CC"

        return "COMPROBANTE DE EGRESO", "CE"

    tipos_egreso = {
        "GASTO",
        "NOMINA AFILIADOS",
        "NÓMINA AFILIADOS",
        "NOMINA OBRA",
        "NÓMINA OBRA",
        "ANTICIPO",
        "PRESTAMO EMPLEADO",
        "PRÉSTAMO EMPLEADO",
        "EGRESO",
        "OTRO EGRESO",
        "PAGO CUENTA POR PAGAR",
    }

    if tipo in tipos_egreso:
        return "COMPROBANTE DE EGRESO", "CE"

    return "COMPROBANTE CONTABLE", "CC"


def _titulo_comprobante_por_prefijo(prefijo):
    prefijo = str(
        prefijo or ""
    ).strip().upper()

    if prefijo == "RC":
        return "RECIBO DE CAJA"

    if prefijo == "CE":
        return "COMPROBANTE DE EGRESO"

    return "COMPROBANTE CONTABLE"


def _asegurar_tabla_secuencias_comprobantes():
    """
    Conserva el ultimo numero utilizado por empresa y por serie.
    Esta tabla impide reutilizar un consecutivo aunque posteriormente
    se elimine el movimiento que lo utilizo.
    """

    db.session.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS crv_comprobante_secuencia (
                empresa_id INTEGER NOT NULL,
                comprobante_tipo VARCHAR(10) NOT NULL,
                ultimo_consecutivo INTEGER NOT NULL DEFAULT 0,
                PRIMARY KEY (empresa_id, comprobante_tipo)
            )
            """
        )
    )


def _siguiente_consecutivo_comprobante(
    empresa_id,
    comprobante_tipo
):
    """
    Obtiene el siguiente numero de la serie CE / RC / CC.

    Si la tabla de secuencias aun no tiene registro para una serie,
    toma como punto de partida el mayor consecutivo historico ya
    asignado en crv_movimiento.
    """

    empresa_id = int(empresa_id)
    comprobante_tipo = str(
        comprobante_tipo or "CC"
    ).strip().upper()

    _asegurar_tabla_secuencias_comprobantes()

    fila = db.session.execute(
        text(
            """
            SELECT ultimo_consecutivo
            FROM crv_comprobante_secuencia
            WHERE
                empresa_id = :empresa_id
                AND comprobante_tipo = :tipo
            """
        ),
        {
            "empresa_id": empresa_id,
            "tipo": comprobante_tipo,
        }
    ).mappings().first()

    if fila is None:
        maximo_historico = db.session.query(
            db.func.max(
                CRVMovimiento.comprobante_consecutivo
            )
        ).filter(
            CRVMovimiento.empresa_id == empresa_id,
            CRVMovimiento.comprobante_tipo == comprobante_tipo
        ).scalar()

        ultimo = int(
            maximo_historico
            or 0
        )

        db.session.execute(
            text(
                """
                INSERT INTO crv_comprobante_secuencia (
                    empresa_id,
                    comprobante_tipo,
                    ultimo_consecutivo
                )
                VALUES (
                    :empresa_id,
                    :tipo,
                    :ultimo
                )
                """
            ),
            {
                "empresa_id": empresa_id,
                "tipo": comprobante_tipo,
                "ultimo": ultimo,
            }
        )
    else:
        ultimo = int(
            fila["ultimo_consecutivo"]
            or 0
        )

    siguiente = ultimo + 1

    db.session.execute(
        text(
            """
            UPDATE crv_comprobante_secuencia
            SET ultimo_consecutivo = :siguiente
            WHERE
                empresa_id = :empresa_id
                AND comprobante_tipo = :tipo
            """
        ),
        {
            "siguiente": siguiente,
            "empresa_id": empresa_id,
            "tipo": comprobante_tipo,
        }
    )

    return siguiente


def _asignar_comprobante_permanente(movimiento):
    """
    Asigna el comprobante una sola vez.

    Si el movimiento ya tiene tipo, consecutivo y codigo, no se
    modifica. De esta forma editar el movimiento o volver a abrir
    el PDF no cambia el numero emitido.
    """

    if movimiento is None:
        return None

    tipo_guardado = str(
        getattr(
            movimiento,
            "comprobante_tipo",
            ""
        )
        or ""
    ).strip().upper()

    consecutivo_guardado = int(
        getattr(
            movimiento,
            "comprobante_consecutivo",
            0
        )
        or 0
    )

    codigo_guardado = str(
        getattr(
            movimiento,
            "comprobante_codigo",
            ""
        )
        or ""
    ).strip().upper()

    if (
        tipo_guardado
        and consecutivo_guardado > 0
        and codigo_guardado
    ):
        return codigo_guardado

    # Si la migracion dejo tipo + numero pero falta solamente el
    # codigo visible, se reconstruye sin consumir un numero nuevo.
    if (
        tipo_guardado
        and consecutivo_guardado > 0
    ):
        codigo = (
            f"{tipo_guardado}-"
            f"{consecutivo_guardado:06d}"
        )

        movimiento.comprobante_codigo = codigo

        return codigo

    _, prefijo = _tipo_comprobante_movimiento(
        movimiento
    )

    siguiente = _siguiente_consecutivo_comprobante(
        movimiento.empresa_id,
        prefijo
    )

    codigo = (
        f"{prefijo}-"
        f"{siguiente:06d}"
    )

    movimiento.comprobante_tipo = prefijo
    movimiento.comprobante_consecutivo = siguiente
    movimiento.comprobante_codigo = codigo

    return codigo



# ============================================================
# CARGA INTELIGENTE - FUNCIONES AUXILIARES
# ============================================================

EXTENSIONES_CARGA_INTELIGENTE = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}

TAMANO_MAXIMO_ARCHIVO_CARGA = 25 * 1024 * 1024


def _carpeta_carga_inteligente(empresa_id):
    carpeta = (
        Path(current_app.instance_path)
        / "crv_carga_inteligente"
        / str(int(empresa_id))
    )

    carpeta.mkdir(
        parents=True,
        exist_ok=True
    )

    return carpeta


def _ruta_lote_carga(empresa_id, lote_id):
    lote_seguro = str(
        lote_id or ""
    ).strip()

    if not lote_seguro:
        return None

    permitidos = (
        "abcdefghijklmnopqrstuvwxyz"
        "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
        "0123456789-_"
    )

    if any(
        caracter not in permitidos
        for caracter in lote_seguro
    ):
        return None

    return (
        _carpeta_carga_inteligente(
            empresa_id
        )
        / lote_seguro
    )


def _leer_manifest_carga(empresa_id, lote_id):
    carpeta = _ruta_lote_carga(
        empresa_id,
        lote_id
    )

    if not carpeta:
        return None

    ruta = carpeta / "lote.json"

    if not ruta.exists():
        return None

    try:
        return json.loads(
            ruta.read_text(
                encoding="utf-8"
            )
        )
    except Exception:
        return None


def _guardar_manifest_carga(
    empresa_id,
    lote_id,
    datos
):
    carpeta = _ruta_lote_carga(
        empresa_id,
        lote_id
    )

    if not carpeta:
        raise ValueError(
            "Identificador de lote no valido."
        )

    carpeta.mkdir(
        parents=True,
        exist_ok=True
    )

    ruta = carpeta / "lote.json"

    ruta.write_text(
        json.dumps(
            datos,
            ensure_ascii=False,
            indent=2
        ),
        encoding="utf-8"
    )


def _listar_lotes_carga_inteligente(
    empresa_id,
    limite=20
):
    raiz = _carpeta_carga_inteligente(
        empresa_id
    )

    lotes = []

    for carpeta in raiz.iterdir():

        if not carpeta.is_dir():
            continue

        manifest = _leer_manifest_carga(
            empresa_id,
            carpeta.name
        )

        if not manifest:
            continue

        lotes.append(
            manifest
        )

    lotes.sort(
        key=lambda x: str(
            x.get(
                "fecha_creacion",
                ""
            )
        ),
        reverse=True
    )

    return lotes[:limite]


def _archivo_carga_valido(archivo):
    if archivo is None:
        return False, "Archivo invalido."

    nombre_original = str(
        archivo.filename or ""
    ).strip()

    if not nombre_original:
        return False, "El archivo no tiene nombre."

    extension = Path(
        nombre_original
    ).suffix.lower()

    if extension not in EXTENSIONES_CARGA_INTELIGENTE:
        return (
            False,
            (
                "Formato no permitido. "
                "Usa PDF, PNG, JPG, JPEG o WEBP."
            )
        )

    return True, ""



def _normalizar_identificacion_carga(valor):
    return re.sub(
        r"[^0-9A-Za-z]",
        "",
        str(valor or "")
    ).upper()


def _limpiar_texto_factura(texto):
    lineas = []

    for linea in str(
        texto or ""
    ).replace(
        "\r",
        "\n"
    ).split(
        "\n"
    ):
        limpia = re.sub(
            r"\s+",
            " ",
            linea
        ).strip()

        if limpia:
            lineas.append(
                limpia
            )

    return "\n".join(
        lineas
    )


def _leer_texto_pdf(ruta):
    if PdfReader is None:
        return "", (
            "Falta instalar pypdf. "
            "Ejecuta: py -m pip install pypdf"
        )

    try:
        lector = PdfReader(
            str(ruta)
        )

        paginas = []

        for pagina in lector.pages:
            try:
                paginas.append(
                    pagina.extract_text()
                    or ""
                )
            except Exception:
                paginas.append(
                    ""
                )

        texto = _limpiar_texto_factura(
            "\n".join(
                paginas
            )
        )

        if not texto:
            return "", (
                "El PDF no contiene texto seleccionable. "
                "Quedará para revisión manual."
            )

        return texto, ""

    except Exception as error:
        return "", (
            "No fue posible leer el PDF: "
            f"{error}"
        )


def _encontrar_valor_despues_etiqueta(
    texto,
    etiquetas
):
    for etiqueta in etiquetas:
        patron = (
            rf"(?im)"
            rf"{etiqueta}"
            rf"[^\n0-9]{{0,20}}"
            rf"(\$?\s*[0-9][0-9\.,]*)"
        )

        coincidencia = re.search(
            patron,
            texto
        )

        if coincidencia:
            return convertir_numero(
                coincidencia.group(1),
                None
            )

    return None


def _extraer_fecha_factura(texto, nombre_archivo=""):
    candidatos = []

    patrones = [
        r"(?<!\d)(\d{2}[-/]\d{2}[-/]\d{4})(?!\d)",
        r"(?<!\d)(\d{4}[-/]\d{2}[-/]\d{2})(?!\d)",
        r"(?<!\d)(\d{2}-[A-Za-z]{3}-\d{4})(?!\d)",
    ]

    for fuente in (
        texto,
        nombre_archivo
    ):
        for patron in patrones:
            candidatos.extend(
                re.findall(
                    patron,
                    str(fuente or "")
                )
            )

    for candidato in candidatos:
        for formato in (
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%d-%b-%Y",
        ):
            try:
                return datetime.strptime(
                    candidato,
                    formato
                ).date().isoformat()
            except Exception:
                continue

    return ""


def _extraer_nit_factura(texto):
    patrones = [
        r"(?im)\bN\.?I\.?T\.?\s*[:#]?\s*([0-9][0-9\.\-\s]{5,20})",
        r"(?im)\bNit\s*[:#]?\s*([0-9][0-9\.\-\s]{5,20})",
    ]

    for patron in patrones:
        coincidencia = re.search(
            patron,
            texto
        )

        if coincidencia:
            valor = coincidencia.group(1).strip()
            valor = re.sub(
                r"\s+",
                "",
                valor
            )
            return valor

    return ""


def _numero_factura_candidato_valido(valor):
    candidato = str(valor or "").strip().upper()

    if not candidato or len(candidato) < 3:
        return False

    prohibidos = {
        "EMAIL", "E-MAIL", "MAIL", "MEDIO", "REMISION", "REMISIÓN",
        "PAGO", "VENTA", "FECHA", "TOTAL", "FACTURA", "DOCUMENTO",
        "CLIENTE", "ADQUIRENTE", "CONTADO", "CREDITO", "CRÉDITO",
    }

    if candidato in prohibidos:
        return False

    if any(palabra in candidato for palabra in ("EMAIL", "REMISION", "REMISIÓN")):
        return False

    # Un número de factura normalmente contiene al menos un dígito.
    if not re.search(r"\d", candidato):
        return False

    return True


def _extraer_numero_factura(texto):
    patrones = [
        r"(?im)FACTURA\s+ELECTR[ÓO]NICA(?:\s+DE\s+VENTA)?\s*(?:No\.?|N[°º])?\s*[:#-]?\s*([A-Z0-9][A-Z0-9\-]{2,29})",
        r"(?im)\bFACTURA\s*(?:No\.?|N[°º])\s*[:#-]?\s*([A-Z0-9][A-Z0-9\-]{2,29})",
        r"(?im)\b(?:No\.?|N[°º])\s*FACTURA\s*[:#-]?\s*([A-Z0-9][A-Z0-9\-]{2,29})",
        r"(?im)\bPREFIJO\s*[:#-]?\s*([A-Z]{1,8})\s+N[ÚU]MERO\s*[:#-]?\s*(\d{2,20})",
    ]

    for patron in patrones:
        coincidencia = re.search(patron, texto)
        if not coincidencia:
            continue

        if coincidencia.lastindex == 2:
            candidato = f"{coincidencia.group(1)}{coincidencia.group(2)}"
        else:
            candidato = coincidencia.group(1)

        candidato = str(candidato or "").strip()

        if _numero_factura_candidato_valido(candidato):
            return candidato

    return ""


def _extraer_tarifa_iva_factura(texto):
    coincidencia = re.search(
        r"(?im)\bIVA\b[^\n%]{0,40}?(\d{1,2}(?:[\.,]\d+)?)\s*%",
        texto
    )

    if coincidencia:
        valor = convertir_numero(
            coincidencia.group(1),
            None
        )

        if valor is not None:
            return float(
                valor
            )

    porcentajes = re.findall(
        r"(?<!\d)(\d{1,2}(?:[\.,]\d+)?)\s*%",
        texto
    )

    for porcentaje in porcentajes:
        valor = convertir_numero(
            porcentaje,
            None
        )

        if valor in (
            0,
            5,
            19
        ):
            return float(
                valor
            )

    return None


def _extraer_valor_iva_factura(texto):
    # Primero intenta identificar una línea de IVA con varios números.
    for linea in texto.splitlines():
        if "IVA" not in linea.upper():
            continue

        numeros = re.findall(
            r"(?<!\w)(?:\$?\s*)?([0-9][0-9\.,]*)",
            linea
        )

        valores = []

        for numero in numeros:
            valor = convertir_numero(
                numero,
                None
            )

            if valor is not None:
                valores.append(
                    float(valor)
                )

        # Se ignoran posibles porcentajes pequeños.
        monetarios = [
            v
            for v in valores
            if v > 100
        ]

        if monetarios:
            return monetarios[-1]

    return _encontrar_valor_despues_etiqueta(
        texto,
        [
            r"\bVALOR\s+IVA\b",
            r"\bIVA\b",
            r"\bIMPUESTO\b",
        ]
    )


def _extraer_total_factura(texto):
    candidatos = []

    for patron in [
        r"(?im)\bTOTAL(?:\s+A\s+PAGAR)?\s*\$?\s*([0-9][0-9\.,]*)",
        r"(?im)\bTOTAL\s*\$?\s*([0-9][0-9\.,]*)",
        r"(?im)\bVALOR\s+TOTAL\s*\$?\s*([0-9][0-9\.,]*)",
    ]:
        for coincidencia in re.findall(
            patron,
            texto
        ):
            valor = convertir_numero(
                coincidencia,
                None
            )

            if valor is not None:
                candidatos.append(
                    float(valor)
                )

    if candidatos:
        return max(
            candidatos
        )

    return None


def _extraer_base_factura(texto):
    valor = _encontrar_valor_despues_etiqueta(
        texto,
        [
            r"\bSUBTOTAL\b",
            r"\bSUBTOT\.?\b",
            r"\bBASE\s+GRAVABLE\b",
            r"\bBASE\b",
        ]
    )

    if valor is not None:
        return float(
            valor
        )

    return None


def _proveedor_candidato_valido(candidato):
    nombre = re.sub(r"\s+", " ", str(candidato or "")).strip()
    mayus = nombre.upper()

    if len(nombre) < 4:
        return False

    descartes = [
        "FACTURA", "ELECTRONICA", "ELECTRÓNICA", "RESPONSABLE",
        "DIRECCION", "DIRECCIÓN", "TELEFONO", "TELÉFONO", "EMAIL",
        "E-MAIL", "CORREO", "CLIENTE", "ADQUIRENTE", "REMISIÓN",
        "REMISION", "MEDIO DE PAGO", "FORMA DE PAGO", "CRV CONSTRUCCIONES",
        "CRV CONSTRUCTORA", "NIT", "TOTAL", "SUBTOTAL", "IVA",
    ]

    if any(palabra in mayus for palabra in descartes):
        return False

    if re.fullmatch(r"[\d\W]+", nombre):
        return False

    # Evita referencias/códigos cortos del encabezado.
    if len(nombre.split()) == 1 and re.search(r"\d", nombre):
        return False

    return True


def _extraer_proveedor_factura(texto, nit=""):
    lineas = [linea.strip() for linea in texto.splitlines() if linea.strip()]

    # 1. Busca expresiones explícitas habituales en facturas.
    patrones_explicitos = [
        r"(?im)^(?:RAZ[ÓO]N\s+SOCIAL|PROVEEDOR|VENDEDOR|EMISOR)\s*[:\-]\s*(.+)$",
        r"(?im)^NOMBRE\s+O\s+RAZ[ÓO]N\s+SOCIAL\s*[:\-]\s*(.+)$",
    ]

    for patron in patrones_explicitos:
        coincidencia = re.search(patron, texto)
        if coincidencia:
            candidato = coincidencia.group(1).strip()
            if _proveedor_candidato_valido(candidato):
                return candidato[:220]

    # 2. Si conocemos el NIT, localiza exactamente la línea que lo contiene.
    nit_normalizado = _normalizar_identificacion_carga(nit)
    posiciones_nit = []

    for indice, linea in enumerate(lineas):
        if nit_normalizado and nit_normalizado in _normalizar_identificacion_carga(linea):
            posiciones_nit.append(indice)

    # Si no coincidió exactamente, usa líneas que mencionan NIT.
    if not posiciones_nit:
        posiciones_nit = [
            indice
            for indice, linea in enumerate(lineas)
            if "NIT" in linea.upper()
        ]

    for indice in posiciones_nit:
        # El nombre del emisor suele estar arriba del NIT. Buscamos hasta 5 líneas.
        for retroceso in (1, 2, 3, 4, 5):
            pos = indice - retroceso
            if pos < 0:
                continue
            candidato = lineas[pos]
            if _proveedor_candidato_valido(candidato):
                return candidato[:220]

    return ""


def _buscar_tercero_por_identificacion_carga(
    empresa_id,
    identificacion
):
    objetivo = _normalizar_identificacion_carga(
        identificacion
    )

    if not objetivo:
        return None

    terceros = (
        CRVTercero.query
        .filter_by(
            empresa_id=empresa_id
        )
        .all()
    )

    for tercero in terceros:
        if (
            _normalizar_identificacion_carga(
                tercero.identificacion
            )
            == objetivo
        ):
            return tercero

    return None


def _es_factura_duplicada_carga(
    empresa_id,
    identificacion,
    numero_documento
):
    numero = str(
        numero_documento or ""
    ).strip().upper()

    nit = _normalizar_identificacion_carga(
        identificacion
    )

    if not numero:
        return False, ""

    movimientos = (
        CRVMovimiento.query
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            db.func.upper(
                CRVMovimiento.numero_documento
            ) == numero
        )
        .all()
    )

    for movimiento in movimientos:
        if (
            not nit
            or _normalizar_identificacion_carga(
                movimiento.tercero_identificacion
            )
            == nit
        ):
            return (
                True,
                "Ya existe un movimiento CRV con esta factura."
            )

    obligaciones = (
        CRVCuentaPagar.query
        .filter(
            CRVCuentaPagar.empresa_id == empresa_id,
            db.func.upper(
                CRVCuentaPagar.numero_documento
            ) == numero
        )
        .all()
    )

    for obligacion in obligaciones:
        proveedor = getattr(
            obligacion,
            "proveedor",
            None
        )

        identificacion_proveedor = getattr(
            proveedor,
            "identificacion",
            ""
        )

        if (
            not nit
            or _normalizar_identificacion_carga(
                identificacion_proveedor
            )
            == nit
        ):
            return (
                True,
                "Ya existe una cuenta por pagar con esta factura."
            )

    return False, ""


def _validar_valores_factura_carga(base, tarifa, iva, total):
    alertas = []

    def numero(valor):
        try:
            return float(valor)
        except (TypeError, ValueError):
            return None

    base = numero(base)
    tarifa = numero(tarifa)
    iva = numero(iva)
    total = numero(total)

    if base is not None and base < 0:
        alertas.append("La base gravable no puede ser negativa.")

    if iva is not None and iva < 0:
        alertas.append("El IVA no puede ser negativo.")

    if total is not None and total <= 0:
        alertas.append("El total debe ser mayor a cero.")

    if tarifa is not None and tarifa < 0:
        alertas.append("La tarifa de IVA no puede ser negativa.")

    if base is not None and tarifa is not None and iva is not None:
        iva_esperado = round(base * tarifa / 100, 2)
        tolerancia = max(2.0, abs(iva_esperado) * 0.002)
        if abs(iva - iva_esperado) > tolerancia:
            alertas.append(
                f"IVA inconsistente: con base {base:,.2f} y tarifa {tarifa:g}% "
                f"se esperan aproximadamente {iva_esperado:,.2f}."
            )

    if base is not None and iva is not None and total is not None:
        total_esperado = round(base + iva, 2)
        tolerancia = max(2.0, abs(total_esperado) * 0.002)
        if abs(total - total_esperado) > tolerancia:
            alertas.append(
                f"Total inconsistente: base + IVA = {total_esperado:,.2f}, "
                f"pero se leyó {total:,.2f}."
            )

    return alertas


def _estado_revision_carga(datos):
    faltantes = []

    if not str(datos.get("proveedor") or "").strip():
        faltantes.append("proveedor")
    if not str(datos.get("identificacion") or "").strip():
        faltantes.append("NIT")
    if not str(datos.get("numero_documento") or "").strip():
        faltantes.append("factura")
    if not str(datos.get("fecha_documento") or "").strip():
        faltantes.append("fecha")
    if datos.get("base_gravable") is None:
        faltantes.append("base")
    if datos.get("valor_iva") is None:
        faltantes.append("IVA")
    if datos.get("total_documento") is None:
        faltantes.append("total")

    alertas = _validar_valores_factura_carga(
        datos.get("base_gravable"),
        datos.get("tarifa_iva"),
        datos.get("valor_iva"),
        datos.get("total_documento"),
    )

    if datos.get("duplicado"):
        estado = "POSIBLE_DUPLICADO"
    elif faltantes or alertas:
        estado = "REVISION_REQUERIDA"
    elif datos.get("revision_manual"):
        estado = "LISTO_PARA_APROBAR"
    else:
        estado = "DATOS_EXTRAIDOS"

    return estado, faltantes, alertas


def _analizar_factura_carga(
    empresa_id,
    ruta,
    nombre_original
):
    extension = ruta.suffix.lower()

    resultado = {
        "estado": "REVISION_REQUERIDA",
        "proveedor": "",
        "identificacion": "",
        "numero_documento": "",
        "fecha_documento": "",
        "base_gravable": None,
        "tarifa_iva": None,
        "valor_iva": None,
        "total_documento": None,
        "tercero_existente": False,
        "tercero_id": None,
        "duplicado": False,
        "revision_manual": False,
        "alertas_validacion": [],
        "observaciones": "",
    }

    if extension != ".pdf":
        resultado["observaciones"] = (
            "La lectura automática de imágenes se agregará en una etapa posterior. "
            "Revisión manual requerida."
        )
        return resultado

    texto, error = _leer_texto_pdf(ruta)
    if error:
        resultado["observaciones"] = error
        return resultado

    nit = _extraer_nit_factura(texto)
    numero_documento = _extraer_numero_factura(texto)
    fecha_documento = _extraer_fecha_factura(texto, nombre_original)
    base = _extraer_base_factura(texto)
    tarifa = _extraer_tarifa_iva_factura(texto)
    iva = _extraer_valor_iva_factura(texto)
    total = _extraer_total_factura(texto)

    # Solo completa valores faltantes cuando la relación matemática es clara.
    if base is None and total is not None and iva is not None and total >= iva:
        base = round(total - iva, 2)

    if iva is None and base is not None and tarifa is not None:
        iva = round(base * tarifa / 100, 2)

    if total is None and base is not None and iva is not None:
        total = round(base + iva, 2)

    tercero = _buscar_tercero_por_identificacion_carga(empresa_id, nit)

    if tercero:
        proveedor = tercero.razon_social or ""
        resultado["tercero_existente"] = True
        resultado["tercero_id"] = tercero.id
    else:
        proveedor = _extraer_proveedor_factura(texto, nit)

    duplicado, observacion_duplicado = _es_factura_duplicada_carga(
        empresa_id,
        nit,
        numero_documento
    )

    resultado.update({
        "proveedor": proveedor,
        "identificacion": nit,
        "numero_documento": numero_documento,
        "fecha_documento": fecha_documento,
        "base_gravable": base,
        "tarifa_iva": tarifa,
        "valor_iva": iva,
        "total_documento": total,
        "duplicado": duplicado,
    })

    estado, faltantes, alertas = _estado_revision_carga(resultado)
    resultado["estado"] = estado
    resultado["alertas_validacion"] = alertas

    observaciones = []
    if observacion_duplicado:
        observaciones.append(observacion_duplicado)
    if faltantes:
        observaciones.append("Revisar: " + ", ".join(faltantes) + ".")
    observaciones.extend(alertas)
    if not tercero and nit:
        observaciones.append("El NIT no existe todavía en Terceros.")

    resultado["observaciones"] = " ".join(observaciones)
    return resultado


# ============================================================
# CARGA INTELIGENTE - PROCESAR LOTE
# ============================================================

@crv_bp.route(
    "/carga-inteligente/<lote_id>/procesar",
    methods=["POST"]
)
def carga_inteligente_procesar(
    lote_id
):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    manifest = (
        _leer_manifest_carga(
            empresa_id,
            lote_id
        )
    )

    if not manifest:
        flash(
            "El lote solicitado no existe.",
            "error"
        )

        return redirect(
            "/crv/carga-inteligente"
        )

    if (
        manifest.get("estado")
        != "PENDIENTE_REVISION"
    ):
        flash(
            "Este lote ya no está disponible para lectura.",
            "warning"
        )

        return redirect(
            (
                "/crv/carga-inteligente"
                f"?lote={lote_id}"
            )
        )

    carpeta = _ruta_lote_carga(
        empresa_id,
        lote_id
    )

    procesados = 0

    for item in manifest.get(
        "archivos",
        []
    ):
        ruta = (
            carpeta
            / str(
                item.get(
                    "nombre_guardado",
                    ""
                )
            )
        )

        if not ruta.exists():
            item[
                "estado"
            ] = "ARCHIVO_NO_ENCONTRADO"

            item[
                "observaciones"
            ] = (
                "El archivo físico no fue encontrado."
            )
            continue

        analisis = _analizar_factura_carga(
            empresa_id,
            ruta,
            item.get(
                "nombre_original",
                ""
            )
        )

        item.update(
            analisis
        )

        item[
            "fecha_lectura"
        ] = datetime.now().isoformat(
            timespec="seconds"
        )

        procesados += 1

    manifest[
        "fecha_ultima_lectura"
    ] = datetime.now().isoformat(
        timespec="seconds"
    )

    _guardar_manifest_carga(
        empresa_id,
        lote_id,
        manifest
    )

    flash(
        (
            f"Lectura finalizada para "
            f"{procesados} documento(s). "
            "Revisa los datos antes de registrar."
        ),
        "success"
    )

    return redirect(
        (
            "/crv/carga-inteligente"
            f"?lote={lote_id}"
        )
    )


# ============================================================
# CARGA INTELIGENTE - SUBIR / BANDEJA
# ============================================================

@crv_bp.route(
    "/carga-inteligente",
    methods=["GET", "POST"]
)
def carga_inteligente():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    if request.method == "POST":

        archivos = request.files.getlist(
            "documentos"
        )

        archivos = [
            archivo
            for archivo in archivos
            if archivo
            and str(
                archivo.filename or ""
            ).strip()
        ]

        if not archivos:

            flash(
                "Selecciona al menos una factura o documento.",
                "error"
            )

            return redirect(
                "/crv/carga-inteligente"
            )

        lote_id = (
            datetime.now().strftime(
                "%Y%m%d-%H%M%S"
            )
            + "-"
            + uuid.uuid4().hex[:8]
        )

        carpeta_lote = (
            _ruta_lote_carga(
                empresa_id,
                lote_id
            )
        )

        carpeta_lote.mkdir(
            parents=True,
            exist_ok=True
        )

        manifest = {
            "lote_id": lote_id,
            "empresa_id": int(
                empresa_id
            ),
            "empresa_nombre": (
                empresa_nombre
            ),
            "fecha_creacion": (
                datetime.now().isoformat(
                    timespec="seconds"
                )
            ),
            "usuario": str(
                session.get(
                    "usuario",
                    ""
                )
            ),
            "estado": "PENDIENTE_REVISION",
            "archivos": [],
        }

        errores = []

        for archivo in archivos:

            valido, mensaje = (
                _archivo_carga_valido(
                    archivo
                )
            )

            if not valido:

                errores.append(
                    (
                        f"{archivo.filename}: "
                        f"{mensaje}"
                    )
                )

                continue

            archivo.stream.seek(
                0,
                2
            )

            tamano = archivo.stream.tell()

            archivo.stream.seek(
                0
            )

            if (
                tamano
                > TAMANO_MAXIMO_ARCHIVO_CARGA
            ):

                errores.append(
                    (
                        f"{archivo.filename}: "
                        "supera el maximo de 25 MB."
                    )
                )

                continue

            nombre_original = str(
                archivo.filename
            ).strip()

            nombre_seguro = (
                secure_filename(
                    nombre_original
                )
                or "documento"
            )

            extension = Path(
                nombre_seguro
            ).suffix.lower()

            archivo_id = (
                uuid.uuid4().hex[:12]
            )

            nombre_guardado = (
                f"{archivo_id}"
                f"{extension}"
            )

            ruta_destino = (
                carpeta_lote
                / nombre_guardado
            )

            archivo.save(
                str(
                    ruta_destino
                )
            )

            manifest[
                "archivos"
            ].append({
                "archivo_id": archivo_id,
                "nombre_original": nombre_original,
                "nombre_guardado": nombre_guardado,
                "extension": extension,
                "tamano_bytes": tamano,
                "estado": "PENDIENTE_LECTURA",
                "proveedor": "",
                "identificacion": "",
                "numero_documento": "",
                "fecha_documento": "",
                "base_gravable": None,
                "tarifa_iva": None,
                "valor_iva": None,
                "total_documento": None,
                "duplicado": False,
                "observaciones": "",
            })

        if not manifest["archivos"]:

            try:
                carpeta_lote.rmdir()
            except Exception:
                pass

            flash(
                (
                    "No se pudo cargar ningun archivo. "
                    + (
                        " ".join(
                            errores
                        )
                        if errores
                        else ""
                    )
                ),
                "error"
            )

            return redirect(
                "/crv/carga-inteligente"
            )

        _guardar_manifest_carga(
            empresa_id,
            lote_id,
            manifest
        )

        flash(
            (
                f"Lote creado con "
                f"{len(manifest['archivos'])} "
                "documento(s). "
                "Aun no se ha creado ningun movimiento."
            ),
            "success"
        )

        if errores:
            flash(
                "Algunos archivos no se cargaron: "
                + " ".join(
                    errores
                ),
                "warning"
            )

        return redirect(
            (
                "/crv/carga-inteligente"
                f"?lote={lote_id}"
            )
        )

    lote_actual_id = request.args.get(
        "lote",
        ""
    ).strip()

    lote_actual = None

    if lote_actual_id:
        lote_actual = (
            _leer_manifest_carga(
                empresa_id,
                lote_actual_id
            )
        )

    lotes = (
        _listar_lotes_carga_inteligente(
            empresa_id
        )
    )

    return render_template(
        "contabilidad/crv_carga_inteligente.html",
        empresa_nombre=empresa_nombre,
        lote=lote_actual,
        lotes=lotes
    )


# ============================================================
# CARGA INTELIGENTE - GUARDAR REVISIÓN MANUAL
# ============================================================

@crv_bp.route(
    "/carga-inteligente/<lote_id>/archivo/<archivo_id>/revisar",
    methods=["POST"]
)
def carga_inteligente_revisar_archivo(lote_id, archivo_id):
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    manifest = _leer_manifest_carga(empresa_id, lote_id)

    if not manifest:
        flash("El lote solicitado no existe.", "error")
        return redirect("/crv/carga-inteligente")

    if manifest.get("estado") != "PENDIENTE_REVISION":
        flash("Este lote ya no puede modificarse.", "warning")
        return redirect(f"/crv/carga-inteligente?lote={lote_id}")

    item = next(
        (
            fila for fila in manifest.get("archivos", [])
            if str(fila.get("archivo_id", "")) == str(archivo_id)
        ),
        None
    )

    if not item:
        flash("El documento no existe dentro del lote.", "error")
        return redirect(f"/crv/carga-inteligente?lote={lote_id}")

    proveedor = request.form.get("proveedor", "").strip()
    identificacion = request.form.get("identificacion", "").strip()
    numero_documento = request.form.get("numero_documento", "").strip()
    fecha_documento = request.form.get("fecha_documento", "").strip()

    base = convertir_numero(request.form.get("base_gravable"), None)
    tarifa = convertir_porcentaje(request.form.get("tarifa_iva"), None)
    iva = convertir_numero(request.form.get("valor_iva"), None)
    total = convertir_numero(request.form.get("total_documento"), None)

    # Fecha debe quedar en ISO YYYY-MM-DD o vacía.
    if fecha_documento and not convertir_fecha(fecha_documento):
        flash("La fecha del documento no es válida.", "error")
        return redirect(f"/crv/carga-inteligente?lote={lote_id}")

    tercero = _buscar_tercero_por_identificacion_carga(
        empresa_id,
        identificacion
    )

    duplicado, observacion_duplicado = _es_factura_duplicada_carga(
        empresa_id,
        identificacion,
        numero_documento
    )

    item.update({
        "proveedor": proveedor,
        "identificacion": identificacion,
        "numero_documento": numero_documento,
        "fecha_documento": fecha_documento,
        "base_gravable": base,
        "tarifa_iva": tarifa,
        "valor_iva": iva,
        "total_documento": total,
        "tercero_existente": bool(tercero),
        "tercero_id": tercero.id if tercero else None,
        "duplicado": duplicado,
        "revision_manual": True,
        "fecha_revision_manual": datetime.now().isoformat(timespec="seconds"),
    })

    estado, faltantes, alertas = _estado_revision_carga(item)
    item["estado"] = estado
    item["alertas_validacion"] = alertas

    observaciones = []
    if observacion_duplicado:
        observaciones.append(observacion_duplicado)
    if faltantes:
        observaciones.append("Revisar: " + ", ".join(faltantes) + ".")
    observaciones.extend(alertas)
    if not tercero and identificacion:
        observaciones.append("El NIT no existe todavía en Terceros.")

    item["observaciones"] = " ".join(observaciones)

    _guardar_manifest_carga(empresa_id, lote_id, manifest)

    if estado == "LISTO_PARA_APROBAR":
        flash(
            "Revisión guardada. Los datos cuadran y el documento queda listo para una futura aprobación.",
            "success"
        )
    elif estado == "POSIBLE_DUPLICADO":
        flash(
            "Revisión guardada, pero el documento sigue marcado como posible duplicado.",
            "warning"
        )
    else:
        flash(
            "Revisión guardada. Aún hay campos o valores que deben corregirse.",
            "warning"
        )

    return redirect(f"/crv/carga-inteligente?lote={lote_id}")


# ============================================================
# CARGA INTELIGENTE - QUITAR ARCHIVO PENDIENTE
# ============================================================

@crv_bp.route(
    "/carga-inteligente/<lote_id>/archivo/<archivo_id>/eliminar",
    methods=["POST"]
)
def carga_inteligente_eliminar_archivo(
    lote_id,
    archivo_id
):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    manifest = (
        _leer_manifest_carga(
            empresa_id,
            lote_id
        )
    )

    if not manifest:

        flash(
            "El lote solicitado no existe.",
            "error"
        )

        return redirect(
            "/crv/carga-inteligente"
        )

    if (
        manifest.get("estado")
        != "PENDIENTE_REVISION"
    ):

        flash(
            "Este lote ya no puede modificarse.",
            "warning"
        )

        return redirect(
            (
                "/crv/carga-inteligente"
                f"?lote={lote_id}"
            )
        )

    carpeta = _ruta_lote_carga(
        empresa_id,
        lote_id
    )

    nuevos_archivos = []
    eliminado = False

    for item in manifest.get(
        "archivos",
        []
    ):

        if (
            str(
                item.get(
                    "archivo_id",
                    ""
                )
            )
            == str(
                archivo_id
            )
        ):

            ruta = (
                carpeta
                / str(
                    item.get(
                        "nombre_guardado",
                        ""
                    )
                )
            )

            if ruta.exists():

                try:
                    ruta.unlink()
                except Exception:
                    pass

            eliminado = True
            continue

        nuevos_archivos.append(
            item
        )

    manifest[
        "archivos"
    ] = nuevos_archivos

    _guardar_manifest_carga(
        empresa_id,
        lote_id,
        manifest
    )

    if eliminado:

        flash(
            "Documento retirado del lote.",
            "success"
        )

    return redirect(
        (
            "/crv/carga-inteligente"
            f"?lote={lote_id}"
        )
    )


# ============================================================
# CARGA INTELIGENTE - ELIMINAR LOTE PENDIENTE
# ============================================================

@crv_bp.route(
    "/carga-inteligente/<lote_id>/eliminar",
    methods=["POST"]
)
def carga_inteligente_eliminar_lote(
    lote_id
):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    manifest = (
        _leer_manifest_carga(
            empresa_id,
            lote_id
        )
    )

    if not manifest:

        return redirect(
            "/crv/carga-inteligente"
        )

    if (
        manifest.get("estado")
        != "PENDIENTE_REVISION"
    ):

        flash(
            "Solo se pueden eliminar lotes pendientes.",
            "warning"
        )

        return redirect(
            (
                "/crv/carga-inteligente"
                f"?lote={lote_id}"
            )
        )

    carpeta = _ruta_lote_carga(
        empresa_id,
        lote_id
    )

    try:
        for archivo in carpeta.iterdir():
            if archivo.is_file():
                archivo.unlink()

        carpeta.rmdir()

        flash(
            "Lote pendiente eliminado correctamente.",
            "success"
        )

    except Exception as error:

        flash(
            (
                "No fue posible eliminar el lote: "
                f"{error}"
            ),
            "error"
        )

    return redirect(
        "/crv/carga-inteligente"
    )


# ============================================================
# DASHBOARD CRV
# ============================================================

@crv_bp.route("/")
@crv_bp.route("/dashboard")
def dashboard():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    mes = request.args.get(
        "mes",
        ""
    ).strip()

    query = CRVMovimiento.query.filter_by(
        empresa_id=empresa_id
    )

    if mes:

        try:

            anio_texto, mes_texto = mes.split("-")

            anio = int(anio_texto)
            numero_mes = int(mes_texto)

            query = query.filter(
                db.extract(
                    "year",
                    CRVMovimiento.fecha
                ) == anio,

                db.extract(
                    "month",
                    CRVMovimiento.fecha
                ) == numero_mes
            )

        except (ValueError, TypeError):
            pass

    movimientos = query.order_by(
        CRVMovimiento.fecha.desc(),
        CRVMovimiento.id.desc()
    ).all()

    # ========================================================
    # CÁLCULO FINANCIERO CRV
    # ========================================================

    ingresos = 0
    compras = 0

    gasto_general = 0
    gasto_nomina_afiliados = 0
    gasto_nomina_obra = 0
    gasto_anticipos = 0
    gasto_prestamos = 0
    gasto_otros = 0

    for movimiento in movimientos:

        tipo = str(
            movimiento.tipo or ""
        ).strip().upper()

        valor = float(
            movimiento.valor_neto
            or movimiento.valor
            or 0
        )

        if tipo == "INGRESO":

            ingresos += valor

        elif tipo == "COMPRA":

            compras += valor

        elif tipo == "GASTO":

            gasto_general += valor

        elif tipo in (
            "NOMINA AFILIADOS",
            "NÓMINA AFILIADOS"
        ):

            gasto_nomina_afiliados += valor

        elif tipo in (
            "NOMINA OBRA",
            "NÓMINA OBRA"
        ):

            gasto_nomina_obra += valor

        elif tipo == "ANTICIPO":

            gasto_anticipos += valor

        elif tipo in (
            "PRESTAMO EMPLEADO",
            "PRÉSTAMO EMPLEADO"
        ):

            gasto_prestamos += valor

        elif tipo in (
            "TRANSFERENCIA",
            "PAGO CUENTA POR PAGAR",
            "SALDO INICIAL CUENTA",
            "COBRO CUENTA POR COBRAR",
        ):

            # Transferencias entre cuentas propias y pagos de CxP
            # no vuelven a afectar el resultado. La compra ya se
            # reconoció cuando se registró la obligación/factura.
            pass

        else:

            gasto_otros += valor

    gastos_totales = (
        gasto_general
        + gasto_nomina_afiliados
        + gasto_nomina_obra
        + gasto_anticipos
        + gasto_prestamos
        + gasto_otros
    )

    egresos_totales = (
        compras
        + gastos_totales
    )

    utilidad = (
        ingresos
        - egresos_totales
    )

    cantidad_terceros = CRVTercero.query.filter_by(
        empresa_id=empresa_id,
        activo=True
    ).count()

    ultimos_movimientos = movimientos[:10]

    return render_template(
        "contabilidad/crv_dashboard.html",

        empresa_nombre=empresa_nombre,
        mes_actual=mes,

        ingresos=ingresos,
        compras=compras,

        gasto_general=gasto_general,
        gasto_nomina_afiliados=gasto_nomina_afiliados,
        gasto_nomina_obra=gasto_nomina_obra,
        gasto_anticipos=gasto_anticipos,
        gasto_prestamos=gasto_prestamos,
        gasto_otros=gasto_otros,

        gastos_totales=gastos_totales,
        egresos_totales=egresos_totales,
        utilidad=utilidad,

        cantidad_movimientos=len(movimientos),
        cantidad_terceros=cantidad_terceros,

        ultimos_movimientos=ultimos_movimientos
    )


# ============================================================
# LISTADO DE MOVIMIENTOS
# ============================================================

@crv_bp.route("/movimientos")
def movimientos():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    mes = request.args.get(
        "mes",
        ""
    ).strip()

    sede_id = request.args.get(
        "sede_id",
        ""
    ).strip()

    tipo = request.args.get(
        "tipo",
        ""
    ).strip().upper()

    tercero = request.args.get(
        "tercero",
        ""
    ).strip()

    query = CRVMovimiento.query.filter_by(
        empresa_id=empresa_id
    )

    if sede_id.isdigit():

        query = query.filter_by(
            sede_id=int(sede_id)
        )

    if tipo:

        query = query.filter(
            CRVMovimiento.tipo == tipo
        )

    if tercero:

        query = query.filter(
            CRVMovimiento.tercero_nombre.ilike(
                f"%{tercero}%"
            )
        )

    if mes:

        try:

            anio_texto, mes_texto = mes.split("-")

            query = query.filter(
                db.extract(
                    "year",
                    CRVMovimiento.fecha
                ) == int(anio_texto),

                db.extract(
                    "month",
                    CRVMovimiento.fecha
                ) == int(mes_texto)
            )

        except (ValueError, TypeError):
            pass

    registros = query.order_by(
        CRVMovimiento.fecha.desc(),
        CRVMovimiento.id.desc()
    ).all()

    sedes = Sede.query.filter_by(
        empresa_id=empresa_id
    ).order_by(
        Sede.nombre.asc()
    ).all()

    return render_template(
        "contabilidad/crv_movimientos.html",

        empresa_nombre=empresa_nombre,

        movimientos=registros,

        sedes=sedes,

        mes_actual=mes,
        sede_actual=sede_id,
        tipo_actual=tipo,
        tercero_actual=tercero
    )


# ============================================================
# NUEVO MOVIMIENTO
# ============================================================

@crv_bp.route(
    "/movimientos/nuevo",
    methods=["GET", "POST"]
)
def nuevo_movimiento():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    sedes = Sede.query.filter_by(
        empresa_id=empresa_id
    ).order_by(
        Sede.nombre.asc()
    ).all()

    terceros = CRVTercero.query.filter_by(
        empresa_id=empresa_id,
        activo=True
    ).order_by(
        CRVTercero.razon_social.asc()
    ).all()

    if request.method == "POST":

        fecha = convertir_fecha(
            request.form.get("fecha")
        )

        if not fecha:

            flash(
                "Debes indicar la fecha contable del movimiento.",
                "error"
            )

            return render_template(
                "contabilidad/crv_nuevo_movimiento.html",
                empresa_nombre=empresa_nombre,
                sedes=sedes,
                terceros=terceros,
                movimiento=None,
                modo="nuevo"
            )

        tipo = request.form.get(
            "tipo",
            ""
        ).strip().upper()

        categoria = request.form.get(
            "categoria",
            ""
        ).strip()

        concepto = request.form.get(
            "concepto",
            ""
        ).strip()

        iva_datos = calcular_iva_formulario(
            request.form
        )

        valor = iva_datos[
            "total_documento"
        ]

        if not tipo:

            flash(
                "Debes seleccionar el tipo de movimiento.",
                "error"
            )

            return redirect(
                "/crv/movimientos/nuevo"
            )

        if not categoria:
            categoria = "OTRO"

        if not concepto:

            flash(
                "Debes escribir el concepto del movimiento.",
                "error"
            )

            return redirect(
                "/crv/movimientos/nuevo"
            )

        tercero_obj = crear_o_actualizar_tercero_desde_formulario(
            empresa_id,
            request.form
        )

        sede_id_texto = request.form.get(
            "sede_id",
            ""
        ).strip()

        sede_id = (
            int(sede_id_texto)
            if sede_id_texto.isdigit()
            else None
        )

        descuento = convertir_numero(
            request.form.get("descuento")
        )

        retencion = convertir_numero(
            request.form.get("retencion")
        )

        otros_descuentos = convertir_numero(
            request.form.get(
                "otros_descuentos"
            )
        )

        valor_neto = max(
            0,
            valor
            - descuento
            - retencion
            - otros_descuentos
        )

        movimiento = CRVMovimiento(

            empresa_id=empresa_id,
            sede_id=sede_id,

            fecha=fecha,

            tipo=tipo,

            categoria=categoria,

            subcategoria=request.form.get(
                "subcategoria",
                ""
            ).strip(),

            concepto=concepto,

            descripcion=request.form.get(
                "descripcion",
                ""
            ).strip(),

            tercero_id=(
                tercero_obj.id
                if tercero_obj
                else None
            ),

            tercero_nombre=(
                tercero_obj.razon_social
                if tercero_obj
                else request.form.get(
                    "tercero_nombre",
                    ""
                ).strip()
            ),

            tercero_identificacion=(
                tercero_obj.identificacion
                if tercero_obj
                else request.form.get(
                    "tercero_identificacion",
                    ""
                ).strip()
            ),

            tercero_direccion=(
                tercero_obj.direccion
                if tercero_obj
                else request.form.get(
                    "tercero_direccion",
                    ""
                ).strip()
            ),

            tercero_email=(
                tercero_obj.email
                if tercero_obj
                else request.form.get(
                    "tercero_email",
                    ""
                ).strip()
            ),

            tercero_telefono=(
                (
                    tercero_obj.celular
                    or tercero_obj.telefono
                )
                if tercero_obj
                else request.form.get(
                    "tercero_celular",
                    ""
                ).strip()
            ),

            tercero_contacto=(
                tercero_obj.contacto
                if tercero_obj
                else request.form.get(
                    "tercero_contacto",
                    ""
                ).strip()
            ),

            tipo_documento=request.form.get(
                "tipo_documento",
                ""
            ).strip(),

            numero_documento=request.form.get(
                "numero_documento",
                ""
            ).strip(),

            numero_soporte=request.form.get(
                "numero_soporte",
                ""
            ).strip(),

            fecha_documento=convertir_fecha(
                request.form.get(
                    "fecha_documento"
                )
            ),

            fecha_vencimiento=convertir_fecha(
                request.form.get(
                    "fecha_vencimiento"
                )
            ),

            base_gravable=iva_datos["base_gravable"],
            tipo_iva=iva_datos["tipo_iva"],
            tarifa_iva=iva_datos["tarifa_iva"],
            valor_iva=iva_datos["valor_iva"],
            valor=valor,
            descuento=descuento,
            retencion=retencion,
            otros_descuentos=otros_descuentos,
            valor_neto=valor_neto,

            forma_pago=request.form.get(
                "forma_pago",
                ""
            ).strip(),

            cuenta=request.form.get(
                "cuenta",
                ""
            ).strip(),

            cuenta_destino=request.form.get(
                "cuenta_destino",
                ""
            ).strip(),

            beneficiario=request.form.get(
                "beneficiario",
                ""
            ).strip(),

            referencia_pago=request.form.get(
                "referencia_pago",
                ""
            ).strip(),

            proyecto=request.form.get(
                "proyecto",
                ""
            ).strip(),

            centro_costo=request.form.get(
                "centro_costo",
                ""
            ).strip(),

            obra=request.form.get(
                "obra",
                ""
            ).strip(),

            observaciones=request.form.get(
                "observaciones",
                ""
            ).strip(),

            revisado=False,

            estado="ACTIVO",

            origen_registro="SISTEMA",

            usuario_registro=str(
                session.get(
                    "usuario",
                    ""
                )
            )
        )

        _asignar_comprobante_permanente(
            movimiento
        )

        db.session.add(movimiento)

        db.session.commit()

        flash(
            "Movimiento CRV registrado correctamente.",
            "success"
        )

        return redirect(
            "/crv/movimientos"
        )

    return render_template(
        "contabilidad/crv_nuevo_movimiento.html",

        empresa_nombre=empresa_nombre,

        sedes=sedes,

        terceros=terceros,

        movimiento=None,

        modo="nuevo"
    )


# ============================================================
# EDITAR MOVIMIENTO
# ============================================================

@crv_bp.route(
    "/movimientos/editar/<int:movimiento_id>",
    methods=["GET", "POST"]
)
def editar_movimiento(movimiento_id):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    movimiento = CRVMovimiento.query.get_or_404(
        movimiento_id
    )

    if movimiento.empresa_id != empresa_id:

        return redirect(
            "/crv/movimientos"
        )

    if movimiento.origen_registro in (
        "CUENTAS_POR_PAGAR",
        "PAGO_CUENTA_POR_PAGAR"
    ):

        flash(
            "Este movimiento fue generado por Cuentas por pagar y debe administrarse desde ese módulo para conservar los saldos correctos.",
            "warning"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    sedes = Sede.query.filter_by(
        empresa_id=empresa_id
    ).order_by(
        Sede.nombre.asc()
    ).all()

    terceros = CRVTercero.query.filter_by(
        empresa_id=empresa_id,
        activo=True
    ).order_by(
        CRVTercero.razon_social.asc()
    ).all()

    if request.method == "POST":

        fecha = convertir_fecha(
            request.form.get("fecha")
        )

        if not fecha:

            flash(
                "La fecha contable es obligatoria.",
                "error"
            )

            return redirect(
                f"/crv/movimientos/editar/{movimiento_id}"
            )

        tercero_obj = crear_o_actualizar_tercero_desde_formulario(
            empresa_id,
            request.form
        )

        movimiento.fecha = fecha

        movimiento.tipo = request.form.get(
            "tipo",
            ""
        ).strip().upper()

        movimiento.categoria = request.form.get(
            "categoria",
            "OTRO"
        ).strip()

        movimiento.subcategoria = request.form.get(
            "subcategoria",
            ""
        ).strip()

        movimiento.concepto = request.form.get(
            "concepto",
            ""
        ).strip()

        movimiento.descripcion = request.form.get(
            "descripcion",
            ""
        ).strip()

        sede_id_texto = request.form.get(
            "sede_id",
            ""
        ).strip()

        movimiento.sede_id = (
            int(sede_id_texto)
            if sede_id_texto.isdigit()
            else None
        )

        if tercero_obj:

            movimiento.tercero_id = tercero_obj.id

            movimiento.tercero_nombre = tercero_obj.razon_social

            movimiento.tercero_identificacion = tercero_obj.identificacion

            movimiento.tercero_direccion = tercero_obj.direccion

            movimiento.tercero_email = tercero_obj.email

            movimiento.tercero_telefono = (
                tercero_obj.celular
                or tercero_obj.telefono
            )

            movimiento.tercero_contacto = tercero_obj.contacto

        movimiento.tipo_documento = request.form.get(
            "tipo_documento",
            ""
        ).strip()

        movimiento.numero_documento = request.form.get(
            "numero_documento",
            ""
        ).strip()

        movimiento.numero_soporte = request.form.get(
            "numero_soporte",
            ""
        ).strip()

        movimiento.fecha_documento = convertir_fecha(
            request.form.get(
                "fecha_documento"
            )
        )

        movimiento.fecha_vencimiento = convertir_fecha(
            request.form.get(
                "fecha_vencimiento"
            )
        )

        iva_datos = calcular_iva_formulario(
            request.form
        )

        movimiento.base_gravable = iva_datos[
            "base_gravable"
        ]
        movimiento.tipo_iva = iva_datos[
            "tipo_iva"
        ]
        movimiento.tarifa_iva = iva_datos[
            "tarifa_iva"
        ]
        movimiento.valor_iva = iva_datos[
            "valor_iva"
        ]
        movimiento.valor = iva_datos[
            "total_documento"
        ]

        movimiento.descuento = convertir_numero(
            request.form.get("descuento")
        )

        movimiento.retencion = convertir_numero(
            request.form.get("retencion")
        )

        movimiento.otros_descuentos = convertir_numero(
            request.form.get(
                "otros_descuentos"
            )
        )

        movimiento.recalcular_valor_neto()

        movimiento.forma_pago = request.form.get(
            "forma_pago",
            ""
        ).strip()

        movimiento.cuenta = request.form.get(
            "cuenta",
            ""
        ).strip()

        movimiento.cuenta_destino = request.form.get(
            "cuenta_destino",
            ""
        ).strip()

        movimiento.beneficiario = request.form.get(
            "beneficiario",
            ""
        ).strip()

        movimiento.referencia_pago = request.form.get(
            "referencia_pago",
            ""
        ).strip()

        movimiento.proyecto = request.form.get(
            "proyecto",
            ""
        ).strip()

        movimiento.centro_costo = request.form.get(
            "centro_costo",
            ""
        ).strip()

        movimiento.obra = request.form.get(
            "obra",
            ""
        ).strip()

        movimiento.observaciones = request.form.get(
            "observaciones",
            ""
        ).strip()

        db.session.commit()

        flash(
            "Movimiento actualizado correctamente.",
            "success"
        )

        return redirect(
            "/crv/movimientos"
        )

    return render_template(
        "contabilidad/crv_nuevo_movimiento.html",

        empresa_nombre=empresa_nombre,

        sedes=sedes,

        terceros=terceros,

        movimiento=movimiento,

        modo="editar"
    )


# ============================================================
# ELIMINAR MOVIMIENTO
# ============================================================

@crv_bp.route(
    "/movimientos/eliminar/<int:movimiento_id>",
    methods=["POST"]
)
def eliminar_movimiento(movimiento_id):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    movimiento = CRVMovimiento.query.get_or_404(
        movimiento_id
    )

    if movimiento.empresa_id != empresa_id:

        return redirect(
            "/crv/movimientos"
        )

    if movimiento.origen_registro in (
        "CUENTAS_POR_PAGAR",
        "PAGO_CUENTA_POR_PAGAR"
    ):

        flash(
            "Este movimiento fue generado por Cuentas por pagar y no debe eliminarse desde Movimientos. Usa el módulo de Cuentas por pagar.",
            "warning"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    db.session.delete(movimiento)

    db.session.commit()

    flash(
        "Movimiento eliminado correctamente.",
        "success"
    )

    return redirect(
        "/crv/movimientos"
    )


# ============================================================
# TERCEROS
# ============================================================

@crv_bp.route(
    "/terceros",
    methods=["GET", "POST"]
)
def terceros():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    if request.method == "POST":

        identificacion = request.form.get(
            "identificacion",
            ""
        ).strip()

        razon_social = request.form.get(
            "razon_social",
            ""
        ).strip()

        if not identificacion:

            flash(
                "La identificación del tercero es obligatoria.",
                "error"
            )

            return redirect(
                "/crv/terceros"
            )

        if not razon_social:

            flash(
                "El nombre o razón social es obligatorio.",
                "error"
            )

            return redirect(
                "/crv/terceros"
            )

        existente = CRVTercero.query.filter_by(
            empresa_id=empresa_id,
            identificacion=identificacion
        ).first()

        if existente:

            flash(
                "Ya existe un tercero con esa identificación.",
                "error"
            )

            return redirect(
                "/crv/terceros"
            )

        tercero = CRVTercero(

            empresa_id=empresa_id,

            tipo_tercero=request.form.get(
                "tipo_tercero",
                "TERCERO"
            ).strip().upper(),

            tipo_documento=request.form.get(
                "tipo_documento",
                ""
            ).strip(),

            identificacion=identificacion,

            razon_social=razon_social,

            nombre_comercial=request.form.get(
                "nombre_comercial",
                ""
            ).strip(),

            direccion=request.form.get(
                "direccion",
                ""
            ).strip(),

            ciudad=request.form.get(
                "ciudad",
                ""
            ).strip(),

            departamento=request.form.get(
                "departamento",
                ""
            ).strip(),

            telefono=request.form.get(
                "telefono",
                ""
            ).strip(),

            celular=request.form.get(
                "celular",
                ""
            ).strip(),

            email=request.form.get(
                "email",
                ""
            ).strip(),

            contacto=request.form.get(
                "contacto",
                ""
            ).strip(),

            banco=request.form.get(
                "banco",
                ""
            ).strip(),

            tipo_cuenta=request.form.get(
                "tipo_cuenta",
                ""
            ).strip(),

            numero_cuenta=request.form.get(
                "numero_cuenta",
                ""
            ).strip(),

            titular_cuenta=request.form.get(
                "titular_cuenta",
                ""
            ).strip(),

            observaciones=request.form.get(
                "observaciones",
                ""
            ).strip(),

            activo=True
        )

        db.session.add(tercero)

        db.session.commit()

        flash(
            "Tercero creado correctamente.",
            "success"
        )

        return redirect(
            "/crv/terceros"
        )

    busqueda = request.args.get(
        "buscar",
        ""
    ).strip()

    query = CRVTercero.query.filter_by(
        empresa_id=empresa_id
    )

    if busqueda:

        patron = f"%{busqueda}%"

        query = query.filter(
            db.or_(
                CRVTercero.razon_social.ilike(
                    patron
                ),

                CRVTercero.identificacion.ilike(
                    patron
                ),

                CRVTercero.nombre_comercial.ilike(
                    patron
                )
            )
        )

    registros = query.order_by(
        CRVTercero.razon_social.asc()
    ).all()

    return render_template(
        "contabilidad/crv_terceros.html",

        empresa_nombre=empresa_nombre,

        terceros=registros,

        busqueda=busqueda
    )


# ============================================================
# ACTIVAR / INACTIVAR TERCERO
# ============================================================

@crv_bp.route(
    "/terceros/estado/<int:tercero_id>",
    methods=["POST"]
)
def cambiar_estado_tercero(tercero_id):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    tercero = CRVTercero.query.get_or_404(
        tercero_id
    )

    if tercero.empresa_id != empresa_id:

        return redirect(
            "/crv/terceros"
        )

    tercero.activo = not tercero.activo

    db.session.commit()

    return redirect(
        "/crv/terceros"
    )
# ============================================================
# EDITAR TERCERO
# ============================================================

@crv_bp.route(
    "/terceros/editar/<int:tercero_id>",
    methods=["POST"]
)
def editar_tercero(tercero_id):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    tercero = CRVTercero.query.get_or_404(
        tercero_id
    )

    # Seguridad:
    # solo permite modificar terceros pertenecientes
    # a la empresa CRV actualmente seleccionada.
    if tercero.empresa_id != empresa_id:

        flash(
            "No tienes permiso para modificar este tercero.",
            "error"
        )

        return redirect(
            "/crv/terceros"
        )

    identificacion = request.form.get(
        "identificacion",
        ""
    ).strip()

    razon_social = request.form.get(
        "razon_social",
        ""
    ).strip()

    # --------------------------------------------------------
    # VALIDACIONES
    # --------------------------------------------------------

    if not identificacion:

        flash(
            "La identificación del tercero es obligatoria.",
            "error"
        )

        return redirect(
            "/crv/terceros"
        )

    if not razon_social:

        flash(
            "El nombre o razón social es obligatorio.",
            "error"
        )

        return redirect(
            "/crv/terceros"
        )

    # Verificar que la nueva identificación no pertenezca
    # a otro tercero de la misma empresa.
    duplicado = CRVTercero.query.filter(
        CRVTercero.empresa_id == empresa_id,
        CRVTercero.identificacion == identificacion,
        CRVTercero.id != tercero.id
    ).first()

    if duplicado:

        flash(
            "Ya existe otro tercero con esa identificación.",
            "error"
        )

        return redirect(
            "/crv/terceros"
        )

    # --------------------------------------------------------
    # INFORMACIÓN GENERAL
    # --------------------------------------------------------

    tercero.tipo_tercero = request.form.get(
        "tipo_tercero",
        "TERCERO"
    ).strip().upper() or "TERCERO"

    tercero.tipo_documento = request.form.get(
        "tipo_documento",
        ""
    ).strip()

    tercero.identificacion = identificacion

    tercero.razon_social = razon_social

    tercero.nombre_comercial = request.form.get(
        "nombre_comercial",
        ""
    ).strip()

    # --------------------------------------------------------
    # INFORMACIÓN DE CONTACTO
    # --------------------------------------------------------

    tercero.direccion = request.form.get(
        "direccion",
        ""
    ).strip()

    tercero.ciudad = request.form.get(
        "ciudad",
        ""
    ).strip()

    tercero.departamento = request.form.get(
        "departamento",
        ""
    ).strip()

    tercero.telefono = request.form.get(
        "telefono",
        ""
    ).strip()

    tercero.celular = request.form.get(
        "celular",
        ""
    ).strip()

    tercero.email = request.form.get(
        "email",
        ""
    ).strip()

    tercero.contacto = request.form.get(
        "contacto",
        ""
    ).strip()

    # --------------------------------------------------------
    # INFORMACIÓN BANCARIA
    # --------------------------------------------------------

    tercero.banco = request.form.get(
        "banco",
        ""
    ).strip()

    tercero.tipo_cuenta = request.form.get(
        "tipo_cuenta",
        ""
    ).strip()

    tercero.numero_cuenta = request.form.get(
        "numero_cuenta",
        ""
    ).strip()

    tercero.titular_cuenta = request.form.get(
        "titular_cuenta",
        ""
    ).strip()

    # --------------------------------------------------------
    # INFORMACIÓN ADICIONAL
    # --------------------------------------------------------

    tercero.observaciones = request.form.get(
        "observaciones",
        ""
    ).strip()

    # --------------------------------------------------------
    # GUARDAR
    # --------------------------------------------------------

    try:

        db.session.commit()

        flash(
            "Datos del tercero actualizados correctamente.",
            "success"
        )

    except Exception:

        db.session.rollback()

        flash(
            "No fue posible actualizar el tercero.",
            "error"
        )

    return redirect(
        "/crv/terceros"
    )

# ============================================================
# CUENTAS POR PAGAR - LISTADO / CREACIÓN
# ============================================================

@crv_bp.route(
    "/cuentas-por-pagar",
    methods=["GET", "POST"]
)
def cuentas_por_pagar():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get("empresa_nombre", "CRV CONSTRUCTORA")

    proveedores = CRVTercero.query.filter(
        CRVTercero.empresa_id == empresa_id,
        CRVTercero.activo.is_(True),
        CRVTercero.tipo_tercero.in_(
            ["PROVEEDOR", "CONTRATISTA", "ENTIDAD", "TERCERO"]
        )
    ).order_by(CRVTercero.razon_social.asc()).all()

    sedes = Sede.query.filter_by(
        empresa_id=empresa_id
    ).order_by(Sede.nombre.asc()).all()

    if request.method == "POST":

        proveedor_id_texto = request.form.get("proveedor_id", "").strip()

        if not proveedor_id_texto.isdigit():
            flash("Debes seleccionar un proveedor válido.", "error")
            return redirect("/crv/cuentas-por-pagar")

        proveedor = CRVTercero.query.get_or_404(int(proveedor_id_texto))

        if proveedor.empresa_id != empresa_id:
            flash("El proveedor seleccionado no pertenece a CRV.", "error")
            return redirect("/crv/cuentas-por-pagar")

        numero_documento = request.form.get("numero_documento", "").strip()
        fecha_documento = convertir_fecha(request.form.get("fecha_documento"))
        concepto = request.form.get("concepto", "").strip()
        iva_datos = calcular_iva_formulario(request.form)
        valor_bruto = iva_datos["total_documento"]

        if not numero_documento:
            flash("El número de documento o factura es obligatorio.", "error")
            return redirect("/crv/cuentas-por-pagar")

        if not fecha_documento:
            flash("La fecha de la factura o documento es obligatoria.", "error")
            return redirect("/crv/cuentas-por-pagar")

        if not concepto:
            flash("El concepto de la obligación es obligatorio.", "error")
            return redirect("/crv/cuentas-por-pagar")

        if valor_bruto <= 0:
            flash("El valor de la factura debe ser mayor a cero.", "error")
            return redirect("/crv/cuentas-por-pagar")

        duplicada = CRVCuentaPagar.query.filter_by(
            empresa_id=empresa_id,
            proveedor_id=proveedor.id,
            numero_documento=numero_documento
        ).first()

        if duplicada:
            flash(
                "Ya existe una cuenta por pagar de este proveedor con el mismo número de documento.",
                "error"
            )
            return redirect(f"/crv/cuentas-por-pagar/{duplicada.id}")

        sede_id_texto = request.form.get("sede_id", "").strip()
        sede_id = int(sede_id_texto) if sede_id_texto.isdigit() else None

        descuento = convertir_numero(request.form.get("descuento"))
        retencion = convertir_numero(request.form.get("retencion"))
        otros_descuentos = convertir_numero(request.form.get("otros_descuentos"))

        cuenta = CRVCuentaPagar(
            empresa_id=empresa_id,
            proveedor_id=proveedor.id,
            sede_id=sede_id,
            tipo_documento=request.form.get("tipo_documento", "").strip(),
            numero_documento=numero_documento,
            fecha_documento=fecha_documento,
            fecha_vencimiento=convertir_fecha(
                request.form.get("fecha_vencimiento")
            ),
            concepto=concepto,
            descripcion=request.form.get("descripcion", "").strip(),
            base_gravable=iva_datos["base_gravable"],
            tipo_iva=iva_datos["tipo_iva"],
            tarifa_iva=iva_datos["tarifa_iva"],
            valor_iva=iva_datos["valor_iva"],
            valor_bruto=valor_bruto,
            descuento=descuento,
            retencion=retencion,
            otros_descuentos=otros_descuentos,
            categoria=request.form.get("categoria", "COMPRA").strip() or "COMPRA",
            subcategoria=request.form.get("subcategoria", "").strip(),
            proyecto=request.form.get("proyecto", "").strip(),
            centro_costo=request.form.get("centro_costo", "").strip(),
            obra=request.form.get("obra", "").strip(),
            observaciones=request.form.get("observaciones", "").strip(),
            origen_registro="CUENTAS_POR_PAGAR",
            usuario_registro=str(session.get("usuario", ""))
        )

        cuenta.valor_neto = cuenta.valor_calculado
        cuenta.total_pagado = 0
        cuenta.saldo_pendiente = cuenta.valor_neto
        cuenta.recalcular_estado()

        try:
            db.session.add(cuenta)
            db.session.flush()

            movimiento_compra = CRVMovimiento(
                empresa_id=empresa_id,
                sede_id=sede_id,
                fecha=fecha_documento,
                tipo="COMPRA",
                categoria=cuenta.categoria or "COMPRA",
                subcategoria=cuenta.subcategoria or "",
                concepto=cuenta.concepto,
                descripcion=cuenta.descripcion or f"Cuenta por pagar {cuenta.numero_documento}",
                tercero_id=proveedor.id,
                tercero_nombre=proveedor.razon_social,
                tercero_identificacion=proveedor.identificacion,
                tercero_direccion=proveedor.direccion or "",
                tercero_email=proveedor.email or "",
                tercero_telefono=proveedor.celular or proveedor.telefono or "",
                tercero_contacto=proveedor.contacto or "",
                tipo_documento=cuenta.tipo_documento or "",
                numero_documento=cuenta.numero_documento,
                numero_soporte=f"CXP-{cuenta.id}",
                fecha_documento=cuenta.fecha_documento,
                fecha_vencimiento=cuenta.fecha_vencimiento,
                base_gravable=cuenta.base_gravable or 0,
                tipo_iva=cuenta.tipo_iva or "NO_APLICA",
                tarifa_iva=cuenta.tarifa_iva or 0,
                valor_iva=cuenta.valor_iva or 0,
                valor=cuenta.valor_bruto or 0,
                descuento=cuenta.descuento or 0,
                retencion=cuenta.retencion or 0,
                otros_descuentos=cuenta.otros_descuentos or 0,
                valor_neto=cuenta.valor_neto or 0,
                forma_pago="CREDITO",
                cuenta="CUENTAS POR PAGAR",
                cuenta_destino="",
                beneficiario=proveedor.razon_social,
                referencia_pago="",
                proyecto=cuenta.proyecto or "",
                centro_costo=cuenta.centro_costo or "",
                obra=cuenta.obra or "",
                observaciones=(
                    "Registro automático generado al crear "
                    f"la cuenta por pagar CXP-{cuenta.id}."
                ),
                revisado=False,
                estado="ACTIVO",
                origen_registro="CUENTAS_POR_PAGAR",
                usuario_registro=str(session.get("usuario", ""))
            )

            _asignar_comprobante_permanente(
                movimiento_compra
            )

            db.session.add(movimiento_compra)
            db.session.commit()

            flash(
                "Cuenta por pagar registrada correctamente. La compra también fue registrada en movimientos.",
                "success"
            )

            return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

        except Exception as error:
            db.session.rollback()
            flash(f"No fue posible registrar la cuenta por pagar: {error}", "error")
            return redirect("/crv/cuentas-por-pagar")

    proveedor_id = request.args.get("proveedor_id", "").strip()
    estado = request.args.get("estado", "").strip().upper()
    mes = request.args.get("mes", "").strip()
    buscar = request.args.get("buscar", "").strip()

    query = CRVCuentaPagar.query.filter_by(empresa_id=empresa_id)

    if proveedor_id.isdigit():
        query = query.filter_by(proveedor_id=int(proveedor_id))

    if estado:
        query = query.filter(CRVCuentaPagar.estado == estado)

    if mes:
        try:
            anio_texto, mes_texto = mes.split("-")
            query = query.filter(
                db.extract("year", CRVCuentaPagar.fecha_documento) == int(anio_texto),
                db.extract("month", CRVCuentaPagar.fecha_documento) == int(mes_texto)
            )
        except (ValueError, TypeError):
            pass

    if buscar:
        patron = f"%{buscar}%"
        query = query.join(
            CRVTercero,
            CRVCuentaPagar.proveedor_id == CRVTercero.id
        ).filter(
            db.or_(
                CRVCuentaPagar.numero_documento.ilike(patron),
                CRVCuentaPagar.concepto.ilike(patron),
                CRVTercero.razon_social.ilike(patron),
                CRVTercero.identificacion.ilike(patron)
            )
        )

    registros = query.order_by(
        CRVCuentaPagar.fecha_documento.desc(),
        CRVCuentaPagar.id.desc()
    ).all()

    hubo_cambios = False

    for cuenta in registros:
        estado_anterior = cuenta.estado
        cuenta.recalcular_estado()
        if estado_anterior != cuenta.estado:
            hubo_cambios = True

    if hubo_cambios:
        db.session.commit()

    total_facturado = sum(float(c.valor_neto or 0) for c in registros)
    total_pagado = sum(float(c.total_pagado or 0) for c in registros)
    total_pendiente = sum(float(c.saldo_pendiente or 0) for c in registros)
    total_vencido = sum(
        float(c.saldo_pendiente or 0)
        for c in registros
        if c.estado == "VENCIDA"
    )

    return render_template(
        "contabilidad/crv_cuentas_pagar.html",
        empresa_nombre=empresa_nombre,
        cuentas=registros,
        proveedores=proveedores,
        sedes=sedes,
        proveedor_actual=proveedor_id,
        estado_actual=estado,
        mes_actual=mes,
        buscar_actual=buscar,
        total_facturado=total_facturado,
        total_pagado=total_pagado,
        total_pendiente=total_pendiente,
        total_vencido=total_vencido
    )



# ============================================================
# CUENTAS POR PAGAR - MOVIMIENTO AUTOMÁTICO DE COMPRA
# ============================================================

def obtener_movimiento_compra_cxp(
    empresa_id,
    cuenta_id
):

    return CRVMovimiento.query.filter_by(
        empresa_id=empresa_id,
        origen_registro="CUENTAS_POR_PAGAR",
        numero_soporte=f"CXP-{cuenta_id}"
    ).order_by(
        CRVMovimiento.id.asc()
    ).first()


def sincronizar_movimiento_compra_cxp(
    cuenta
):

    proveedor = cuenta.proveedor

    movimiento = obtener_movimiento_compra_cxp(
        cuenta.empresa_id,
        cuenta.id
    )

    if movimiento is None:

        movimiento = CRVMovimiento(
            empresa_id=cuenta.empresa_id,
            origen_registro="CUENTAS_POR_PAGAR",
            numero_soporte=f"CXP-{cuenta.id}",
            revisado=False,
            estado="ACTIVO"
        )

        db.session.add(movimiento)

    movimiento.sede_id = cuenta.sede_id
    movimiento.fecha = cuenta.fecha_documento
    movimiento.tipo = "COMPRA"
    movimiento.categoria = cuenta.categoria or "COMPRA"
    movimiento.subcategoria = cuenta.subcategoria or ""
    movimiento.concepto = cuenta.concepto
    movimiento.descripcion = (
        cuenta.descripcion
        or f"Cuenta por pagar {cuenta.numero_documento}"
    )

    movimiento.tercero_id = (
        proveedor.id
        if proveedor
        else None
    )

    movimiento.tercero_nombre = (
        proveedor.razon_social
        if proveedor
        else ""
    )

    movimiento.tercero_identificacion = (
        proveedor.identificacion
        if proveedor
        else ""
    )

    movimiento.tercero_direccion = (
        proveedor.direccion
        if proveedor
        else ""
    )

    movimiento.tercero_email = (
        proveedor.email
        if proveedor
        else ""
    )

    movimiento.tercero_telefono = (
        (
            proveedor.celular
            or proveedor.telefono
            or ""
        )
        if proveedor
        else ""
    )

    movimiento.tercero_contacto = (
        proveedor.contacto
        if proveedor
        else ""
    )

    movimiento.tipo_documento = cuenta.tipo_documento or ""
    movimiento.numero_documento = cuenta.numero_documento
    movimiento.numero_soporte = f"CXP-{cuenta.id}"
    movimiento.fecha_documento = cuenta.fecha_documento
    movimiento.fecha_vencimiento = cuenta.fecha_vencimiento

    movimiento.base_gravable = cuenta.base_gravable or 0
    movimiento.tipo_iva = cuenta.tipo_iva or "NO_APLICA"
    movimiento.tarifa_iva = cuenta.tarifa_iva or 0
    movimiento.valor_iva = cuenta.valor_iva or 0
    movimiento.valor = cuenta.valor_bruto or 0
    movimiento.descuento = cuenta.descuento or 0
    movimiento.retencion = cuenta.retencion or 0
    movimiento.otros_descuentos = cuenta.otros_descuentos or 0
    movimiento.valor_neto = cuenta.valor_neto or 0

    movimiento.forma_pago = "CREDITO"
    movimiento.cuenta = "CUENTAS POR PAGAR"
    movimiento.cuenta_destino = ""
    movimiento.beneficiario = (
        proveedor.razon_social
        if proveedor
        else ""
    )
    movimiento.referencia_pago = ""

    movimiento.proyecto = cuenta.proyecto or ""
    movimiento.centro_costo = cuenta.centro_costo or ""
    movimiento.obra = cuenta.obra or ""

    movimiento.observaciones = (
        "Registro automático sincronizado desde "
        f"la cuenta por pagar CXP-{cuenta.id}."
    )

    movimiento.origen_registro = "CUENTAS_POR_PAGAR"
    movimiento.estado = "ACTIVO"
    movimiento.usuario_registro = str(
        session.get(
            "usuario",
            ""
        )
    )

    _asignar_comprobante_permanente(
        movimiento
    )

    return movimiento


# ============================================================
# CUENTA POR PAGAR - EDITAR OBLIGACIÓN
# ============================================================

@crv_bp.route(
    "/cuentas-por-pagar/<int:cuenta_id>/editar",
    methods=["POST"]
)
def editar_cuenta_por_pagar(cuenta_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    cuenta = CRVCuentaPagar.query.get_or_404(
        cuenta_id
    )

    if cuenta.empresa_id != empresa_id:

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    proveedor_id_texto = request.form.get(
        "proveedor_id",
        ""
    ).strip()

    if not proveedor_id_texto.isdigit():

        flash(
            "Debes seleccionar un proveedor válido.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    proveedor = CRVTercero.query.get_or_404(
        int(proveedor_id_texto)
    )

    if proveedor.empresa_id != empresa_id:

        flash(
            "El proveedor seleccionado no pertenece a CRV.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    numero_documento = request.form.get(
        "numero_documento",
        ""
    ).strip()

    fecha_documento = convertir_fecha(
        request.form.get(
            "fecha_documento"
        )
    )

    fecha_vencimiento = convertir_fecha(
        request.form.get(
            "fecha_vencimiento"
        )
    )

    concepto = request.form.get(
        "concepto",
        ""
    ).strip()

    iva_datos = calcular_iva_formulario(
        request.form
    )

    valor_bruto = iva_datos[
        "total_documento"
    ]

    descuento = convertir_numero(
        request.form.get(
            "descuento"
        )
    )

    retencion = convertir_numero(
        request.form.get(
            "retencion"
        )
    )

    otros_descuentos = convertir_numero(
        request.form.get(
            "otros_descuentos"
        )
    )

    if not numero_documento:

        flash(
            "El número de documento o factura es obligatorio.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    if not fecha_documento:

        flash(
            "La fecha del documento es obligatoria.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    if not concepto:

        flash(
            "El concepto de la obligación es obligatorio.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    if valor_bruto <= 0:

        flash(
            "El valor bruto debe ser mayor a cero.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    duplicada = CRVCuentaPagar.query.filter(
        CRVCuentaPagar.empresa_id == empresa_id,
        CRVCuentaPagar.proveedor_id == proveedor.id,
        CRVCuentaPagar.numero_documento == numero_documento,
        CRVCuentaPagar.id != cuenta.id
    ).first()

    if duplicada:

        flash(
            "Ya existe otra obligación de este proveedor con el mismo número de documento.",
            "error"
        )

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    nuevo_valor_neto = max(
        0,
        valor_bruto
        - descuento
        - retencion
        - otros_descuentos
    )

    total_pagado = sum(
        float(p.valor_pagado or 0)
        for p in cuenta.pagos
    )

    if nuevo_valor_neto < total_pagado:

        flash(
            "No puedes dejar el valor neto de la obligación por debajo de lo que ya se ha pagado. Primero reversa o corrige los pagos.",
            "error"
        )

        return redirect(
            f"/crv/cuentas-por-pagar/{cuenta.id}"
        )

    sede_id_texto = request.form.get(
        "sede_id",
        ""
    ).strip()

    sede_id = (
        int(sede_id_texto)
        if sede_id_texto.isdigit()
        else None
    )

    try:

        cuenta.proveedor_id = proveedor.id
        cuenta.sede_id = sede_id

        cuenta.tipo_documento = request.form.get(
            "tipo_documento",
            ""
        ).strip()

        cuenta.numero_documento = numero_documento
        cuenta.fecha_documento = fecha_documento
        cuenta.fecha_vencimiento = fecha_vencimiento

        cuenta.concepto = concepto

        cuenta.descripcion = request.form.get(
            "descripcion",
            ""
        ).strip()

        cuenta.base_gravable = iva_datos["base_gravable"]
        cuenta.tipo_iva = iva_datos["tipo_iva"]
        cuenta.tarifa_iva = iva_datos["tarifa_iva"]
        cuenta.valor_iva = iva_datos["valor_iva"]
        cuenta.valor_bruto = valor_bruto
        cuenta.descuento = descuento
        cuenta.retencion = retencion
        cuenta.otros_descuentos = otros_descuentos
        cuenta.valor_neto = nuevo_valor_neto

        cuenta.categoria = request.form.get(
            "categoria",
            "COMPRA"
        ).strip() or "COMPRA"

        cuenta.subcategoria = request.form.get(
            "subcategoria",
            ""
        ).strip()

        cuenta.proyecto = request.form.get(
            "proyecto",
            ""
        ).strip()

        cuenta.centro_costo = request.form.get(
            "centro_costo",
            ""
        ).strip()

        cuenta.obra = request.form.get(
            "obra",
            ""
        ).strip()

        cuenta.observaciones = request.form.get(
            "observaciones",
            ""
        ).strip()

        cuenta.usuario_registro = str(
            session.get(
                "usuario",
                ""
            )
        )

        cuenta.recalcular_saldos()

        sincronizar_movimiento_compra_cxp(
            cuenta
        )

        db.session.commit()

        flash(
            "Cuenta por pagar actualizada correctamente. El movimiento de compra también fue sincronizado.",
            "success"
        )

    except Exception as error:

        db.session.rollback()

        flash(
            f"No fue posible actualizar la cuenta por pagar: {error}",
            "error"
        )

    return redirect(
        f"/crv/cuentas-por-pagar/{cuenta.id}"
    )


# ============================================================
# CUENTA POR PAGAR - ELIMINAR OBLIGACIÓN
# ============================================================

@crv_bp.route(
    "/cuentas-por-pagar/<int:cuenta_id>/eliminar",
    methods=["POST"]
)
def eliminar_cuenta_por_pagar(cuenta_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    cuenta = CRVCuentaPagar.query.get_or_404(
        cuenta_id
    )

    if cuenta.empresa_id != empresa_id:

        return redirect(
            "/crv/cuentas-por-pagar"
        )

    cantidad_pagos = CRVPagoCuentaPagar.query.filter_by(
        empresa_id=empresa_id,
        cuenta_pagar_id=cuenta.id
    ).count()

    if cantidad_pagos > 0:

        flash(
            "No puedes eliminar esta obligación porque tiene pagos registrados. Primero entra al detalle y reversa todos los pagos.",
            "error"
        )

        return redirect(
            f"/crv/cuentas-por-pagar/{cuenta.id}"
        )

    try:

        movimiento_compra = obtener_movimiento_compra_cxp(
            empresa_id,
            cuenta.id
        )

        if movimiento_compra:

            db.session.delete(
                movimiento_compra
            )

        db.session.delete(
            cuenta
        )

        db.session.commit()

        flash(
            "Cuenta por pagar eliminada correctamente. El movimiento de compra asociado también fue eliminado.",
            "success"
        )

    except Exception as error:

        db.session.rollback()

        flash(
            f"No fue posible eliminar la cuenta por pagar: {error}",
            "error"
        )

        return redirect(
            f"/crv/cuentas-por-pagar/{cuenta.id}"
        )

    return redirect(
        "/crv/cuentas-por-pagar"
    )


# ============================================================
# CUENTA POR PAGAR - DETALLE
# ============================================================

@crv_bp.route("/cuentas-por-pagar/<int:cuenta_id>")
def cuenta_por_pagar_detalle(cuenta_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get("empresa_nombre", "CRV CONSTRUCTORA")

    cuenta = CRVCuentaPagar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-pagar")

    cuenta.recalcular_saldos()
    db.session.commit()

    pagos = CRVPagoCuentaPagar.query.filter_by(
        empresa_id=empresa_id,
        cuenta_pagar_id=cuenta.id
    ).order_by(
        CRVPagoCuentaPagar.fecha.desc(),
        CRVPagoCuentaPagar.id.desc()
    ).all()

    return render_template(
        "contabilidad/crv_cuenta_pagar_detalle.html",
        empresa_nombre=empresa_nombre,
        cuenta=cuenta,
        pagos=pagos
    )


# ============================================================
# CUENTA POR PAGAR - REGISTRAR PAGO
# ============================================================

@crv_bp.route(
    "/cuentas-por-pagar/<int:cuenta_id>/pagar",
    methods=["POST"]
)
def pagar_cuenta_por_pagar(cuenta_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaPagar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-pagar")

    cuenta.recalcular_saldos()

    if cuenta.saldo_pendiente <= 0:
        flash("Esta obligación ya se encuentra pagada.", "warning")
        return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

    fecha_pago = convertir_fecha(request.form.get("fecha"))
    valor_pagado = convertir_numero(request.form.get("valor_pagado"))

    if not fecha_pago:
        flash("La fecha real del pago es obligatoria.", "error")
        return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

    if valor_pagado <= 0:
        flash("El valor del pago debe ser mayor a cero.", "error")
        return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

    if valor_pagado > float(cuenta.saldo_pendiente or 0):
        flash(
            "El pago no puede superar el saldo pendiente de la obligación.",
            "error"
        )
        return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

    proveedor = cuenta.proveedor

    pago = CRVPagoCuentaPagar(
        empresa_id=empresa_id,
        cuenta_pagar_id=cuenta.id,
        fecha=fecha_pago,
        valor_pagado=valor_pagado,
        forma_pago=request.form.get("forma_pago", "").strip(),
        cuenta_origen=request.form.get("cuenta_origen", "").strip(),
        cuenta_destino=request.form.get("cuenta_destino", "").strip(),
        beneficiario=request.form.get(
            "beneficiario",
            proveedor.razon_social if proveedor else ""
        ).strip(),
        referencia_pago=request.form.get("referencia_pago", "").strip(),
        numero_soporte=request.form.get("numero_soporte", "").strip(),
        observaciones=request.form.get("observaciones", "").strip(),
        usuario_registro=str(session.get("usuario", ""))
    )

    try:
        db.session.add(pago)
        db.session.flush()

        movimiento_pago = CRVMovimiento(
            empresa_id=empresa_id,
            sede_id=cuenta.sede_id,
            fecha=fecha_pago,
            tipo="PAGO CUENTA POR PAGAR",
            categoria="PAGO PROVEEDOR",
            subcategoria=cuenta.categoria or "",
            concepto=f"Pago cuenta por pagar {cuenta.numero_documento}",
            descripcion=(
                f"Abono/pago a "
                f"{proveedor.razon_social if proveedor else 'proveedor'} "
                f"- documento {cuenta.numero_documento}"
            ),
            tercero_id=proveedor.id if proveedor else None,
            tercero_nombre=proveedor.razon_social if proveedor else "",
            tercero_identificacion=proveedor.identificacion if proveedor else "",
            tercero_direccion=proveedor.direccion if proveedor else "",
            tercero_email=proveedor.email if proveedor else "",
            tercero_telefono=(
                (proveedor.celular or proveedor.telefono or "")
                if proveedor else ""
            ),
            tercero_contacto=proveedor.contacto if proveedor else "",
            tipo_documento="PAGO CXP",
            numero_documento=cuenta.numero_documento,
            numero_soporte=pago.numero_soporte or f"CXP-PAGO-{pago.id}",
            fecha_documento=fecha_pago,
            fecha_vencimiento=None,
            valor=valor_pagado,
            descuento=0,
            retencion=0,
            otros_descuentos=0,
            valor_neto=valor_pagado,
            forma_pago=pago.forma_pago or "",
            cuenta=pago.cuenta_origen or "",
            cuenta_destino=pago.cuenta_destino or "",
            beneficiario=pago.beneficiario or "",
            referencia_pago=pago.referencia_pago or "",
            proyecto=cuenta.proyecto or "",
            centro_costo=cuenta.centro_costo or "",
            obra=cuenta.obra or "",
            observaciones=(
                pago.observaciones
                or f"Movimiento automático generado desde la cuenta por pagar CXP-{cuenta.id}."
            ),
            revisado=False,
            estado="ACTIVO",
            origen_registro="PAGO_CUENTA_POR_PAGAR",
            usuario_registro=str(session.get("usuario", ""))
        )

        _asignar_comprobante_permanente(
            movimiento_pago
        )

        db.session.add(movimiento_pago)
        db.session.flush()

        pago.movimiento_id = movimiento_pago.id

        cuenta.recalcular_saldos()

        db.session.commit()

        flash(
            "Pago registrado correctamente. El saldo fue actualizado y el movimiento quedó registrado.",
            "success"
        )

    except Exception as error:
        db.session.rollback()
        flash(f"No fue posible registrar el pago: {error}", "error")

    return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")


# ============================================================
# CUENTA POR PAGAR - REVERSAR PAGO
# ============================================================

@crv_bp.route(
    "/cuentas-por-pagar/<int:cuenta_id>/pagos/<int:pago_id>/eliminar",
    methods=["POST"]
)
def eliminar_pago_cuenta_por_pagar(cuenta_id, pago_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaPagar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-pagar")

    pago = CRVPagoCuentaPagar.query.get_or_404(pago_id)

    if pago.empresa_id != empresa_id or pago.cuenta_pagar_id != cuenta.id:
        return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

    try:
        if pago.movimiento_id:
            movimiento = CRVMovimiento.query.filter_by(
                id=pago.movimiento_id,
                empresa_id=empresa_id
            ).first()

            if movimiento:
                db.session.delete(movimiento)

        db.session.delete(pago)
        db.session.flush()

        cuenta.recalcular_saldos()
        db.session.commit()

        flash(
            "Pago reversado correctamente. El movimiento asociado también fue eliminado.",
            "success"
        )

    except Exception as error:
        db.session.rollback()
        flash(f"No fue posible reversar el pago: {error}", "error")

    return redirect(f"/crv/cuentas-por-pagar/{cuenta.id}")

# ============================================================
# REPORTES CONTADOR - FUNCIONES AUXILIARES
# ============================================================

def _reporte_filtros():
    return {
        "fecha_desde": request.args.get("fecha_desde", "").strip(),
        "fecha_hasta": request.args.get("fecha_hasta", "").strip(),
        "mes": request.args.get("mes", "").strip(),
        "anio": request.args.get("anio", "").strip(),
        "tipo": request.args.get("tipo", "").strip().upper(),
        "tercero": request.args.get("tercero", "").strip(),
        "proyecto": request.args.get("proyecto", "").strip(),
        "centro_costo": request.args.get("centro_costo", "").strip(),
    }


def _aplicar_filtros_movimientos(query, filtros):
    fecha_desde = convertir_fecha(filtros.get("fecha_desde"))
    fecha_hasta = convertir_fecha(filtros.get("fecha_hasta"))

    if fecha_desde:
        query = query.filter(CRVMovimiento.fecha >= fecha_desde)

    if fecha_hasta:
        query = query.filter(CRVMovimiento.fecha <= fecha_hasta)

    mes = filtros.get("mes", "")
    if mes:
        try:
            anio_texto, mes_texto = mes.split("-")
            query = query.filter(
                db.extract("year", CRVMovimiento.fecha) == int(anio_texto),
                db.extract("month", CRVMovimiento.fecha) == int(mes_texto)
            )
        except (ValueError, TypeError):
            pass

    anio = filtros.get("anio", "")
    if anio.isdigit():
        query = query.filter(
            db.extract("year", CRVMovimiento.fecha) == int(anio)
        )

    tipo = filtros.get("tipo", "")
    if tipo:
        query = query.filter(CRVMovimiento.tipo == tipo)

    tercero = filtros.get("tercero", "")
    if tercero:
        patron = f"%{tercero}%"
        query = query.filter(
            db.or_(
                CRVMovimiento.tercero_nombre.ilike(patron),
                CRVMovimiento.tercero_identificacion.ilike(patron)
            )
        )

    proyecto = filtros.get("proyecto", "")
    if proyecto:
        query = query.filter(
            CRVMovimiento.proyecto.ilike(f"%{proyecto}%")
        )

    centro_costo = filtros.get("centro_costo", "")
    if centro_costo:
        query = query.filter(
            CRVMovimiento.centro_costo.ilike(f"%{centro_costo}%")
        )

    return query


def _aplicar_filtros_cxp(query, filtros):
    fecha_desde = convertir_fecha(filtros.get("fecha_desde"))
    fecha_hasta = convertir_fecha(filtros.get("fecha_hasta"))

    if fecha_desde:
        query = query.filter(CRVCuentaPagar.fecha_documento >= fecha_desde)

    if fecha_hasta:
        query = query.filter(CRVCuentaPagar.fecha_documento <= fecha_hasta)

    mes = filtros.get("mes", "")
    if mes:
        try:
            anio_texto, mes_texto = mes.split("-")
            query = query.filter(
                db.extract("year", CRVCuentaPagar.fecha_documento) == int(anio_texto),
                db.extract("month", CRVCuentaPagar.fecha_documento) == int(mes_texto)
            )
        except (ValueError, TypeError):
            pass

    anio = filtros.get("anio", "")
    if anio.isdigit():
        query = query.filter(
            db.extract("year", CRVCuentaPagar.fecha_documento) == int(anio)
        )

    tercero = filtros.get("tercero", "")
    if tercero:
        patron = f"%{tercero}%"
        query = query.join(
            CRVTercero,
            CRVCuentaPagar.proveedor_id == CRVTercero.id
        ).filter(
            db.or_(
                CRVTercero.razon_social.ilike(patron),
                CRVTercero.identificacion.ilike(patron)
            )
        )

    proyecto = filtros.get("proyecto", "")
    if proyecto:
        query = query.filter(
            CRVCuentaPagar.proyecto.ilike(f"%{proyecto}%")
        )

    centro_costo = filtros.get("centro_costo", "")
    if centro_costo:
        query = query.filter(
            CRVCuentaPagar.centro_costo.ilike(f"%{centro_costo}%")
        )

    return query


def _calcular_resumen_reporte(movimientos, cuentas):
    ingresos = 0.0
    compras = 0.0
    gasto_general = 0.0
    nomina = 0.0
    anticipos = 0.0
    prestamos = 0.0
    otros = 0.0

    for movimiento in movimientos:
        tipo = str(movimiento.tipo or "").strip().upper()
        valor = float(movimiento.valor_neto or movimiento.valor or 0)

        if tipo == "INGRESO":
            ingresos += valor
        elif tipo == "COMPRA":
            compras += valor
        elif tipo == "GASTO":
            gasto_general += valor
        elif tipo in ("NOMINA AFILIADOS", "NÓMINA AFILIADOS", "NOMINA OBRA", "NÓMINA OBRA"):
            nomina += valor
        elif tipo == "ANTICIPO":
            anticipos += valor
        elif tipo in ("PRESTAMO EMPLEADO", "PRÉSTAMO EMPLEADO"):
            prestamos += valor
        elif tipo in (
            "TRANSFERENCIA",
            "PAGO CUENTA POR PAGAR",
            "SALDO INICIAL CUENTA",
            "COBRO CUENTA POR COBRAR",
        ):
            pass
        else:
            otros += valor

    gastos = gasto_general + nomina + anticipos + prestamos + otros
    egresos = compras + gastos
    resultado = ingresos - egresos

    cxp_facturado = sum(float(c.valor_neto or 0) for c in cuentas)
    cxp_pagado = sum(float(c.total_pagado or 0) for c in cuentas)
    cxp_pendiente = sum(float(c.saldo_pendiente or 0) for c in cuentas)
    cxp_vencido = sum(
        float(c.saldo_pendiente or 0)
        for c in cuentas
        if c.estado == "VENCIDA"
    )

    return {
        "ingresos": ingresos,
        "compras": compras,
        "gastos": gastos,
        "nomina": nomina,
        "anticipos": anticipos,
        "prestamos": prestamos,
        "otros": otros,
        "egresos": egresos,
        "resultado": resultado,
        "cxp_facturado": cxp_facturado,
        "cxp_pagado": cxp_pagado,
        "cxp_pendiente": cxp_pendiente,
        "cxp_vencido": cxp_vencido,
        "cantidad_movimientos": len(movimientos),
        "cantidad_cxp": len(cuentas),
    }


def _datos_reportes_crv(empresa_id, filtros):
    mov_query = CRVMovimiento.query.filter_by(empresa_id=empresa_id)
    mov_query = _aplicar_filtros_movimientos(mov_query, filtros)
    movimientos = mov_query.order_by(
        CRVMovimiento.fecha.asc(),
        CRVMovimiento.id.asc()
    ).all()

    cxp_query = CRVCuentaPagar.query.filter_by(empresa_id=empresa_id)
    cxp_query = _aplicar_filtros_cxp(cxp_query, filtros)
    cuentas = cxp_query.order_by(
        CRVCuentaPagar.fecha_documento.asc(),
        CRVCuentaPagar.id.asc()
    ).all()

    terceros = CRVTercero.query.filter_by(
        empresa_id=empresa_id
    ).order_by(CRVTercero.razon_social.asc()).all()

    resumen = _calcular_resumen_reporte(movimientos, cuentas)
    iva_resumen = resumen_iva_movimientos(movimientos)

    # Resumen por proveedor basado en CxP.
    proveedor_map = {}
    for cuenta in cuentas:
        proveedor = cuenta.proveedor
        clave = proveedor.id if proveedor else f"SIN-{cuenta.id}"
        if clave not in proveedor_map:
            proveedor_map[clave] = {
                "nombre": proveedor.razon_social if proveedor else "SIN PROVEEDOR",
                "identificacion": proveedor.identificacion if proveedor else "",
                "facturado": 0.0,
                "pagado": 0.0,
                "saldo": 0.0,
                "vencido": 0.0,
            }
        fila = proveedor_map[clave]
        fila["facturado"] += float(cuenta.valor_neto or 0)
        fila["pagado"] += float(cuenta.total_pagado or 0)
        fila["saldo"] += float(cuenta.saldo_pendiente or 0)
        if cuenta.estado == "VENCIDA":
            fila["vencido"] += float(cuenta.saldo_pendiente or 0)

    proveedores_resumen = sorted(
        proveedor_map.values(),
        key=lambda x: x["nombre"].upper()
    )

    # Resumen por proyecto / centro de costo con impacto contable.
    proyectos_map = {}
    tipos_sin_impacto = {
        "TRANSFERENCIA",
        "PAGO CUENTA POR PAGAR",
        "SALDO INICIAL CUENTA",
        "COBRO CUENTA POR COBRAR",
    }
    for movimiento in movimientos:
        tipo = str(movimiento.tipo or "").strip().upper()
        if tipo in tipos_sin_impacto or tipo == "INGRESO":
            continue
        proyecto = (movimiento.proyecto or "SIN PROYECTO").strip() or "SIN PROYECTO"
        centro = (movimiento.centro_costo or "SIN CENTRO DE COSTO").strip() or "SIN CENTRO DE COSTO"
        clave = (proyecto, centro)
        if clave not in proyectos_map:
            proyectos_map[clave] = {
                "proyecto": proyecto,
                "centro_costo": centro,
                "total": 0.0,
                "movimientos": 0,
            }
        proyectos_map[clave]["total"] += float(
            movimiento.valor_neto or movimiento.valor or 0
        )
        proyectos_map[clave]["movimientos"] += 1

    proyectos_resumen = sorted(
        proyectos_map.values(),
        key=lambda x: (x["proyecto"].upper(), x["centro_costo"].upper())
    )

    return {
        "movimientos": movimientos,
        "cuentas": cuentas,
        "terceros": terceros,
        "resumen": resumen,
        "iva_resumen": iva_resumen,
        "proveedores_resumen": proveedores_resumen,
        "proyectos_resumen": proyectos_resumen,
    }


def _periodo_texto(filtros):
    if filtros.get("fecha_desde") or filtros.get("fecha_hasta"):
        return f"{filtros.get('fecha_desde') or 'Inicio'} a {filtros.get('fecha_hasta') or 'Hoy'}"
    if filtros.get("mes"):
        return filtros["mes"]
    if filtros.get("anio"):
        return filtros["anio"]
    return "Todos los registros"


# ============================================================
# REPORTES CONTADOR - PANTALLA PRINCIPAL
# ============================================================

@crv_bp.route("/reportes-contador")
def reportes_contador():
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get("empresa_nombre", "CRV CONSTRUCTORA")
    filtros = _reporte_filtros()
    datos = _datos_reportes_crv(empresa_id, filtros)

    tipos = [
        fila[0]
        for fila in db.session.query(CRVMovimiento.tipo)
        .filter(CRVMovimiento.empresa_id == empresa_id)
        .distinct()
        .order_by(CRVMovimiento.tipo.asc())
        .all()
        if fila[0]
    ]

    proyectos = [
        fila[0]
        for fila in db.session.query(CRVMovimiento.proyecto)
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            CRVMovimiento.proyecto.isnot(None),
            CRVMovimiento.proyecto != ""
        )
        .distinct()
        .order_by(CRVMovimiento.proyecto.asc())
        .all()
        if fila[0]
    ]

    centros = [
        fila[0]
        for fila in db.session.query(CRVMovimiento.centro_costo)
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            CRVMovimiento.centro_costo.isnot(None),
            CRVMovimiento.centro_costo != ""
        )
        .distinct()
        .order_by(CRVMovimiento.centro_costo.asc())
        .all()
        if fila[0]
    ]

    return render_template(
        "contabilidad/crv_reportes.html",
        empresa_nombre=empresa_nombre,
        filtros=filtros,
        periodo=_periodo_texto(filtros),
        tipos=tipos,
        proyectos=proyectos,
        centros=centros,
        movimientos=datos["movimientos"],
        cuentas=datos["cuentas"],
        terceros=datos["terceros"],
        resumen=datos["resumen"],
        iva_resumen=datos["iva_resumen"],
        proveedores_resumen=datos["proveedores_resumen"],
        proyectos_resumen=datos["proyectos_resumen"]
    )


# ============================================================
# REPORTES CONTADOR - EXCEL PROFESIONAL
# ============================================================

@crv_bp.route("/reportes-contador/excel")
def reportes_contador_excel():
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    filtros = _reporte_filtros()
    datos = _datos_reportes_crv(empresa_id, filtros)
    resumen = datos["resumen"]

    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"

    navy = "17304F"
    gold = "D8B541"
    green = "0F8A47"
    red = "D92D20"
    light = "EEF3F8"
    white = "FFFFFF"
    gray = "667085"
    thin_gray = Side(style="thin", color="D9E1E8")

    def titulo_hoja(sheet, titulo, subtitulo=""):
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:H1")
        c = sheet["A1"]
        c.value = titulo
        c.font = Font(size=18, bold=True, color=navy)
        c.alignment = Alignment(horizontal="left")
        if subtitulo:
            sheet.merge_cells("A2:H2")
            c2 = sheet["A2"]
            c2.value = subtitulo
            c2.font = Font(size=10, color=gray)
        sheet.row_dimensions[1].height = 28

    def encabezado(sheet, fila, headers):
        for col, valor in enumerate(headers, 1):
            cell = sheet.cell(fila, col, valor)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[fila].height = 28

    def moneda(cell):
        cell.number_format = '$#,##0;[Red]($#,##0);-'
        cell.alignment = Alignment(horizontal="right")

    def auto_ancho(sheet, max_width=38):
        for col in range(1, sheet.max_column + 1):
            max_len = 0
            letter = get_column_letter(col)
            for row in range(1, min(sheet.max_row, 500) + 1):
                value = sheet.cell(row, col).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)

    # RESUMEN
    titulo_hoja(ws, "CRV CONSTRUCTORA - REPORTE CONTABLE", f"Periodo: {_periodo_texto(filtros)}")
    ws["A4"] = "Indicador"
    ws["B4"] = "Valor"
    encabezado(ws, 4, ["Indicador", "Valor"])

    resumen_filas = [
        ("Ingresos", resumen["ingresos"]),
        ("Compras", resumen["compras"]),
        ("Gastos", resumen["gastos"]),
        ("Egresos totales", resumen["egresos"]),
        ("Resultado", resumen["resultado"]),
        ("CxP facturado", resumen["cxp_facturado"]),
        ("CxP pagado", resumen["cxp_pagado"]),
        ("CxP pendiente", resumen["cxp_pendiente"]),
        ("CxP vencido", resumen["cxp_vencido"]),
        ("Base gravable IVA", datos["iva_resumen"]["base_gravable"]),
        ("Base IVA generado", datos["iva_resumen"]["base_generado"]),
        ("Base IVA descontable", datos["iva_resumen"]["base_descontable"]),
        ("Base IVA mayor valor costo/gasto", datos["iva_resumen"]["base_mayor_valor"]),
        ("Base exenta", datos["iva_resumen"]["base_exenta"]),
        ("Base excluida", datos["iva_resumen"]["base_excluida"]),
        ("IVA generado", datos["iva_resumen"]["iva_generado"]),
        ("IVA descontable", datos["iva_resumen"]["iva_descontable"]),
        ("IVA mayor valor costo/gasto", datos["iva_resumen"]["iva_mayor_valor"]),
        ("Saldo IVA preliminar", datos["iva_resumen"]["saldo_iva_preliminar"]),
    ]
    for i, (label, valor) in enumerate(resumen_filas, 5):
        ws.cell(i, 1, label)
        ws.cell(i, 2, valor)
        moneda(ws.cell(i, 2))
        if label == "Resultado":
            ws.cell(i, 2).font = Font(bold=True, color=green if valor >= 0 else red)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20

    # MOVIMIENTOS
    mov_ws = wb.create_sheet("MOVIMIENTOS")
    titulo_hoja(mov_ws, "MOVIMIENTOS CONTABLES CRV", f"Periodo: {_periodo_texto(filtros)}")
    mov_headers = [
        "Fecha", "Tipo", "Categoría", "Subcategoría", "Tercero", "NIT/ID",
        "Documento", "Concepto", "Descripción", "Cuenta", "Cuenta destino",
        "Proyecto", "Centro costo", "Obra", "Base gravable", "Tipo IVA",
        "Tarifa IVA %", "Valor IVA", "Total documento", "Descuento",
        "Retención", "Otros descuentos", "Valor neto", "Forma pago",
        "Soporte", "Referencia", "Origen"
    ]
    encabezado(mov_ws, 4, mov_headers)
    for r, m in enumerate(datos["movimientos"], 5):
        vals = [
            m.fecha, m.tipo, m.categoria, m.subcategoria, m.nombre_tercero,
            m.tercero_identificacion, m.numero_documento, m.concepto, m.descripcion,
            m.cuenta, m.cuenta_destino, m.proyecto, m.centro_costo, m.obra,
            float(m.base_gravable or 0), m.tipo_iva or "NO_APLICA",
            float(m.tarifa_iva or 0), float(m.valor_iva or 0), float(m.valor or 0),
            float(m.descuento or 0), float(m.retencion or 0),
            float(m.otros_descuentos or 0), float(m.valor_neto or m.valor or 0),
            m.forma_pago, m.numero_soporte, m.referencia_pago, m.origen_registro
        ]
        for c, v in enumerate(vals, 1):
            mov_ws.cell(r, c, v)
        for c in (15, 18, 19, 20, 21, 22, 23):
            moneda(mov_ws.cell(r, c))
    mov_ws.freeze_panes = "A5"
    mov_ws.auto_filter.ref = f"A4:{get_column_letter(len(mov_headers))}{max(4, mov_ws.max_row)}"
    auto_ancho(mov_ws)

    # CUENTAS POR PAGAR
    cxp_ws = wb.create_sheet("CUENTAS_POR_PAGAR")
    titulo_hoja(cxp_ws, "CUENTAS POR PAGAR CRV", f"Periodo documento: {_periodo_texto(filtros)}")
    cxp_headers = [
        "Fecha", "Proveedor", "NIT/ID", "Tipo documento", "Documento", "Concepto",
        "Proyecto", "Centro costo", "Vencimiento", "Base gravable", "Tipo IVA",
        "Tarifa IVA %", "Valor IVA", "Total documento", "Descuento",
        "Retención", "Otros descuentos", "Neto", "Pagado", "Saldo", "Estado"
    ]
    encabezado(cxp_ws, 4, cxp_headers)
    for r, cta in enumerate(datos["cuentas"], 5):
        prov = cta.proveedor
        vals = [
            cta.fecha_documento,
            prov.razon_social if prov else "",
            prov.identificacion if prov else "",
            cta.tipo_documento,
            cta.numero_documento,
            cta.concepto,
            cta.proyecto,
            cta.centro_costo,
            cta.fecha_vencimiento,
            float(cta.base_gravable or 0),
            cta.tipo_iva or "NO_APLICA",
            float(cta.tarifa_iva or 0),
            float(cta.valor_iva or 0),
            float(cta.valor_bruto or 0),
            float(cta.descuento or 0),
            float(cta.retencion or 0),
            float(cta.otros_descuentos or 0),
            float(cta.valor_neto or 0),
            float(cta.total_pagado or 0),
            float(cta.saldo_pendiente or 0),
            cta.estado,
        ]
        for col, v in enumerate(vals, 1):
            cxp_ws.cell(r, col, v)
        for col in (10, 13, 14, 15, 16, 17, 18, 19, 20):
            moneda(cxp_ws.cell(r, col))
    cxp_ws.freeze_panes = "A5"
    cxp_ws.auto_filter.ref = f"A4:{get_column_letter(len(cxp_headers))}{max(4, cxp_ws.max_row)}"
    auto_ancho(cxp_ws)

    # IVA / IMPUESTOS
    iva_ws = wb.create_sheet("IVA")
    titulo_hoja(iva_ws, "RESUMEN DE IVA CRV", f"Periodo: {_periodo_texto(filtros)}")
    encabezado(iva_ws, 4, ["Concepto", "Valor"])
    iva_filas = [
        ("Base gravable IVA", datos["iva_resumen"]["base_gravable"]),
        ("Base IVA generado", datos["iva_resumen"]["base_generado"]),
        ("Base IVA descontable", datos["iva_resumen"]["base_descontable"]),
        ("Base IVA mayor valor costo/gasto", datos["iva_resumen"]["base_mayor_valor"]),
        ("Base exenta", datos["iva_resumen"]["base_exenta"]),
        ("Base excluida", datos["iva_resumen"]["base_excluida"]),
        ("IVA generado", datos["iva_resumen"]["iva_generado"]),
        ("IVA descontable", datos["iva_resumen"]["iva_descontable"]),
        ("IVA mayor valor costo/gasto", datos["iva_resumen"]["iva_mayor_valor"]),
        ("IVA total registrado", datos["iva_resumen"]["iva_total"]),
        ("Saldo IVA preliminar", datos["iva_resumen"]["saldo_iva_preliminar"]),
    ]
    for r, (concepto, valor) in enumerate(iva_filas, 5):
        iva_ws.cell(r, 1, concepto)
        iva_ws.cell(r, 2, valor)
        moneda(iva_ws.cell(r, 2))

    encabezado(iva_ws, 18, [
        "Fecha", "Tipo", "Tercero", "Documento", "Concepto",
        "Tipo IVA", "Tarifa %", "Base gravable", "Valor IVA", "Total documento"
    ])
    rr = 19
    for m in datos["movimientos"]:
        tipo_iva_mov = str(
            m.tipo_iva or "NO_APLICA"
        ).strip().upper()

        # La hoja tributaria incluye movimientos con tratamiento IVA,
        # incluso EXENTOS y EXCLUIDOS aunque su valor IVA sea cero.
        if tipo_iva_mov == "NO_APLICA":
            continue
        vals = [
            m.fecha, m.tipo, m.nombre_tercero, m.numero_documento or m.numero_soporte,
            m.concepto, m.tipo_iva, float(m.tarifa_iva or 0),
            float(m.base_gravable or 0), float(m.valor_iva or 0), float(m.valor or 0)
        ]
        for c, v in enumerate(vals, 1):
            iva_ws.cell(rr, c, v)
        for c in (8, 9, 10):
            moneda(iva_ws.cell(rr, c))
        rr += 1
    auto_ancho(iva_ws)

    # PROVEEDORES
    pr_ws = wb.create_sheet("PROVEEDORES")
    titulo_hoja(pr_ws, "RESUMEN POR PROVEEDOR", f"Periodo: {_periodo_texto(filtros)}")
    encabezado(pr_ws, 4, ["Proveedor", "NIT/ID", "Facturado", "Pagado", "Saldo", "Vencido"])
    for r, p in enumerate(datos["proveedores_resumen"], 5):
        vals = [p["nombre"], p["identificacion"], p["facturado"], p["pagado"], p["saldo"], p["vencido"]]
        for c, v in enumerate(vals, 1):
            pr_ws.cell(r, c, v)
        for c in range(3, 7):
            moneda(pr_ws.cell(r, c))
    auto_ancho(pr_ws)

    # PROYECTOS
    proy_ws = wb.create_sheet("PROYECTOS")
    titulo_hoja(proy_ws, "COSTOS POR PROYECTO Y CENTRO DE COSTO", f"Periodo: {_periodo_texto(filtros)}")
    encabezado(proy_ws, 4, ["Proyecto", "Centro de costo", "Movimientos", "Costo / egreso"])
    for r, p in enumerate(datos["proyectos_resumen"], 5):
        proy_ws.cell(r, 1, p["proyecto"])
        proy_ws.cell(r, 2, p["centro_costo"])
        proy_ws.cell(r, 3, p["movimientos"])
        proy_ws.cell(r, 4, p["total"])
        moneda(proy_ws.cell(r, 4))
    auto_ancho(proy_ws)

    # TERCEROS
    ter_ws = wb.create_sheet("TERCEROS")
    titulo_hoja(ter_ws, "DIRECTORIO DE TERCEROS CRV", "Información registrada en el sistema")
    ter_headers = [
        "Tipo", "Tipo documento", "Identificación", "Razón social", "Nombre comercial",
        "Dirección", "Ciudad", "Departamento", "Teléfono", "Celular", "Email",
        "Contacto", "Banco", "Tipo cuenta", "Número cuenta", "Titular cuenta", "Activo"
    ]
    encabezado(ter_ws, 4, ter_headers)
    for r, t in enumerate(datos["terceros"], 5):
        vals = [
            t.tipo_tercero, t.tipo_documento, t.identificacion, t.razon_social,
            t.nombre_comercial, t.direccion, t.ciudad, t.departamento, t.telefono,
            t.celular, t.email, t.contacto, t.banco, t.tipo_cuenta, t.numero_cuenta,
            t.titular_cuenta, "SI" if t.activo else "NO"
        ]
        for c, v in enumerate(vals, 1):
            ter_ws.cell(r, c, v)
    ter_ws.freeze_panes = "A5"
    ter_ws.auto_filter.ref = f"A4:{get_column_letter(len(ter_headers))}{max(4, ter_ws.max_row)}"
    auto_ancho(ter_ws)

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre = f"CRV_reporte_contable_{date.today().isoformat()}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ============================================================
# REPORTES CONTADOR - PDF PROFESIONAL
# ============================================================

@crv_bp.route("/reportes-contador/pdf")
def reportes_contador_pdf():
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get("empresa_nombre", "CRV CONSTRUCTORA")
    filtros = _reporte_filtros()
    datos = _datos_reportes_crv(empresa_id, filtros)
    resumen = datos["resumen"]

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Reporte contable CRV"
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CRVTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=18, leading=21, textColor=colors.HexColor("#17304F"), alignment=TA_CENTER
    )
    subtitle_style = ParagraphStyle(
        "CRVSub", parent=styles["Normal"], fontName="Helvetica",
        fontSize=8.5, leading=11, textColor=colors.HexColor("#667085"), alignment=TA_CENTER
    )
    section_style = ParagraphStyle(
        "CRVSection", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=12, leading=14, textColor=colors.HexColor("#17304F"), spaceBefore=5, spaceAfter=7
    )
    small = ParagraphStyle(
        "CRVSmall", parent=styles["Normal"], fontName="Helvetica",
        fontSize=6.5, leading=8, textColor=colors.HexColor("#25364A"), alignment=TA_LEFT
    )

    story = []
    logo_path = Path(current_app.root_path) / "static" / "logo_crv.png"
    if logo_path.exists():
        story.append(Image(str(logo_path), width=24 * mm, height=18 * mm))
    story.append(Paragraph(f"{empresa_nombre} - REPORTE CONTABLE", title_style))
    story.append(Paragraph(f"Periodo: {_periodo_texto(filtros)} · Generado: {date.today().strftime('%d/%m/%Y')}", subtitle_style))
    story.append(Spacer(1, 5 * mm))

    def money(v):
        return f"${float(v or 0):,.0f}"

    # KPI summary
    kpi_data = [
        ["Ingresos", "Compras", "Gastos", "Resultado", "CxP pendiente", "CxP vencido"],
        [
            money(resumen["ingresos"]), money(resumen["compras"]), money(resumen["gastos"]),
            money(resumen["resultado"]), money(resumen["cxp_pendiente"]), money(resumen["cxp_vencido"])
        ]
    ]
    kpi = Table(kpi_data, colWidths=[43 * mm] * 6, rowHeights=[8 * mm, 11 * mm])
    kpi.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17304F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 11),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F7F9FC")),
        ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E1E8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4EAF0")),
    ]))
    story.append(kpi)
    story.append(Spacer(1, 5 * mm))

    story.append(Paragraph("Resumen de IVA", section_style))
    iva_tabla = Table(
        [
            [
                "Base IVA generado",
                "Base IVA descontable",
                "Base exenta",
                "Base excluida",
            ],
            [
                money(datos["iva_resumen"]["base_generado"]),
                money(datos["iva_resumen"]["base_descontable"]),
                money(datos["iva_resumen"]["base_exenta"]),
                money(datos["iva_resumen"]["base_excluida"]),
            ],
            [
                "IVA generado",
                "IVA descontable",
                "IVA mayor valor costo/gasto",
                "Saldo IVA preliminar",
            ],
            [
                money(datos["iva_resumen"]["iva_generado"]),
                money(datos["iva_resumen"]["iva_descontable"]),
                money(datos["iva_resumen"]["iva_mayor_valor"]),
                money(datos["iva_resumen"]["saldo_iva_preliminar"]),
            ],
        ],
        colWidths=[62 * mm] * 4
    )
    iva_tabla.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#D8B541")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.HexColor("#17304F")),
        ("FONTNAME", (0,0), (-1,-1), "Helvetica-Bold"),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#D9E1E8")),
        ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#FFF9E6")),
        ("TOPPADDING", (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(iva_tabla)
    story.append(Spacer(1, 5 * mm))

    # CxP summary by provider
    story.append(Paragraph("Resumen por proveedor", section_style))
    prov_rows = [["Proveedor", "NIT/ID", "Facturado", "Pagado", "Saldo", "Vencido"]]
    for p in datos["proveedores_resumen"]:
        prov_rows.append([
            Paragraph(p["nombre"], small), p["identificacion"], money(p["facturado"]),
            money(p["pagado"]), money(p["saldo"]), money(p["vencido"])
        ])
    if len(prov_rows) == 1:
        prov_rows.append(["Sin registros", "", "-", "-", "-", "-"])
    prov_table = Table(prov_rows, repeatRows=1, colWidths=[68*mm, 30*mm, 35*mm, 35*mm, 35*mm, 35*mm])
    prov_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17304F")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 6.5),
        ("ALIGN", (2,1), (-1,-1), "RIGHT"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E1E8")),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story.append(prov_table)
    story.append(PageBreak())

    # Movements detail
    story.append(Paragraph("Listado general de movimientos", section_style))
    mov_rows = [["Fecha", "Tipo", "Tercero", "Documento", "Concepto", "Cuenta", "Proyecto", "C. costo", "Neto"]]
    for m in datos["movimientos"]:
        mov_rows.append([
            m.fecha.strftime("%d/%m/%Y") if m.fecha else "",
            Paragraph(str(m.tipo or ""), small),
            Paragraph(str(m.nombre_tercero or ""), small),
            Paragraph(str(m.numero_documento or ""), small),
            Paragraph(str(m.concepto or ""), small),
            Paragraph(str(m.cuenta or ""), small),
            Paragraph(str(m.proyecto or ""), small),
            Paragraph(str(m.centro_costo or ""), small),
            money(m.valor_neto or m.valor or 0),
        ])
    if len(mov_rows) == 1:
        mov_rows.append(["", "Sin movimientos", "", "", "", "", "", "", "-"])
    mov_table = Table(mov_rows, repeatRows=1, colWidths=[19*mm, 25*mm, 40*mm, 24*mm, 45*mm, 27*mm, 27*mm, 27*mm, 26*mm])
    mov_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17304F")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 6),
        ("ALIGN", (-1,1), (-1,-1), "RIGHT"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E1E8")),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
        ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(mov_table)
    story.append(PageBreak())

    # CxP detail
    story.append(Paragraph("Detalle de cuentas por pagar", section_style))
    cxp_rows = [["Fecha", "Proveedor", "Documento", "Vence", "Neto", "Pagado", "Saldo", "Estado"]]
    for cta in datos["cuentas"]:
        prov = cta.proveedor
        cxp_rows.append([
            cta.fecha_documento.strftime("%d/%m/%Y") if cta.fecha_documento else "",
            Paragraph(prov.razon_social if prov else "", small),
            Paragraph(str(cta.numero_documento or ""), small),
            cta.fecha_vencimiento.strftime("%d/%m/%Y") if cta.fecha_vencimiento else "",
            money(cta.valor_neto), money(cta.total_pagado), money(cta.saldo_pendiente), cta.estado
        ])
    if len(cxp_rows) == 1:
        cxp_rows.append(["", "Sin cuentas por pagar", "", "", "-", "-", "-", ""])
    cxp_table = Table(cxp_rows, repeatRows=1, colWidths=[20*mm, 60*mm, 32*mm, 24*mm, 29*mm, 29*mm, 29*mm, 25*mm])
    cxp_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17304F")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 6.5),
        ("ALIGN", (4,1), (6,-1), "RIGHT"),
        ("ALIGN", (7,1), (7,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E1E8")),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story.append(cxp_table)

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawString(12*mm, 7*mm, "CRV CONSTRUCTORA - Reporte generado por el sistema contable")
        canvas.drawRightString(landscape(A4)[0] - 12*mm, 7*mm, f"Página {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buffer.seek(0)

    nombre = f"CRV_reporte_contable_{date.today().isoformat()}.pdf"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre,
        mimetype="application/pdf"
    )
# ============================================================
# BANCOS Y CAJA - FUNCIONES AUXILIARES
# ============================================================

def _nombre_cuenta_valido(valor):
    nombre = str(valor or "").strip()
    if not nombre:
        return ""
    if nombre.upper() in (
        "CUENTAS POR PAGAR",
        "CUENTAS POR COBRAR",
    ):
        return ""
    return nombre


def _valor_movimiento_bancario(movimiento):
    return float(
        movimiento.valor_neto
        or movimiento.valor
        or 0
    )


def _cuentas_financieras_crv(empresa_id):
    """
    Obtiene las cuentas financieras usadas por CRV sin crear
    una tabla adicional. Esto permite que los pagos ya registrados,
    movimientos históricos y transferencias aparezcan en el módulo.
    """
    movimientos = CRVMovimiento.query.filter_by(
        empresa_id=empresa_id
    ).order_by(
        CRVMovimiento.fecha.asc(),
        CRVMovimiento.id.asc()
    ).all()

    cuentas = {}

    for m in movimientos:
        tipo = str(m.tipo or "").strip().upper()

        # La cuenta de origen sí representa una cuenta propia de CRV.
        origen = _nombre_cuenta_valido(m.cuenta)

        if origen and origen not in cuentas:
            cuentas[origen] = {
                "nombre": origen,
                "saldo": 0.0,
                "ingresos": 0.0,
                "egresos": 0.0,
                "transferencias_entrada": 0.0,
                "transferencias_salida": 0.0,
                "movimientos": 0,
                "sin_conciliar": 0,
            }

        # La cuenta destino SOLO se considera cuenta propia cuando
        # el movimiento es una transferencia interna entre cuentas CRV.
        # En pagos a proveedores, cuenta_destino corresponde a la cuenta
        # bancaria del proveedor y NO debe aparecer como cuenta de la empresa.
        if tipo == "TRANSFERENCIA":
            destino = _nombre_cuenta_valido(m.cuenta_destino)

            if destino and destino not in cuentas:
                cuentas[destino] = {
                    "nombre": destino,
                    "saldo": 0.0,
                    "ingresos": 0.0,
                    "egresos": 0.0,
                    "transferencias_entrada": 0.0,
                    "transferencias_salida": 0.0,
                    "movimientos": 0,
                    "sin_conciliar": 0,
                }

    for m in movimientos:
        tipo = str(m.tipo or "").strip().upper()
        origen = _nombre_cuenta_valido(m.cuenta)
        destino = _nombre_cuenta_valido(m.cuenta_destino)
        valor = _valor_movimiento_bancario(m)

        if tipo == "SALDO INICIAL CUENTA":
            if origen:
                cuentas[origen]["saldo"] += valor
                cuentas[origen]["ingresos"] += valor
                cuentas[origen]["movimientos"] += 1
                if not bool(m.revisado):
                    cuentas[origen]["sin_conciliar"] += 1
            continue

        if tipo in ("INGRESO", "COBRO CUENTA POR COBRAR"):
            if origen:
                cuentas[origen]["saldo"] += valor
                cuentas[origen]["ingresos"] += valor
                cuentas[origen]["movimientos"] += 1
                if not bool(m.revisado):
                    cuentas[origen]["sin_conciliar"] += 1
            continue

        if tipo == "TRANSFERENCIA":
            if origen:
                cuentas[origen]["saldo"] -= valor
                cuentas[origen]["transferencias_salida"] += valor
                cuentas[origen]["movimientos"] += 1
                if not bool(m.revisado):
                    cuentas[origen]["sin_conciliar"] += 1

            if destino:
                cuentas[destino]["saldo"] += valor
                cuentas[destino]["transferencias_entrada"] += valor
                cuentas[destino]["movimientos"] += 1
                if not bool(m.revisado):
                    cuentas[destino]["sin_conciliar"] += 1
            continue

        # La COMPRA de una CxP no mueve banco al registrarse.
        if tipo == "COMPRA":
            forma_pago = str(m.forma_pago or "").strip().upper()
            if forma_pago in ("CREDITO", "CRÉDITO") or not origen:
                continue

        # Todos los demás movimientos con cuenta de origen
        # representan salida real de esa cuenta.
        if origen:
            cuentas[origen]["saldo"] -= valor
            cuentas[origen]["egresos"] += valor
            cuentas[origen]["movimientos"] += 1
            if not bool(m.revisado):
                cuentas[origen]["sin_conciliar"] += 1

    return sorted(
        cuentas.values(),
        key=lambda x: x["nombre"].upper()
    )


def _movimiento_es_bancario(movimiento):
    tipo = str(movimiento.tipo or "").strip().upper()

    # Todo movimiento con cuenta de origen pertenece al control bancario.
    if _nombre_cuenta_valido(movimiento.cuenta):
        return True

    # Una cuenta destino solo pertenece al control bancario cuando se trata
    # de una transferencia interna entre cuentas propias de CRV.
    if (
        tipo == "TRANSFERENCIA"
        and _nombre_cuenta_valido(movimiento.cuenta_destino)
    ):
        return True

    return False


def _impacto_movimiento_cuenta(movimiento, cuenta_nombre):
    """
    Retorna el impacto de un movimiento en una cuenta concreta.
    Positivo = entrada.
    Negativo = salida.
    """
    tipo = str(movimiento.tipo or "").strip().upper()
    origen = _nombre_cuenta_valido(movimiento.cuenta)
    destino = _nombre_cuenta_valido(movimiento.cuenta_destino)
    valor = _valor_movimiento_bancario(movimiento)

    impacto = 0.0

    if tipo == "SALDO INICIAL CUENTA":
        if origen == cuenta_nombre:
            impacto += valor
        return impacto

    if tipo in ("INGRESO", "COBRO CUENTA POR COBRAR"):
        if origen == cuenta_nombre:
            impacto += valor
        return impacto

    if tipo == "TRANSFERENCIA":
        if origen == cuenta_nombre:
            impacto -= valor
        if destino == cuenta_nombre:
            impacto += valor
        return impacto

    if tipo == "COMPRA":
        forma_pago = str(movimiento.forma_pago or "").strip().upper()
        if forma_pago in ("CREDITO", "CRÉDITO"):
            return 0.0

    if origen == cuenta_nombre:
        impacto -= valor

    return impacto


# ============================================================
# BANCOS Y CAJA - PANTALLA PRINCIPAL
# ============================================================

@crv_bp.route(
    "/bancos-caja",
    methods=["GET", "POST"]
)
def bancos_caja():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    if request.method == "POST":

        accion = request.form.get(
            "accion",
            ""
        ).strip().lower()

        # ----------------------------------------------------
        # CREAR CUENTA / SALDO INICIAL
        # ----------------------------------------------------

        if accion == "crear_cuenta":

            nombre = request.form.get(
                "nombre_cuenta",
                ""
            ).strip()

            tipo_cuenta = request.form.get(
                "tipo_cuenta",
                "BANCO"
            ).strip().upper()

            fecha = convertir_fecha(
                request.form.get("fecha_apertura")
            )

            saldo_inicial = convertir_numero(
                request.form.get("saldo_inicial")
            )

            observaciones = request.form.get(
                "observaciones",
                ""
            ).strip()

            if not nombre:
                flash(
                    "Debes escribir el nombre de la cuenta.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            if not fecha:
                flash(
                    "Debes indicar la fecha del saldo inicial.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            cuentas_existentes = {
                fila["nombre"].strip().upper()
                for fila in _cuentas_financieras_crv(empresa_id)
            }

            if nombre.strip().upper() in cuentas_existentes:
                flash(
                    "Ya existe una cuenta financiera con ese nombre.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            registro = CRVMovimiento(
                empresa_id=empresa_id,
                sede_id=None,
                fecha=fecha,
                tipo="SALDO INICIAL CUENTA",
                categoria="BANCOS Y CAJA",
                subcategoria=tipo_cuenta,
                concepto=f"Saldo inicial - {nombre}",
                descripcion=(
                    f"Apertura/control de cuenta financiera {nombre}"
                ),
                tercero_id=None,
                tercero_nombre="",
                tercero_identificacion="",
                tercero_direccion="",
                tercero_email="",
                tercero_telefono="",
                tercero_contacto="",
                tipo_documento="SALDO INICIAL",
                numero_documento="",
                numero_soporte=(
                    f"BC-INICIAL-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
                ),
                fecha_documento=fecha,
                fecha_vencimiento=None,
                valor=saldo_inicial,
                descuento=0,
                retencion=0,
                otros_descuentos=0,
                valor_neto=saldo_inicial,
                forma_pago=tipo_cuenta,
                cuenta=nombre,
                cuenta_destino="",
                beneficiario="CRV CONSTRUCTORA",
                referencia_pago="",
                proyecto="",
                centro_costo="",
                obra="",
                observaciones=observaciones,
                revisado=False,
                estado="ACTIVO",
                origen_registro="BANCOS_CAJA_SALDO_INICIAL",
                usuario_registro=str(session.get("usuario", ""))
            )

            try:
                _asignar_comprobante_permanente(
                    registro
                )

                db.session.add(registro)
                db.session.commit()

                flash(
                    "Cuenta financiera creada correctamente.",
                    "success"
                )

            except Exception as error:
                db.session.rollback()
                flash(
                    f"No fue posible crear la cuenta: {error}",
                    "error"
                )

            return redirect("/crv/bancos-caja")

        # ----------------------------------------------------
        # TRANSFERENCIA ENTRE CUENTAS
        # ----------------------------------------------------

        if accion == "transferencia":

            fecha = convertir_fecha(
                request.form.get("fecha")
            )

            cuenta_origen = request.form.get(
                "cuenta_origen",
                ""
            ).strip()

            cuenta_destino = request.form.get(
                "cuenta_destino",
                ""
            ).strip()

            valor = convertir_numero(
                request.form.get("valor")
            )

            referencia = request.form.get(
                "referencia_pago",
                ""
            ).strip()

            observaciones = request.form.get(
                "observaciones",
                ""
            ).strip()

            if not fecha:
                flash(
                    "La fecha real de la transferencia es obligatoria.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            if not cuenta_origen or not cuenta_destino:
                flash(
                    "Debes seleccionar cuenta de origen y destino.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            if cuenta_origen == cuenta_destino:
                flash(
                    "La cuenta de origen y destino deben ser diferentes.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            if valor <= 0:
                flash(
                    "El valor de la transferencia debe ser mayor a cero.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            cuentas_actuales = {
                fila["nombre"]: fila
                for fila in _cuentas_financieras_crv(empresa_id)
            }

            if (
                cuenta_origen not in cuentas_actuales
                or cuenta_destino not in cuentas_actuales
            ):
                flash(
                    "Una de las cuentas seleccionadas no existe.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            saldo_origen = float(
                cuentas_actuales[cuenta_origen]["saldo"] or 0
            )

            if valor > saldo_origen:
                flash(
                    "La transferencia supera el saldo calculado de la cuenta de origen.",
                    "error"
                )
                return redirect("/crv/bancos-caja")

            movimiento = CRVMovimiento(
                empresa_id=empresa_id,
                sede_id=None,
                fecha=fecha,
                tipo="TRANSFERENCIA",
                categoria="BANCOS Y CAJA",
                subcategoria="TRANSFERENCIA INTERNA",
                concepto=(
                    f"Transferencia {cuenta_origen} → {cuenta_destino}"
                ),
                descripcion=(
                    f"Transferencia entre cuentas propias de CRV: "
                    f"{cuenta_origen} a {cuenta_destino}"
                ),
                tercero_id=None,
                tercero_nombre="CRV CONSTRUCTORA",
                tercero_identificacion="",
                tercero_direccion="",
                tercero_email="",
                tercero_telefono="",
                tercero_contacto="",
                tipo_documento="TRANSFERENCIA INTERNA",
                numero_documento="",
                numero_soporte=(
                    f"BC-TRF-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
                ),
                fecha_documento=fecha,
                fecha_vencimiento=None,
                valor=valor,
                descuento=0,
                retencion=0,
                otros_descuentos=0,
                valor_neto=valor,
                forma_pago="TRANSFERENCIA",
                cuenta=cuenta_origen,
                cuenta_destino=cuenta_destino,
                beneficiario="CRV CONSTRUCTORA",
                referencia_pago=referencia,
                proyecto="",
                centro_costo="",
                obra="",
                observaciones=observaciones,
                revisado=False,
                estado="ACTIVO",
                origen_registro="BANCOS_CAJA_TRANSFERENCIA",
                usuario_registro=str(session.get("usuario", ""))
            )

            try:
                _asignar_comprobante_permanente(
                    movimiento
                )

                db.session.add(movimiento)
                db.session.commit()

                flash(
                    "Transferencia registrada correctamente. No afecta el resultado financiero.",
                    "success"
                )

            except Exception as error:
                db.session.rollback()
                flash(
                    f"No fue posible registrar la transferencia: {error}",
                    "error"
                )

            return redirect("/crv/bancos-caja")

    # --------------------------------------------------------
    # CONSULTA
    # --------------------------------------------------------

    mes = request.args.get(
        "mes",
        ""
    ).strip()

    cuenta_filtro = request.args.get(
        "cuenta",
        ""
    ).strip()

    cuentas = _cuentas_financieras_crv(
        empresa_id
    )

    movimientos_query = CRVMovimiento.query.filter_by(
        empresa_id=empresa_id
    )

    movimientos = movimientos_query.order_by(
        CRVMovimiento.fecha.desc(),
        CRVMovimiento.id.desc()
    ).all()

    movimientos_bancarios = [
        m
        for m in movimientos
        if _movimiento_es_bancario(m)
    ]

    if mes:
        try:
            anio_texto, mes_texto = mes.split("-")
            anio = int(anio_texto)
            numero_mes = int(mes_texto)

            movimientos_bancarios = [
                m
                for m in movimientos_bancarios
                if (
                    m.fecha
                    and m.fecha.year == anio
                    and m.fecha.month == numero_mes
                )
            ]

        except (ValueError, TypeError):
            pass

    if cuenta_filtro:
        movimientos_bancarios = [
            m
            for m in movimientos_bancarios
            if (
                _nombre_cuenta_valido(m.cuenta) == cuenta_filtro
                or _nombre_cuenta_valido(m.cuenta_destino) == cuenta_filtro
            )
        ]

    entradas_periodo = 0.0
    salidas_periodo = 0.0
    transferencias_periodo = 0.0
    sin_conciliar_periodo = 0

    for m in movimientos_bancarios:
        tipo = str(m.tipo or "").strip().upper()
        valor = _valor_movimiento_bancario(m)

        if tipo == "INGRESO":
            entradas_periodo += valor

        elif tipo == "SALDO INICIAL CUENTA":
            entradas_periodo += valor

        elif tipo == "TRANSFERENCIA":
            transferencias_periodo += valor

        else:
            salidas_periodo += valor

        if not bool(m.revisado):
            sin_conciliar_periodo += 1

    saldo_total = sum(
        float(c["saldo"] or 0)
        for c in cuentas
    )

    return render_template(
        "contabilidad/crv_bancos_caja.html",
        empresa_nombre=empresa_nombre,
        cuentas=cuentas,
        movimientos=movimientos_bancarios,
        mes_actual=mes,
        cuenta_actual=cuenta_filtro,
        saldo_total=saldo_total,
        entradas_periodo=entradas_periodo,
        salidas_periodo=salidas_periodo,
        transferencias_periodo=transferencias_periodo,
        sin_conciliar_periodo=sin_conciliar_periodo
    )


# ============================================================
# BANCOS Y CAJA - MARCAR / DESMARCAR CONCILIADO
# ============================================================

@crv_bp.route(
    "/bancos-caja/conciliar/<int:movimiento_id>",
    methods=["POST"]
)
def bancos_caja_conciliar(movimiento_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    movimiento = CRVMovimiento.query.get_or_404(
        movimiento_id
    )

    if movimiento.empresa_id != empresa_id:
        return redirect("/crv/bancos-caja")

    if not _movimiento_es_bancario(movimiento):
        flash(
            "Ese registro no corresponde a Bancos y caja.",
            "error"
        )
        return redirect("/crv/bancos-caja")

    movimiento.revisado = not bool(
        movimiento.revisado
    )

    db.session.commit()

    flash(
        "Estado de conciliación actualizado.",
        "success"
    )

    return redirect(
        request.referrer
        or "/crv/bancos-caja"
    )


# ============================================================
# BANCOS Y CAJA - ELIMINAR REGISTRO PROPIO DEL MÓDULO
# ============================================================

@crv_bp.route(
    "/bancos-caja/eliminar/<int:movimiento_id>",
    methods=["POST"]
)
def bancos_caja_eliminar(movimiento_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    movimiento = CRVMovimiento.query.get_or_404(
        movimiento_id
    )

    if movimiento.empresa_id != empresa_id:
        return redirect("/crv/bancos-caja")

    if movimiento.origen_registro not in (
        "BANCOS_CAJA_SALDO_INICIAL",
        "BANCOS_CAJA_TRANSFERENCIA",
    ):
        flash(
            "Este movimiento proviene de otro módulo y debe corregirse desde su origen.",
            "warning"
        )
        return redirect("/crv/bancos-caja")

    try:
        db.session.delete(movimiento)
        db.session.commit()

        flash(
            "Registro bancario eliminado correctamente.",
            "success"
        )

    except Exception as error:
        db.session.rollback()
        flash(
            f"No fue posible eliminar el registro: {error}",
            "error"
        )

    return redirect("/crv/bancos-caja")

# ============================================================
# EXPORTACIONES CRV - FUNCIONES AUXILIARES
# ============================================================

def _excel_estilos_base():
    return {
        "navy": "17304F",
        "gold": "D8B541",
        "green": "0F8A47",
        "red": "D92D20",
        "light": "EEF3F8",
        "white": "FFFFFF",
        "gray": "667085",
    }


def _excel_titulo(sheet, titulo, subtitulo="", columnas=10):
    estilos = _excel_estilos_base()
    navy = estilos["navy"]
    gray = estilos["gray"]

    sheet.sheet_view.showGridLines = False

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(2, columnas)
    )

    c = sheet.cell(1, 1)
    c.value = titulo
    c.font = Font(
        size=18,
        bold=True,
        color=navy
    )
    c.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    if subtitulo:
        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=max(2, columnas)
        )
        c2 = sheet.cell(2, 1)
        c2.value = subtitulo
        c2.font = Font(
            size=10,
            color=gray
        )

    sheet.row_dimensions[1].height = 28


def _excel_encabezado(sheet, fila, headers):
    estilos = _excel_estilos_base()

    for col, valor in enumerate(headers, 1):
        cell = sheet.cell(fila, col, valor)
        cell.fill = PatternFill(
            "solid",
            fgColor=estilos["navy"]
        )
        cell.font = Font(
            bold=True,
            color=estilos["white"]
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    sheet.row_dimensions[fila].height = 28


def _excel_moneda(cell):
    cell.number_format = '$#,##0;[Red]($#,##0);-'
    cell.alignment = Alignment(
        horizontal="right"
    )


def _excel_auto_ancho(sheet, max_width=40):
    for col in range(
        1,
        sheet.max_column + 1
    ):
        letter = get_column_letter(col)
        max_len = 0

        for row in range(
            1,
            min(sheet.max_row, 1000) + 1
        ):
            value = sheet.cell(
                row,
                col
            ).value

            if value is not None:
                max_len = max(
                    max_len,
                    len(str(value))
                )

        sheet.column_dimensions[
            letter
        ].width = min(
            max(max_len + 2, 10),
            max_width
        )


def _enviar_excel(wb, nombre_archivo):
    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


def _pdf_estilos_crv():
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CRVExportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=21,
        textColor=colors.HexColor("#17304F"),
        alignment=TA_CENTER
    )

    subtitle_style = ParagraphStyle(
        "CRVExportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#667085"),
        alignment=TA_CENTER
    )

    section_style = ParagraphStyle(
        "CRVExportSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        textColor=colors.HexColor("#17304F"),
        spaceBefore=5,
        spaceAfter=7
    )

    small_style = ParagraphStyle(
        "CRVExportSmall",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=6.7,
        leading=8,
        textColor=colors.HexColor("#25364A"),
        alignment=TA_LEFT
    )

    return (
        title_style,
        subtitle_style,
        section_style,
        small_style
    )


def _pdf_encabezado_historia(story, titulo, subtitulo=""):
    (
        title_style,
        subtitle_style,
        _,
        _
    ) = _pdf_estilos_crv()

    logo_path = (
        Path(current_app.root_path)
        / "static"
        / "logo_crv.png"
    )

    if logo_path.exists():
        story.append(
            Image(
                str(logo_path),
                width=24 * mm,
                height=18 * mm
            )
        )

    story.append(
        Paragraph(
            titulo,
            title_style
        )
    )

    if subtitulo:
        story.append(
            Paragraph(
                subtitulo,
                subtitle_style
            )
        )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )


def _pdf_tabla(headers, filas, widths=None):
    _, _, _, small = _pdf_estilos_crv()

    data = [
        [
            Paragraph(
                str(h),
                small
            )
            for h in headers
        ]
    ]

    for fila in filas:
        data.append([
            Paragraph(
                str(v if v not in (None, "") else "-"),
                small
            )
            for v in fila
        ])

    tabla = Table(
        data,
        colWidths=widths,
        repeatRows=1,
        hAlign="LEFT"
    )

    tabla.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#17304F")
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#D9E1E8")
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F8FAFC")
                ]
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                4
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                4
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                4
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                4
            ),
        ])
    )

    return tabla


def _enviar_pdf(story, nombre_archivo, horizontal=True):
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=(
            landscape(A4)
            if horizontal
            else A4
        ),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm
    )

    doc.build(story)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype="application/pdf"
    )


def _money_export(v):
    return f"${float(v or 0):,.0f}"


# ============================================================
# CENTRO DE EXPORTACIONES
# ============================================================

@crv_bp.route("/exportaciones")
def exportaciones_crv():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    return render_template(
        "contabilidad/crv_exportaciones.html",
        empresa_nombre=session.get(
            "empresa_nombre",
            "CRV CONSTRUCTORA"
        )
    )


# ============================================================
# MOVIMIENTOS - EXCEL
# ============================================================

@crv_bp.route("/movimientos/excel")
def movimientos_excel():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    mes = request.args.get(
        "mes",
        ""
    ).strip()

    query = CRVMovimiento.query.filter_by(
        empresa_id=empresa_id
    )

    if mes:
        try:
            anio, numero_mes = mes.split("-")
            query = query.filter(
                db.extract(
                    "year",
                    CRVMovimiento.fecha
                ) == int(anio),
                db.extract(
                    "month",
                    CRVMovimiento.fecha
                ) == int(numero_mes)
            )
        except ValueError:
            pass

    movimientos = query.order_by(
        CRVMovimiento.fecha.asc(),
        CRVMovimiento.id.asc()
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "MOVIMIENTOS"

    headers = [
        "Fecha",
        "Tipo",
        "Categoría",
        "Subcategoría",
        "Tercero",
        "NIT/ID",
        "Documento",
        "Concepto",
        "Descripción",
        "Cuenta origen CRV",
        "Cuenta destino / beneficiario",
        "Proyecto",
        "Centro costo",
        "Obra",
        "Base gravable",
        "Tipo IVA",
        "Tarifa IVA %",
        "Valor IVA",
        "Total documento",
        "Descuento",
        "Retención",
        "Otros descuentos",
        "Valor neto",
        "Forma pago",
        "Soporte",
        "Referencia",
        "Revisado"
    ]

    _excel_titulo(
        ws,
        "CRV CONSTRUCTORA - MOVIMIENTOS",
        (
            f"Periodo: {mes}"
            if mes
            else "Todos los registros"
        ),
        len(headers)
    )

    _excel_encabezado(
        ws,
        4,
        headers
    )

    for fila, m in enumerate(
        movimientos,
        5
    ):
        valores = [
            m.fecha,
            m.tipo,
            m.categoria,
            m.subcategoria,
            m.nombre_tercero,
            m.tercero_identificacion,
            m.numero_documento,
            m.concepto,
            m.descripcion,
            m.cuenta,
            m.cuenta_destino,
            m.proyecto,
            m.centro_costo,
            m.obra,
            float(m.base_gravable or 0),
            m.tipo_iva or "NO_APLICA",
            float(m.tarifa_iva or 0),
            float(m.valor_iva or 0),
            float(m.valor or 0),
            float(m.descuento or 0),
            float(m.retencion or 0),
            float(m.otros_descuentos or 0),
            float(
                m.valor_neto
                or m.valor
                or 0
            ),
            m.forma_pago,
            m.numero_soporte,
            m.referencia_pago,
            "SI" if m.revisado else "NO"
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            ws.cell(
                fila,
                col,
                valor
            )

        for col in range(
            15,
            20
        ):
            _excel_moneda(
                ws.cell(
                    fila,
                    col
                )
            )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers))}"
        f"{max(4, ws.max_row)}"
    )

    _excel_auto_ancho(ws)

    return _enviar_excel(
        wb,
        (
            f"CRV_movimientos_"
            f"{date.today().isoformat()}.xlsx"
        )
    )



# ============================================================
# MOVIMIENTO INDIVIDUAL - COMPROBANTE PDF
# ============================================================

def _texto_pdf_seguro(valor):
    return escape(
        str(valor or "")
    )


@crv_bp.route(
    "/movimientos/<int:movimiento_id>/comprobante-pdf"
)
def movimiento_comprobante_pdf(movimiento_id):

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    movimiento = (
        CRVMovimiento.query
        .filter_by(
            id=movimiento_id,
            empresa_id=empresa_id
        )
        .first_or_404()
    )

    prefijo = str(
        movimiento.comprobante_tipo
        or ""
    ).strip().upper()

    consecutivo = str(
        movimiento.comprobante_codigo
        or ""
    ).strip().upper()

    # La migracion y las rutas de registro asignan el consecutivo.
    # Abrir el PDF no crea ni modifica numeracion.
    if not prefijo or not consecutivo:
        titulo_calculado, prefijo_calculado = (
            _tipo_comprobante_movimiento(
                movimiento
            )
        )

        prefijo = (
            prefijo
            or prefijo_calculado
        )

        titulo_comprobante = (
            _titulo_comprobante_por_prefijo(
                prefijo
            )
        )

        consecutivo = (
            consecutivo
            or "SIN-CONSECUTIVO"
        )
    else:
        titulo_comprobante = (
            _titulo_comprobante_por_prefijo(
                prefijo
            )
        )

    def dinero(valor):
        return (
            f"$"
            f"{float(valor or 0):,.0f}"
        )

    def fecha_texto(valor):
        if not valor:
            return "-"

        try:
            return valor.strftime(
                "%d/%m/%Y"
            )
        except Exception:
            return str(valor)

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=(
            f"{titulo_comprobante} "
            f"{consecutivo}"
        ),
        author=empresa_nombre
    )

    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        "ComprobanteTitulo",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=19,
        textColor=colors.HexColor(
            "#17304F"
        ),
        alignment=TA_CENTER,
        spaceAfter=3
    )

    subtitulo_style = ParagraphStyle(
        "ComprobanteSubtitulo",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor(
            "#667085"
        ),
        alignment=TA_CENTER
    )

    label_style = ParagraphStyle(
        "ComprobanteLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.2,
        leading=9,
        textColor=colors.HexColor(
            "#17304F"
        )
    )

    value_style = ParagraphStyle(
        "ComprobanteValue",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor(
            "#25364A"
        )
    )

    value_bold_style = ParagraphStyle(
        "ComprobanteValueBold",
        parent=value_style,
        fontName="Helvetica-Bold"
    )

    obs_style = ParagraphStyle(
        "ComprobanteObs",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor(
            "#25364A"
        )
    )

    story = []

    logo_path = (
        Path(current_app.root_path)
        / "static"
        / "logo_crv.png"
    )

    if logo_path.exists():
        story.append(
            Image(
                str(logo_path),
                width=25 * mm,
                height=19 * mm
            )
        )

    story.append(
        Paragraph(
            _texto_pdf_seguro(
                empresa_nombre
            ),
            titulo_style
        )
    )

    story.append(
        Paragraph(
            titulo_comprobante,
            titulo_style
        )
    )

    story.append(
        Paragraph(
            (
                f"Consecutivo: <b>{consecutivo}</b>"
                f" &nbsp;&nbsp;|&nbsp;&nbsp; "
                f"Fecha contable: "
                f"<b>{fecha_texto(movimiento.fecha)}</b>"
            ),
            subtitulo_style
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )

    # --------------------------------------------------------
    # IDENTIFICACIÓN DEL TERCERO Y DOCUMENTO
    # --------------------------------------------------------

    tercero_nombre = (
        movimiento.nombre_tercero
        or movimiento.tercero_nombre
        or movimiento.beneficiario
        or "-"
    )

    identificacion = (
        movimiento.tercero_identificacion
        or "-"
    )

    documento = (
        movimiento.numero_documento
        or movimiento.numero_soporte
        or "-"
    )

    tipo_documento = (
        movimiento.tipo_documento
        or "-"
    )

    datos_generales = [
        [
            Paragraph(
                "Tercero / beneficiario",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    tercero_nombre
                ),
                value_bold_style
            ),
            Paragraph(
                "NIT / Identificación",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    identificacion
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Tipo de documento",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    tipo_documento
                ),
                value_style
            ),
            Paragraph(
                "Documento / soporte",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    documento
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Fecha documento",
                label_style
            ),
            Paragraph(
                fecha_texto(
                    movimiento.fecha_documento
                ),
                value_style
            ),
            Paragraph(
                "Fecha vencimiento",
                label_style
            ),
            Paragraph(
                fecha_texto(
                    movimiento.fecha_vencimiento
                ),
                value_style
            ),
        ],
    ]

    tabla_generales = Table(
        datos_generales,
        colWidths=[
            34 * mm,
            58 * mm,
            34 * mm,
            49 * mm
        ]
    )

    tabla_generales.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (0, -1),
                colors.HexColor("#EEF3F8")
            ),
            (
                "BACKGROUND",
                (2, 0),
                (2, -1),
                colors.HexColor("#EEF3F8")
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#D9E1E8")
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    story.append(
        tabla_generales
    )

    story.append(
        Spacer(
            1,
            4 * mm
        )
    )

    # --------------------------------------------------------
    # CONCEPTO / CLASIFICACIÓN
    # --------------------------------------------------------

    clasificacion = [
        [
            Paragraph(
                "Tipo",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.tipo
                ),
                value_style
            ),
            Paragraph(
                "Categoría",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.categoria
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Concepto",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.concepto
                ),
                value_bold_style
            ),
            Paragraph(
                "Subcategoría",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.subcategoria
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Proyecto",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.proyecto
                ),
                value_style
            ),
            Paragraph(
                "Centro de costo",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.centro_costo
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Obra / frente",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.obra
                ),
                value_style
            ),
            Paragraph(
                "Origen del registro",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.origen_registro
                ),
                value_style
            ),
        ],
    ]

    tabla_clasificacion = Table(
        clasificacion,
        colWidths=[
            27 * mm,
            65 * mm,
            30 * mm,
            53 * mm
        ]
    )

    tabla_clasificacion.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (0, -1),
                colors.HexColor("#FFF7DB")
            ),
            (
                "BACKGROUND",
                (2, 0),
                (2, -1),
                colors.HexColor("#FFF7DB")
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#D9E1E8")
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP"
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    story.append(
        tabla_clasificacion
    )

    story.append(
        Spacer(
            1,
            4 * mm
        )
    )

    # --------------------------------------------------------
    # VALORES / IVA
    # --------------------------------------------------------

    total_documento = float(
        movimiento.valor or 0
    )

    valor_neto = float(
        movimiento.valor_neto
        or movimiento.valor
        or 0
    )

    total_descuentos = (
        float(movimiento.descuento or 0)
        + float(movimiento.retencion or 0)
        + float(
            movimiento.otros_descuentos
            or 0
        )
    )

    iva_detalle = (
        f"{float(movimiento.tarifa_iva or 0):g}%"
        f" · "
        f"{_texto_pdf_seguro(movimiento.tipo_iva or 'NO_APLICA')}"
    )

    valores = [
        [
            "Base gravable",
            dinero(
                movimiento.base_gravable
            ),
            "IVA",
            dinero(
                movimiento.valor_iva
            ),
        ],
        [
            "Tratamiento IVA",
            iva_detalle,
            "Total documento",
            dinero(
                total_documento
            ),
        ],
        [
            "Descuento",
            dinero(
                movimiento.descuento
            ),
            "Retención",
            dinero(
                movimiento.retencion
            ),
        ],
        [
            "Otros descuentos",
            dinero(
                movimiento.otros_descuentos
            ),
            "Total descuentos",
            dinero(
                total_descuentos
            ),
        ],
        [
            "VALOR NETO",
            dinero(
                valor_neto
            ),
            "",
            "",
        ],
    ]

    datos_valores = []

    for fila in valores:
        datos_valores.append([
            Paragraph(
                str(fila[0]),
                label_style
            ),
            Paragraph(
                str(fila[1]),
                (
                    value_bold_style
                    if fila[0] == "VALOR NETO"
                    else value_style
                )
            ),
            Paragraph(
                str(fila[2]),
                label_style
            ),
            Paragraph(
                str(fila[3]),
                value_style
            ),
        ])

    tabla_valores = Table(
        datos_valores,
        colWidths=[
            34 * mm,
            53 * mm,
            34 * mm,
            54 * mm
        ]
    )

    tabla_valores.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (0, -2),
                colors.HexColor("#EEF3F8")
            ),
            (
                "BACKGROUND",
                (2, 0),
                (2, -2),
                colors.HexColor("#EEF3F8")
            ),
            (
                "BACKGROUND",
                (0, -1),
                (1, -1),
                colors.HexColor("#FFF4C7")
            ),
            (
                "SPAN",
                (1, -1),
                (3, -1)
            ),
            (
                "GRID",
                (0, 0),
                (-1, -2),
                0.4,
                colors.HexColor("#D9E1E8")
            ),
            (
                "BOX",
                (0, -1),
                (3, -1),
                0.6,
                colors.HexColor("#D8B541")
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    story.append(
        tabla_valores
    )

    story.append(
        Spacer(
            1,
            4 * mm
        )
    )

    # --------------------------------------------------------
    # PAGO / CUENTAS
    # --------------------------------------------------------

    pago = [
        [
            Paragraph(
                "Forma de pago",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.forma_pago
                ),
                value_style
            ),
            Paragraph(
                "Cuenta CRV",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.cuenta
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Cuenta destino",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.cuenta_destino
                ),
                value_style
            ),
            Paragraph(
                "Beneficiario",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.beneficiario
                    or tercero_nombre
                ),
                value_style
            ),
        ],
        [
            Paragraph(
                "Referencia",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.referencia_pago
                ),
                value_style
            ),
            Paragraph(
                "Soporte",
                label_style
            ),
            Paragraph(
                _texto_pdf_seguro(
                    movimiento.numero_soporte
                ),
                value_style
            ),
        ],
    ]

    tabla_pago = Table(
        pago,
        colWidths=[
            31 * mm,
            56 * mm,
            31 * mm,
            57 * mm
        ]
    )

    tabla_pago.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (0, -1),
                colors.HexColor("#EEF3F8")
            ),
            (
                "BACKGROUND",
                (2, 0),
                (2, -1),
                colors.HexColor("#EEF3F8")
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#D9E1E8")
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    story.append(
        tabla_pago
    )

    # --------------------------------------------------------
    # DESCRIPCIÓN Y OBSERVACIONES
    # --------------------------------------------------------

    descripcion = str(
        movimiento.descripcion
        or ""
    ).strip()

    observaciones = str(
        movimiento.observaciones
        or ""
    ).strip()

    if descripcion:
        story.append(
            Spacer(
                1,
                4 * mm
            )
        )

        story.append(
            Paragraph(
                "<b>Descripción:</b> "
                + _texto_pdf_seguro(
                    descripcion
                ),
                obs_style
            )
        )

    if observaciones:
        story.append(
            Spacer(
                1,
                2 * mm
            )
        )

        story.append(
            Paragraph(
                "<b>Observaciones:</b> "
                + _texto_pdf_seguro(
                    observaciones
                ),
                obs_style
            )
        )

    # --------------------------------------------------------
    # FIRMAS
    # --------------------------------------------------------

    story.append(
        Spacer(
            1,
            13 * mm
        )
    )

    firmas = Table(
        [
            [
                "____________________________",
                "____________________________",
                "____________________________",
            ],
            [
                "Elaborado por",
                "Revisado por",
                (
                    "Recibido por"
                    if titulo_comprobante
                    == "COMPROBANTE DE EGRESO"
                    else "Responsable"
                ),
            ],
            [
                _texto_pdf_seguro(
                    movimiento.usuario_registro
                    or session.get(
                        "usuario",
                        ""
                    )
                ),
                "",
                "",
            ],
        ],
        colWidths=[
            58 * mm,
            58 * mm,
            58 * mm
        ]
    )

    firmas.setStyle(
        TableStyle([
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),
            (
                "FONTNAME",
                (0, 1),
                (-1, 1),
                "Helvetica-Bold"
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7.5
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, -1),
                colors.HexColor("#25364A")
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                2
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                2
            ),
        ])
    )

    story.append(
        firmas
    )

    def pie_pagina(
        canvas,
        doc_obj
    ):
        canvas.saveState()

        ancho, _ = A4

        canvas.setStrokeColor(
            colors.HexColor(
                "#D8B541"
            )
        )

        canvas.setLineWidth(
            0.8
        )

        canvas.line(
            15 * mm,
            10 * mm,
            ancho - 15 * mm,
            10 * mm
        )

        canvas.setFont(
            "Helvetica",
            6.5
        )

        canvas.setFillColor(
            colors.HexColor(
                "#667085"
            )
        )

        canvas.drawString(
            15 * mm,
            6.5 * mm,
            (
                "Documento interno generado por "
                "el sistema contable de CRV."
            )
        )

        canvas.drawRightString(
            ancho - 15 * mm,
            6.5 * mm,
            (
                f"{consecutivo} · "
                f"Página {doc_obj.page}"
            )
        )

        canvas.restoreState()

    doc.build(
        story,
        onFirstPage=pie_pagina,
        onLaterPages=pie_pagina
    )

    buffer.seek(
        0
    )

    nombre_archivo = (
        f"{consecutivo}_"
        f"{date.today().isoformat()}.pdf"
    )

    return send_file(
        buffer,
        as_attachment=False,
        download_name=nombre_archivo,
        mimetype="application/pdf"
    )



# ============================================================
# MOVIMIENTOS - PDF
# ============================================================

@crv_bp.route("/movimientos/pdf")
def movimientos_pdf():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    movimientos = (
        CRVMovimiento.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )

    filas = []

    for m in movimientos:
        filas.append([
            (
                m.fecha.strftime("%d/%m/%Y")
                if m.fecha
                else ""
            ),
            m.tipo or "",
            m.nombre_tercero or "",
            m.numero_documento or "",
            m.concepto or "",
            m.cuenta or "",
            m.proyecto or "",
            _money_export(
                m.valor_neto
                or m.valor
                or 0
            )
        ])

    story = []

    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - MOVIMIENTOS",
        (
            "Listado general de movimientos contables · "
            f"Generado {date.today().strftime('%d/%m/%Y')}"
        )
    )

    story.append(
        _pdf_tabla(
            [
                "Fecha",
                "Tipo",
                "Tercero",
                "Documento",
                "Concepto",
                "Cuenta",
                "Proyecto",
                "Valor"
            ],
            filas,
            [
                20 * mm,
                28 * mm,
                42 * mm,
                26 * mm,
                48 * mm,
                35 * mm,
                32 * mm,
                26 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            f"CRV_movimientos_"
            f"{date.today().isoformat()}.pdf"
        )
    )


# ============================================================
# TERCEROS - EXCEL
# ============================================================

@crv_bp.route("/terceros/excel")
def terceros_excel():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    terceros = (
        CRVTercero.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVTercero.razon_social.asc()
        )
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "TERCEROS"

    headers = [
        "Tipo tercero",
        "Tipo documento",
        "Identificación",
        "Razón social",
        "Nombre comercial",
        "Dirección",
        "Ciudad",
        "Departamento",
        "Teléfono",
        "Celular",
        "Email",
        "Contacto",
        "Banco",
        "Tipo cuenta",
        "Número cuenta",
        "Titular cuenta",
        "Estado"
    ]

    _excel_titulo(
        ws,
        "CRV CONSTRUCTORA - TERCEROS",
        "Directorio de proveedores, clientes, empleados y terceros",
        len(headers)
    )

    _excel_encabezado(
        ws,
        4,
        headers
    )

    for fila, t in enumerate(
        terceros,
        5
    ):
        valores = [
            t.tipo_tercero,
            t.tipo_documento,
            t.identificacion,
            t.razon_social,
            t.nombre_comercial,
            t.direccion,
            t.ciudad,
            t.departamento,
            t.telefono,
            t.celular,
            t.email,
            t.contacto,
            t.banco,
            t.tipo_cuenta,
            t.numero_cuenta,
            t.titular_cuenta,
            "ACTIVO" if t.activo else "INACTIVO"
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            ws.cell(
                fila,
                col,
                valor
            )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers))}"
        f"{max(4, ws.max_row)}"
    )

    _excel_auto_ancho(ws)

    return _enviar_excel(
        wb,
        (
            f"CRV_terceros_"
            f"{date.today().isoformat()}.xlsx"
        )
    )


# ============================================================
# TERCEROS - PDF
# ============================================================

@crv_bp.route("/terceros/pdf")
def terceros_pdf():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    terceros = (
        CRVTercero.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVTercero.razon_social.asc()
        )
        .all()
    )

    filas = []

    for t in terceros:
        filas.append([
            t.tipo_tercero or "",
            t.razon_social or "",
            t.identificacion or "",
            (
                t.celular
                or t.telefono
                or ""
            ),
            t.email or "",
            t.ciudad or "",
            t.banco or "",
            (
                "ACTIVO"
                if t.activo
                else "INACTIVO"
            )
        ])

    story = []

    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - TERCEROS",
        (
            "Directorio general · "
            f"Generado {date.today().strftime('%d/%m/%Y')}"
        )
    )

    story.append(
        _pdf_tabla(
            [
                "Tipo",
                "Nombre / Razón social",
                "NIT/ID",
                "Contacto",
                "Correo",
                "Ciudad",
                "Banco",
                "Estado"
            ],
            filas,
            [
                25 * mm,
                52 * mm,
                28 * mm,
                30 * mm,
                48 * mm,
                28 * mm,
                30 * mm,
                22 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            f"CRV_terceros_"
            f"{date.today().isoformat()}.pdf"
        )
    )


# ============================================================
# CUENTAS POR PAGAR - EXCEL
# ============================================================

@crv_bp.route("/cuentas-por-pagar/excel")
def cuentas_pagar_excel():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    cuentas = (
        CRVCuentaPagar.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVCuentaPagar.fecha_documento.asc(),
            CRVCuentaPagar.id.asc()
        )
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "CUENTAS_POR_PAGAR"

    headers = [
        "Fecha",
        "Proveedor",
        "NIT/ID",
        "Tipo documento",
        "Documento",
        "Concepto",
        "Proyecto",
        "Centro costo",
        "Vencimiento",
        "Base gravable",
        "Tipo IVA",
        "Tarifa IVA %",
        "Valor IVA",
        "Total documento",
        "Descuento",
        "Retención",
        "Otros descuentos",
        "Neto",
        "Pagado",
        "Saldo",
        "Estado"
    ]

    _excel_titulo(
        ws,
        "CRV CONSTRUCTORA - CUENTAS POR PAGAR",
        (
            "Cartera de proveedores y obligaciones "
            f"· Generado {date.today().strftime('%d/%m/%Y')}"
        ),
        len(headers)
    )

    _excel_encabezado(
        ws,
        4,
        headers
    )

    for fila, cta in enumerate(
        cuentas,
        5
    ):
        prov = cta.proveedor

        valores = [
            cta.fecha_documento,
            (
                prov.razon_social
                if prov
                else ""
            ),
            (
                prov.identificacion
                if prov
                else ""
            ),
            cta.tipo_documento,
            cta.numero_documento,
            cta.concepto,
            cta.proyecto,
            cta.centro_costo,
            cta.fecha_vencimiento,
            float(cta.base_gravable or 0),
            cta.tipo_iva or "NO_APLICA",
            float(cta.tarifa_iva or 0),
            float(cta.valor_iva or 0),
            float(cta.valor_bruto or 0),
            float(cta.descuento or 0),
            float(cta.retencion or 0),
            float(cta.otros_descuentos or 0),
            float(cta.valor_neto or 0),
            float(cta.total_pagado or 0),
            float(cta.saldo_pendiente or 0),
            cta.estado
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            ws.cell(
                fila,
                col,
                valor
            )

        for col in (10, 13, 14, 15, 16, 17, 18, 19, 20):
            _excel_moneda(
                ws.cell(
                    fila,
                    col
                )
            )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers))}"
        f"{max(4, ws.max_row)}"
    )

    _excel_auto_ancho(ws)

    return _enviar_excel(
        wb,
        (
            f"CRV_cuentas_por_pagar_"
            f"{date.today().isoformat()}.xlsx"
        )
    )


# ============================================================
# CUENTAS POR PAGAR - PDF
# ============================================================

@crv_bp.route("/cuentas-por-pagar/pdf")
def cuentas_pagar_pdf():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    cuentas = (
        CRVCuentaPagar.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVCuentaPagar.fecha_documento.asc(),
            CRVCuentaPagar.id.asc()
        )
        .all()
    )

    filas = []

    for cta in cuentas:
        prov = cta.proveedor

        filas.append([
            (
                cta.fecha_documento.strftime(
                    "%d/%m/%Y"
                )
                if cta.fecha_documento
                else ""
            ),
            (
                prov.razon_social
                if prov
                else ""
            ),
            cta.numero_documento or "",
            cta.concepto or "",
            (
                cta.fecha_vencimiento.strftime(
                    "%d/%m/%Y"
                )
                if cta.fecha_vencimiento
                else ""
            ),
            _money_export(
                cta.base_gravable
            ),
            _money_export(
                cta.valor_iva
            ),
            _money_export(
                cta.valor_bruto
            ),
            _money_export(
                cta.valor_neto
            ),
            _money_export(
                cta.total_pagado
            ),
            _money_export(
                cta.saldo_pendiente
            ),
            cta.estado or ""
        ])

    story = []

    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - CUENTAS POR PAGAR",
        (
            "Cartera general de proveedores · "
            f"Generado {date.today().strftime('%d/%m/%Y')}"
        )
    )

    story.append(
        _pdf_tabla(
            [
                "Fecha",
                "Proveedor",
                "Documento",
                "Concepto",
                "Vence",
                "Base",
                "IVA",
                "Total",
                "Neto",
                "Pagado",
                "Saldo",
                "Estado"
            ],
            filas,
            [
                18 * mm,
                38 * mm,
                22 * mm,
                38 * mm,
                18 * mm,
                21 * mm,
                19 * mm,
                21 * mm,
                21 * mm,
                21 * mm,
                21 * mm,
                20 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            f"CRV_cuentas_por_pagar_"
            f"{date.today().isoformat()}.pdf"
        )
    )


# ============================================================
# BANCOS Y CAJA - EXCEL
# ============================================================

@crv_bp.route("/bancos-caja/excel")
def bancos_caja_excel():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    cuentas = _cuentas_financieras_crv(
        empresa_id
    )

    movimientos = (
        CRVMovimiento.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )

    movimientos = [
        m
        for m in movimientos
        if _movimiento_es_bancario(m)
    ]

    wb = Workbook()

    ws = wb.active
    ws.title = "SALDOS"

    _excel_titulo(
        ws,
        "CRV CONSTRUCTORA - BANCOS Y CAJA",
        (
            "Saldos calculados de cuentas propias "
            f"· Generado {date.today().strftime('%d/%m/%Y')}"
        ),
        9
    )

    headers_saldos = [
        "Cuenta CRV",
        "Saldo",
        "Entradas",
        "Egresos",
        "Transferencias entrada",
        "Transferencias salida",
        "Movimientos",
        "Sin conciliar",
        "Estado"
    ]

    _excel_encabezado(
        ws,
        4,
        headers_saldos
    )

    for fila, c in enumerate(
        cuentas,
        5
    ):
        valores = [
            c["nombre"],
            c["saldo"],
            c["ingresos"],
            c["egresos"],
            c["transferencias_entrada"],
            c["transferencias_salida"],
            c["movimientos"],
            c["sin_conciliar"],
            "ACTIVA"
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            ws.cell(
                fila,
                col,
                valor
            )

        for col in range(
            2,
            7
        ):
            _excel_moneda(
                ws.cell(
                    fila,
                    col
                )
            )

    _excel_auto_ancho(ws)

    mov_ws = wb.create_sheet(
        "MOVIMIENTOS_BANCARIOS"
    )

    mov_headers = [
        "Fecha",
        "Tipo",
        "Cuenta origen CRV",
        "Cuenta destino / beneficiario",
        "Concepto",
        "Tercero / beneficiario",
        "Documento",
        "Soporte",
        "Referencia",
        "Valor",
        "Conciliado"
    ]

    _excel_titulo(
        mov_ws,
        "MOVIMIENTOS BANCARIOS Y DE CAJA",
        "Únicamente movimientos asociados a cuentas propias de CRV",
        len(mov_headers)
    )

    _excel_encabezado(
        mov_ws,
        4,
        mov_headers
    )

    for fila, m in enumerate(
        movimientos,
        5
    ):
        valores = [
            m.fecha,
            m.tipo,
            m.cuenta,
            m.cuenta_destino,
            m.concepto,
            (
                m.nombre_tercero
                or m.beneficiario
                or ""
            ),
            m.numero_documento,
            m.numero_soporte,
            m.referencia_pago,
            float(
                m.valor_neto
                or m.valor
                or 0
            ),
            (
                "SI"
                if m.revisado
                else "NO"
            )
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            mov_ws.cell(
                fila,
                col,
                valor
            )

        _excel_moneda(
            mov_ws.cell(
                fila,
                10
            )
        )

    mov_ws.freeze_panes = "A5"
    mov_ws.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(mov_headers))}"
        f"{max(4, mov_ws.max_row)}"
    )

    _excel_auto_ancho(mov_ws)

    return _enviar_excel(
        wb,
        (
            f"CRV_bancos_y_caja_"
            f"{date.today().isoformat()}.xlsx"
        )
    )


# ============================================================
# BANCOS Y CAJA - PDF
# ============================================================

@crv_bp.route("/bancos-caja/pdf")
def bancos_caja_pdf():

    bloqueo = proteger_crv()

    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    cuentas = _cuentas_financieras_crv(
        empresa_id
    )

    movimientos = (
        CRVMovimiento.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )

    movimientos = [
        m
        for m in movimientos
        if _movimiento_es_bancario(m)
    ]

    story = []

    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - BANCOS Y CAJA",
        (
            "Reporte para conciliación bancaria · "
            f"Generado {date.today().strftime('%d/%m/%Y')}"
        )
    )

    filas_saldos = []

    for c in cuentas:
        filas_saldos.append([
            c["nombre"],
            _money_export(c["saldo"]),
            _money_export(c["ingresos"]),
            _money_export(c["egresos"]),
            c["movimientos"],
            c["sin_conciliar"]
        ])

    story.append(
        Paragraph(
            "Resumen por cuenta",
            _pdf_estilos_crv()[2]
        )
    )

    story.append(
        _pdf_tabla(
            [
                "Cuenta CRV",
                "Saldo",
                "Entradas",
                "Egresos",
                "Movimientos",
                "Sin conciliar"
            ],
            filas_saldos,
            [
                65 * mm,
                34 * mm,
                34 * mm,
                34 * mm,
                27 * mm,
                27 * mm
            ]
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )

    story.append(
        Paragraph(
            "Movimientos bancarios",
            _pdf_estilos_crv()[2]
        )
    )

    filas_mov = []

    for m in movimientos:
        filas_mov.append([
            (
                m.fecha.strftime("%d/%m/%Y")
                if m.fecha
                else ""
            ),
            m.tipo or "",
            m.cuenta or "",
            (
                m.cuenta_destino
                if str(
                    m.tipo or ""
                ).strip().upper() == "TRANSFERENCIA"
                else (
                    m.nombre_tercero
                    or m.beneficiario
                    or ""
                )
            ),
            m.concepto or "",
            _money_export(
                m.valor_neto
                or m.valor
                or 0
            ),
            (
                "CONCILIADO"
                if m.revisado
                else "PENDIENTE"
            )
        ])

    story.append(
        _pdf_tabla(
            [
                "Fecha",
                "Tipo",
                "Cuenta origen",
                "Destino / beneficiario",
                "Concepto",
                "Valor",
                "Conciliación"
            ],
            filas_mov,
            [
                20 * mm,
                30 * mm,
                38 * mm,
                48 * mm,
                65 * mm,
                28 * mm,
                28 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            f"CRV_bancos_y_caja_"
            f"{date.today().isoformat()}.pdf"
        )
    )

# ============================================================
# CUENTAS POR COBRAR - AUXILIARES
# ============================================================

def _clientes_crv(empresa_id):
    return CRVTercero.query.filter(
        CRVTercero.empresa_id == empresa_id,
        CRVTercero.activo == True,
        CRVTercero.tipo_tercero.in_([
            "CLIENTE",
            "TERCERO",
            "ENTIDAD",
            "OTRO",
        ])
    ).order_by(CRVTercero.razon_social.asc()).all()


def _sincronizar_movimiento_cxc(cuenta):
    movimiento = CRVMovimiento.query.filter_by(
        empresa_id=cuenta.empresa_id,
        origen_registro=f"CUENTA_POR_COBRAR:{cuenta.id}"
    ).first()

    cliente = cuenta.cliente

    if movimiento is None:
        movimiento = CRVMovimiento(
            empresa_id=cuenta.empresa_id,
            origen_registro=f"CUENTA_POR_COBRAR:{cuenta.id}",
            usuario_registro=str(session.get("usuario", "")),
            estado="ACTIVO"
        )
        db.session.add(movimiento)

    movimiento.sede_id = cuenta.sede_id
    movimiento.fecha = cuenta.fecha_documento
    movimiento.tipo = "INGRESO"
    movimiento.categoria = cuenta.categoria or "VENTA / CUENTA POR COBRAR"
    movimiento.subcategoria = cuenta.subcategoria or ""
    movimiento.concepto = cuenta.concepto
    movimiento.descripcion = cuenta.descripcion or ""
    movimiento.tercero_id = cliente.id if cliente else None
    movimiento.tercero_nombre = cliente.razon_social if cliente else ""
    movimiento.tercero_identificacion = cliente.identificacion if cliente else ""
    movimiento.tercero_direccion = cliente.direccion if cliente else ""
    movimiento.tercero_email = cliente.email if cliente else ""
    movimiento.tercero_telefono = ((cliente.celular or cliente.telefono or "") if cliente else "")
    movimiento.tercero_contacto = cliente.contacto if cliente else ""
    movimiento.tipo_documento = cuenta.tipo_documento or "CUENTA POR COBRAR"
    movimiento.numero_documento = cuenta.numero_documento
    movimiento.numero_soporte = f"CXC-{cuenta.id}"
    movimiento.fecha_documento = cuenta.fecha_documento
    movimiento.fecha_vencimiento = cuenta.fecha_vencimiento
    movimiento.base_gravable = float(cuenta.base_gravable or 0)
    movimiento.tipo_iva = cuenta.tipo_iva or "NO_APLICA"
    movimiento.tarifa_iva = float(cuenta.tarifa_iva or 0)
    movimiento.valor_iva = float(cuenta.valor_iva or 0)
    movimiento.valor = float(cuenta.valor_bruto or 0)
    movimiento.descuento = float(cuenta.descuento or 0)
    movimiento.retencion = float(cuenta.retencion or 0)
    movimiento.otros_descuentos = float(cuenta.otros_descuentos or 0)
    movimiento.valor_neto = float(cuenta.valor_neto or 0)
    movimiento.forma_pago = "CREDITO"
    movimiento.cuenta = "CUENTAS POR COBRAR"
    movimiento.cuenta_destino = ""
    movimiento.beneficiario = "CRV CONSTRUCTORA"
    movimiento.referencia_pago = ""
    movimiento.proyecto = cuenta.proyecto or ""
    movimiento.centro_costo = cuenta.centro_costo or ""
    movimiento.obra = cuenta.obra or ""
    movimiento.observaciones = cuenta.observaciones or ""
    movimiento.revisado = False

    _asignar_comprobante_permanente(
        movimiento
    )

    return movimiento


# ============================================================
# CUENTAS POR COBRAR - LISTADO / CREAR
# ============================================================

@crv_bp.route("/cuentas-por-cobrar", methods=["GET", "POST"])
def cuentas_por_cobrar():
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get("empresa_nombre", "CRV CONSTRUCTORA")

    if request.method == "POST":
        cliente_id = request.form.get("cliente_id", "").strip()
        numero_documento = request.form.get("numero_documento", "").strip()
        fecha_documento = convertir_fecha(request.form.get("fecha_documento"))

        if not cliente_id.isdigit() or not numero_documento or not fecha_documento:
            flash("Cliente, documento y fecha son obligatorios.", "error")
            return redirect("/crv/cuentas-por-cobrar")

        cliente = CRVTercero.query.filter_by(
            id=int(cliente_id),
            empresa_id=empresa_id
        ).first()

        if cliente is None:
            flash("El cliente seleccionado no pertenece a CRV.", "error")
            return redirect("/crv/cuentas-por-cobrar")

        existe = CRVCuentaCobrar.query.filter_by(
            empresa_id=empresa_id,
            numero_documento=numero_documento
        ).first()

        if existe:
            flash("Ya existe una cuenta por cobrar con ese documento.", "error")
            return redirect("/crv/cuentas-por-cobrar")

        sede_texto = request.form.get("sede_id", "").strip()
        sede_id = int(sede_texto) if sede_texto.isdigit() else None

        iva_datos = calcular_iva_formulario(
            request.form
        )

        cuenta = CRVCuentaCobrar(
            empresa_id=empresa_id,
            cliente_id=cliente.id,
            sede_id=sede_id,
            tipo_documento=request.form.get("tipo_documento", "FACTURA").strip(),
            numero_documento=numero_documento,
            fecha_documento=fecha_documento,
            fecha_vencimiento=convertir_fecha(request.form.get("fecha_vencimiento")),
            concepto=request.form.get("concepto", "").strip() or "Cuenta por cobrar",
            descripcion=request.form.get("descripcion", "").strip(),
            base_gravable=iva_datos["base_gravable"],
            tipo_iva=iva_datos["tipo_iva"],
            tarifa_iva=iva_datos["tarifa_iva"],
            valor_iva=iva_datos["valor_iva"],
            valor_bruto=iva_datos["total_documento"],
            descuento=convertir_numero(request.form.get("descuento")),
            retencion=convertir_numero(request.form.get("retencion")),
            otros_descuentos=convertir_numero(request.form.get("otros_descuentos")),
            categoria=request.form.get("categoria", "VENTA / CUENTA POR COBRAR").strip(),
            subcategoria=request.form.get("subcategoria", "").strip(),
            proyecto=request.form.get("proyecto", "").strip(),
            centro_costo=request.form.get("centro_costo", "").strip(),
            obra=request.form.get("obra", "").strip(),
            observaciones=request.form.get("observaciones", "").strip(),
            origen_registro="SISTEMA",
            usuario_registro=str(session.get("usuario", ""))
        )

        if cuenta.valor_bruto <= 0:
            flash("El valor bruto debe ser mayor a cero.", "error")
            return redirect("/crv/cuentas-por-cobrar")

        try:
            db.session.add(cuenta)
            db.session.flush()
            cuenta.recalcular_saldos()
            _sincronizar_movimiento_cxc(cuenta)
            db.session.commit()
            flash("Cuenta por cobrar registrada correctamente.", "success")
            return redirect(f"/crv/cuentas-por-cobrar/{cuenta.id}")
        except Exception as error:
            db.session.rollback()
            flash(f"No fue posible guardar la cuenta por cobrar: {error}", "error")
            return redirect("/crv/cuentas-por-cobrar")

    mes = request.args.get("mes", "").strip()
    cliente_filtro = request.args.get("cliente_id", "").strip()
    estado = request.args.get("estado", "").strip().upper()
    buscar = request.args.get("buscar", "").strip()

    query = CRVCuentaCobrar.query.filter_by(empresa_id=empresa_id)

    if mes:
        try:
            anio, numero_mes = mes.split("-")
            query = query.filter(
                db.extract("year", CRVCuentaCobrar.fecha_documento) == int(anio),
                db.extract("month", CRVCuentaCobrar.fecha_documento) == int(numero_mes)
            )
        except ValueError:
            pass

    if cliente_filtro.isdigit():
        query = query.filter_by(cliente_id=int(cliente_filtro))

    cuentas = query.order_by(
        CRVCuentaCobrar.fecha_documento.desc(),
        CRVCuentaCobrar.id.desc()
    ).all()

    for cuenta in cuentas:
        cuenta.recalcular_saldos()

    if estado:
        cuentas = [c for c in cuentas if c.estado == estado]

    if buscar:
        buscar_low = buscar.lower()
        cuentas = [
            c for c in cuentas
            if buscar_low in (c.numero_documento or "").lower()
            or buscar_low in (c.concepto or "").lower()
            or buscar_low in ((c.cliente.razon_social if c.cliente else "").lower())
            or buscar_low in ((c.cliente.identificacion if c.cliente else "").lower())
        ]

    total_facturado = sum(float(c.valor_neto or 0) for c in cuentas)
    total_cobrado = sum(float(c.total_cobrado or 0) for c in cuentas)
    total_pendiente = sum(float(c.saldo_pendiente or 0) for c in cuentas)
    total_vencido = sum(
        float(c.saldo_pendiente or 0)
        for c in cuentas
        if c.estado == "VENCIDA"
    )

    db.session.commit()

    return render_template(
        "contabilidad/crv_cuentas_cobrar.html",
        empresa_nombre=empresa_nombre,
        cuentas=cuentas,
        clientes=_clientes_crv(empresa_id),
        sedes=Sede.query.filter_by(empresa_id=empresa_id).order_by(Sede.nombre.asc()).all(),
        mes_actual=mes,
        cliente_actual=cliente_filtro,
        estado_actual=estado,
        buscar_actual=buscar,
        total_facturado=total_facturado,
        total_cobrado=total_cobrado,
        total_pendiente=total_pendiente,
        total_vencido=total_vencido
    )


# ============================================================
# CUENTAS POR COBRAR - DETALLE
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/<int:cuenta_id>")
def cuenta_cobrar_detalle(cuenta_id):
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaCobrar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-cobrar")

    cuenta.recalcular_saldos()
    db.session.commit()

    return render_template(
        "contabilidad/crv_cuenta_cobrar_detalle.html",
        empresa_nombre=session.get("empresa_nombre", "CRV CONSTRUCTORA"),
        cuenta=cuenta,
        cobros=sorted(cuenta.cobros, key=lambda x: (x.fecha, x.id), reverse=True),
        cuentas_financieras=_cuentas_financieras_crv(empresa_id)
    )


# ============================================================
# CUENTAS POR COBRAR - EDITAR
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/<int:cuenta_id>/editar", methods=["POST"])
def cuenta_cobrar_editar(cuenta_id):
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaCobrar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-cobrar")

    cliente_id = request.form.get("cliente_id", "").strip()
    cliente = CRVTercero.query.filter_by(
        id=int(cliente_id) if cliente_id.isdigit() else -1,
        empresa_id=empresa_id
    ).first()

    if cliente is None:
        flash("Cliente inválido.", "error")
        return redirect("/crv/cuentas-por-cobrar")

    iva_datos = calcular_iva_formulario(request.form)
    nuevo_bruto = iva_datos["total_documento"]
    descuento = convertir_numero(request.form.get("descuento"))
    retencion = convertir_numero(request.form.get("retencion"))
    otros = convertir_numero(request.form.get("otros_descuentos"))
    nuevo_neto = max(0, nuevo_bruto - descuento - retencion - otros)

    if nuevo_neto < float(cuenta.total_cobrado or 0):
        flash("No puedes dejar el valor neto por debajo de lo ya cobrado.", "error")
        return redirect("/crv/cuentas-por-cobrar")

    sede_texto = request.form.get("sede_id", "").strip()

    cuenta.cliente_id = cliente.id
    cuenta.sede_id = int(sede_texto) if sede_texto.isdigit() else None
    cuenta.tipo_documento = request.form.get("tipo_documento", "").strip()
    cuenta.numero_documento = request.form.get("numero_documento", "").strip()
    cuenta.fecha_documento = convertir_fecha(request.form.get("fecha_documento")) or cuenta.fecha_documento
    cuenta.fecha_vencimiento = convertir_fecha(request.form.get("fecha_vencimiento"))
    cuenta.concepto = request.form.get("concepto", "").strip()
    cuenta.descripcion = request.form.get("descripcion", "").strip()
    cuenta.base_gravable = iva_datos["base_gravable"]
    cuenta.tipo_iva = iva_datos["tipo_iva"]
    cuenta.tarifa_iva = iva_datos["tarifa_iva"]
    cuenta.valor_iva = iva_datos["valor_iva"]
    cuenta.valor_bruto = nuevo_bruto
    cuenta.descuento = descuento
    cuenta.retencion = retencion
    cuenta.otros_descuentos = otros
    cuenta.categoria = request.form.get("categoria", "").strip()
    cuenta.subcategoria = request.form.get("subcategoria", "").strip()
    cuenta.proyecto = request.form.get("proyecto", "").strip()
    cuenta.centro_costo = request.form.get("centro_costo", "").strip()
    cuenta.obra = request.form.get("obra", "").strip()
    cuenta.observaciones = request.form.get("observaciones", "").strip()
    cuenta.recalcular_saldos()
    _sincronizar_movimiento_cxc(cuenta)

    try:
        db.session.commit()
        flash("Cuenta por cobrar actualizada correctamente.", "success")
    except Exception as error:
        db.session.rollback()
        flash(f"No fue posible actualizar: {error}", "error")

    return redirect("/crv/cuentas-por-cobrar")


# ============================================================
# CUENTAS POR COBRAR - ELIMINAR
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/<int:cuenta_id>/eliminar", methods=["POST"])
def cuenta_cobrar_eliminar(cuenta_id):
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaCobrar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-cobrar")

    if cuenta.cobros:
        flash("No se puede eliminar una cuenta con cobros. Primero revierte los cobros.", "warning")
        return redirect("/crv/cuentas-por-cobrar")

    movimiento = CRVMovimiento.query.filter_by(
        empresa_id=empresa_id,
        origen_registro=f"CUENTA_POR_COBRAR:{cuenta.id}"
    ).first()

    try:
        if movimiento:
            db.session.delete(movimiento)
        db.session.delete(cuenta)
        db.session.commit()
        flash("Cuenta por cobrar eliminada correctamente.", "success")
    except Exception as error:
        db.session.rollback()
        flash(f"No fue posible eliminar la cuenta: {error}", "error")

    return redirect("/crv/cuentas-por-cobrar")


# ============================================================
# CUENTAS POR COBRAR - REGISTRAR COBRO
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/<int:cuenta_id>/cobrar", methods=["POST"])
def cuenta_cobrar_registrar_cobro(cuenta_id):
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaCobrar.query.get_or_404(cuenta_id)

    if cuenta.empresa_id != empresa_id:
        return redirect("/crv/cuentas-por-cobrar")

    cuenta.recalcular_saldos()

    fecha = convertir_fecha(request.form.get("fecha"))
    valor = convertir_numero(request.form.get("valor_cobrado"))
    cuenta_receptora = request.form.get("cuenta_receptora", "").strip()

    if not fecha or valor <= 0:
        flash("Fecha y valor del cobro son obligatorios.", "error")
        return redirect(f"/crv/cuentas-por-cobrar/{cuenta.id}")

    if valor > float(cuenta.saldo_pendiente or 0):
        flash("El cobro no puede superar el saldo pendiente.", "error")
        return redirect(f"/crv/cuentas-por-cobrar/{cuenta.id}")

    if not cuenta_receptora:
        flash("Selecciona la cuenta propia de CRV que recibió el dinero.", "error")
        return redirect(f"/crv/cuentas-por-cobrar/{cuenta.id}")

    cliente = cuenta.cliente

    cobro = CRVCobroCuentaCobrar(
        empresa_id=empresa_id,
        cuenta_cobrar_id=cuenta.id,
        fecha=fecha,
        valor_cobrado=valor,
        forma_pago=request.form.get("forma_pago", "").strip(),
        cuenta_receptora=cuenta_receptora,
        cuenta_origen_cliente=request.form.get("cuenta_origen_cliente", "").strip(),
        pagador=request.form.get("pagador", "").strip() or (cliente.razon_social if cliente else ""),
        referencia_pago=request.form.get("referencia_pago", "").strip(),
        numero_soporte=request.form.get("numero_soporte", "").strip(),
        observaciones=request.form.get("observaciones", "").strip(),
        usuario_registro=str(session.get("usuario", ""))
    )

    try:
        db.session.add(cobro)
        db.session.flush()

        movimiento = CRVMovimiento(
            empresa_id=empresa_id,
            sede_id=cuenta.sede_id,
            fecha=fecha,
            tipo="COBRO CUENTA POR COBRAR",
            categoria="COBRO CLIENTE",
            subcategoria=cuenta.categoria or "",
            concepto=f"Cobro cuenta por cobrar {cuenta.numero_documento}",
            descripcion=f"Cobro recibido de {cliente.razon_social if cliente else 'cliente'} - documento {cuenta.numero_documento}",
            tercero_id=cliente.id if cliente else None,
            tercero_nombre=cliente.razon_social if cliente else "",
            tercero_identificacion=cliente.identificacion if cliente else "",
            tercero_direccion=cliente.direccion if cliente else "",
            tercero_email=cliente.email if cliente else "",
            tercero_telefono=((cliente.celular or cliente.telefono or "") if cliente else ""),
            tercero_contacto=cliente.contacto if cliente else "",
            tipo_documento="COBRO CXC",
            numero_documento=cuenta.numero_documento,
            numero_soporte=cobro.numero_soporte or f"CXC-COBRO-{cobro.id}",
            fecha_documento=fecha,
            fecha_vencimiento=None,
            valor=valor,
            descuento=0,
            retencion=0,
            otros_descuentos=0,
            valor_neto=valor,
            forma_pago=cobro.forma_pago or "",
            cuenta=cuenta_receptora,
            cuenta_destino="",
            beneficiario="CRV CONSTRUCTORA",
            referencia_pago=cobro.referencia_pago or "",
            proyecto=cuenta.proyecto or "",
            centro_costo=cuenta.centro_costo or "",
            obra=cuenta.obra or "",
            observaciones=cobro.observaciones or "Cobro automático generado desde cuentas por cobrar.",
            revisado=False,
            estado="ACTIVO",
            origen_registro="COBRO_CUENTA_POR_COBRAR",
            usuario_registro=str(session.get("usuario", ""))
        )

        _asignar_comprobante_permanente(
            movimiento
        )

        db.session.add(movimiento)
        db.session.flush()
        cobro.movimiento_id = movimiento.id
        cuenta.recalcular_saldos()
        db.session.commit()
        flash("Cobro registrado correctamente.", "success")
    except Exception as error:
        db.session.rollback()
        flash(f"No fue posible registrar el cobro: {error}", "error")

    return redirect(f"/crv/cuentas-por-cobrar/{cuenta.id}")


# ============================================================
# CUENTAS POR COBRAR - REVERSAR COBRO
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/<int:cuenta_id>/cobros/<int:cobro_id>/eliminar", methods=["POST"])
def cuenta_cobrar_reversar_cobro(cuenta_id, cobro_id):
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuenta = CRVCuentaCobrar.query.get_or_404(cuenta_id)
    cobro = CRVCobroCuentaCobrar.query.get_or_404(cobro_id)

    if cuenta.empresa_id != empresa_id or cobro.empresa_id != empresa_id or cobro.cuenta_cobrar_id != cuenta.id:
        return redirect("/crv/cuentas-por-cobrar")

    try:
        if cobro.movimiento:
            db.session.delete(cobro.movimiento)
        db.session.delete(cobro)
        db.session.flush()
        cuenta.recalcular_saldos()
        db.session.commit()
        flash("Cobro reversado correctamente.", "success")
    except Exception as error:
        db.session.rollback()
        flash(f"No fue posible reversar el cobro: {error}", "error")

    return redirect(f"/crv/cuentas-por-cobrar/{cuenta.id}")


# ============================================================
# CUENTAS POR COBRAR - EXCEL
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/excel")
def cuentas_cobrar_excel():
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuentas = CRVCuentaCobrar.query.filter_by(empresa_id=empresa_id).order_by(
        CRVCuentaCobrar.fecha_documento.asc(), CRVCuentaCobrar.id.asc()
    ).all()

    for c in cuentas:
        c.recalcular_saldos()

    wb = Workbook()
    ws = wb.active
    ws.title = "CUENTAS_POR_COBRAR"

    headers = [
        "Fecha", "Cliente", "NIT/ID", "Tipo documento", "Documento",
        "Concepto", "Proyecto", "Centro costo", "Vencimiento",
        "Base gravable", "Tipo IVA", "Tarifa IVA %", "Valor IVA", "Total documento",
        "Descuento", "Retención", "Otros descuentos",
        "Neto", "Cobrado", "Saldo", "Estado"
    ]

    _excel_titulo(ws, "CRV CONSTRUCTORA - CUENTAS POR COBRAR",
                  f"Cartera de clientes · Generado {date.today().strftime('%d/%m/%Y')}", len(headers))
    _excel_encabezado(ws, 4, headers)

    for fila, c in enumerate(cuentas, 5):
        cli = c.cliente
        valores = [
            c.fecha_documento,
            cli.razon_social if cli else "",
            cli.identificacion if cli else "",
            c.tipo_documento,
            c.numero_documento,
            c.concepto,
            c.proyecto,
            c.centro_costo,
            c.fecha_vencimiento,
            float(c.base_gravable or 0),
            c.tipo_iva or "NO_APLICA",
            float(c.tarifa_iva or 0),
            float(c.valor_iva or 0),
            float(c.valor_bruto or 0),
            float(c.descuento or 0),
            float(c.retencion or 0),
            float(c.otros_descuentos or 0),
            float(c.valor_neto or 0),
            float(c.total_cobrado or 0),
            float(c.saldo_pendiente or 0),
            c.estado,
        ]
        for col, valor in enumerate(valores, 1):
            ws.cell(fila, col, valor)
        for col in (10, 13, 14, 15, 16, 17, 18, 19, 20):
            _excel_moneda(ws.cell(fila, col))

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, ws.max_row)}"
    _excel_auto_ancho(ws)

    return _enviar_excel(wb, f"CRV_cuentas_por_cobrar_{date.today().isoformat()}.xlsx")


# ============================================================
# CUENTAS POR COBRAR - PDF
# ============================================================

@crv_bp.route("/cuentas-por-cobrar/pdf")
def cuentas_cobrar_pdf():
    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    cuentas = CRVCuentaCobrar.query.filter_by(empresa_id=empresa_id).order_by(
        CRVCuentaCobrar.fecha_documento.asc(), CRVCuentaCobrar.id.asc()
    ).all()

    filas = []
    for c in cuentas:
        c.recalcular_saldos()
        cli = c.cliente
        filas.append([
            c.fecha_documento.strftime("%d/%m/%Y") if c.fecha_documento else "",
            cli.razon_social if cli else "",
            c.numero_documento or "",
            c.concepto or "",
            c.fecha_vencimiento.strftime("%d/%m/%Y") if c.fecha_vencimiento else "",
            _money_export(c.base_gravable),
            _money_export(c.valor_iva),
            _money_export(c.valor_bruto),
            _money_export(c.valor_neto),
            _money_export(c.total_cobrado),
            _money_export(c.saldo_pendiente),
            c.estado or ""
        ])

    story = []
    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - CUENTAS POR COBRAR",
        f"Cartera general de clientes · Generado {date.today().strftime('%d/%m/%Y')}"
    )
    story.append(_pdf_tabla(
        ["Fecha", "Cliente", "Documento", "Concepto", "Vence", "Base", "IVA", "Total", "Neto", "Cobrado", "Saldo", "Estado"],
        filas,
        [18*mm, 38*mm, 22*mm, 38*mm, 18*mm, 21*mm, 19*mm, 21*mm, 21*mm, 21*mm, 21*mm, 20*mm]
    ))

    return _enviar_pdf(story, f"CRV_cuentas_por_cobrar_{date.today().isoformat()}.pdf")

# ============================================================
# PROYECTOS / OBRAS - FUNCIONES AUXILIARES
# ============================================================

def _nombre_proyecto_normalizado(valor):
    return str(valor or "").strip()


def _movimientos_proyecto(empresa_id, nombre_proyecto):
    nombre = _nombre_proyecto_normalizado(nombre_proyecto)

    return (
        CRVMovimiento.query
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            db.func.lower(
                db.func.trim(
                    CRVMovimiento.proyecto
                )
            ) == nombre.lower()
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )


def _cuentas_pagar_proyecto(empresa_id, nombre_proyecto):
    nombre = _nombre_proyecto_normalizado(nombre_proyecto)

    return (
        CRVCuentaPagar.query
        .filter(
            CRVCuentaPagar.empresa_id == empresa_id,
            db.func.lower(
                db.func.trim(
                    CRVCuentaPagar.proyecto
                )
            ) == nombre.lower()
        )
        .order_by(
            CRVCuentaPagar.fecha_documento.asc(),
            CRVCuentaPagar.id.asc()
        )
        .all()
    )


def _cuentas_cobrar_proyecto(empresa_id, nombre_proyecto):
    nombre = _nombre_proyecto_normalizado(nombre_proyecto)

    return (
        CRVCuentaCobrar.query
        .filter(
            CRVCuentaCobrar.empresa_id == empresa_id,
            db.func.lower(
                db.func.trim(
                    CRVCuentaCobrar.proyecto
                )
            ) == nombre.lower()
        )
        .order_by(
            CRVCuentaCobrar.fecha_documento.asc(),
            CRVCuentaCobrar.id.asc()
        )
        .all()
    )


def _resumen_proyecto(proyecto):
    empresa_id = proyecto.empresa_id
    movimientos = _movimientos_proyecto(
        empresa_id,
        proyecto.nombre
    )

    ingresos = 0.0
    compras = 0.0
    gastos = 0.0

    tipos_gasto = {
        "GASTO",
        "NOMINA AFILIADOS",
        "NÓMINA AFILIADOS",
        "NOMINA OBRA",
        "NÓMINA OBRA",
        "ANTICIPO",
        "PRESTAMO EMPLEADO",
        "PRÉSTAMO EMPLEADO",
        "EGRESO",
        "OTRO EGRESO",
        "OTRO",
    }

    tipos_sin_resultado = {
        "TRANSFERENCIA",
        "PAGO CUENTA POR PAGAR",
        "COBRO CUENTA POR COBRAR",
        "SALDO INICIAL CUENTA",
    }

    for m in movimientos:
        tipo = str(
            m.tipo or ""
        ).strip().upper()

        valor = float(
            m.valor_neto
            or m.valor
            or 0
        )

        if tipo == "INGRESO":
            ingresos += valor

        elif tipo == "COMPRA":
            compras += valor

        elif tipo in tipos_gasto:
            gastos += valor

        elif tipo in tipos_sin_resultado:
            pass

        else:
            # Conservamos el mismo criterio del dashboard:
            # un egreso no clasificado se considera gasto.
            gastos += valor

    cuentas_pagar = _cuentas_pagar_proyecto(
        empresa_id,
        proyecto.nombre
    )

    for cuenta in cuentas_pagar:
        cuenta.recalcular_saldos()

    cuentas_cobrar = _cuentas_cobrar_proyecto(
        empresa_id,
        proyecto.nombre
    )

    for cuenta in cuentas_cobrar:
        cuenta.recalcular_saldos()

    cxp_facturado = sum(
        float(c.valor_neto or 0)
        for c in cuentas_pagar
    )

    cxp_pagado = sum(
        float(c.total_pagado or 0)
        for c in cuentas_pagar
    )

    cxp_pendiente = sum(
        float(c.saldo_pendiente or 0)
        for c in cuentas_pagar
    )

    cxc_facturado = sum(
        float(c.valor_neto or 0)
        for c in cuentas_cobrar
    )

    cxc_cobrado = sum(
        float(c.total_cobrado or 0)
        for c in cuentas_cobrar
    )

    cxc_pendiente = sum(
        float(c.saldo_pendiente or 0)
        for c in cuentas_cobrar
    )

    costos = compras + gastos
    resultado = ingresos - costos

    presupuesto = float(
        proyecto.presupuesto or 0
    )

    porcentaje_presupuesto = (
        (costos / presupuesto) * 100
        if presupuesto > 0
        else 0
    )

    disponible_presupuesto = (
        presupuesto - costos
        if presupuesto > 0
        else 0
    )

    return {
        "movimientos": movimientos,
        "cuentas_pagar": cuentas_pagar,
        "cuentas_cobrar": cuentas_cobrar,
        "ingresos": ingresos,
        "compras": compras,
        "gastos": gastos,
        "costos": costos,
        "resultado": resultado,
        "cxp_facturado": cxp_facturado,
        "cxp_pagado": cxp_pagado,
        "cxp_pendiente": cxp_pendiente,
        "cxc_facturado": cxc_facturado,
        "cxc_cobrado": cxc_cobrado,
        "cxc_pendiente": cxc_pendiente,
        "presupuesto": presupuesto,
        "porcentaje_presupuesto": porcentaje_presupuesto,
        "disponible_presupuesto": disponible_presupuesto,
        "cantidad_movimientos": len(movimientos),
    }


def _clientes_proyectos_crv(empresa_id):
    return (
        CRVTercero.query
        .filter(
            CRVTercero.empresa_id == empresa_id,
            CRVTercero.activo.is_(True),
            CRVTercero.tipo_tercero.in_(
                ["CLIENTE", "TERCERO"]
            )
        )
        .order_by(
            CRVTercero.razon_social.asc()
        )
        .all()
    )


# ============================================================
# PROYECTOS / OBRAS - LISTADO Y CREACIÓN
# ============================================================

@crv_bp.route(
    "/proyectos",
    methods=["GET", "POST"]
)
def proyectos_crv():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")
    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    if request.method == "POST":

        codigo = request.form.get(
            "codigo",
            ""
        ).strip().upper()

        nombre = request.form.get(
            "nombre",
            ""
        ).strip().upper()

        if not codigo or not nombre:
            flash(
                "Código y nombre del proyecto son obligatorios.",
                "error"
            )
            return redirect("/crv/proyectos")

        existe_codigo = (
            CRVProyecto.query
            .filter_by(
                empresa_id=empresa_id,
                codigo=codigo
            )
            .first()
        )

        if existe_codigo:
            flash(
                "Ya existe un proyecto con ese código.",
                "error"
            )
            return redirect("/crv/proyectos")

        existe_nombre = (
            CRVProyecto.query
            .filter(
                CRVProyecto.empresa_id == empresa_id,
                db.func.lower(
                    CRVProyecto.nombre
                ) == nombre.lower()
            )
            .first()
        )

        if existe_nombre:
            flash(
                "Ya existe un proyecto con ese nombre.",
                "error"
            )
            return redirect("/crv/proyectos")

        cliente_texto = request.form.get(
            "cliente_id",
            ""
        ).strip()

        cliente_id = (
            int(cliente_texto)
            if cliente_texto.isdigit()
            else None
        )

        proyecto = CRVProyecto(
            empresa_id=empresa_id,
            codigo=codigo,
            nombre=nombre,
            tipo=request.form.get(
                "tipo",
                ""
            ).strip().upper(),
            cliente_id=cliente_id,
            ubicacion=request.form.get(
                "ubicacion",
                ""
            ).strip(),
            fecha_inicio=convertir_fecha(
                request.form.get(
                    "fecha_inicio"
                )
            ),
            fecha_fin_estimada=convertir_fecha(
                request.form.get(
                    "fecha_fin_estimada"
                )
            ),
            presupuesto=convertir_numero(
                request.form.get(
                    "presupuesto"
                )
            ),
            valor_contrato=convertir_numero(
                request.form.get(
                    "valor_contrato"
                )
            ),
            responsable=request.form.get(
                "responsable",
                ""
            ).strip(),
            centro_costo=request.form.get(
                "centro_costo",
                ""
            ).strip(),
            estado=request.form.get(
                "estado",
                "ACTIVO"
            ).strip().upper() or "ACTIVO",
            descripcion=request.form.get(
                "descripcion",
                ""
            ).strip(),
            observaciones=request.form.get(
                "observaciones",
                ""
            ).strip(),
            activo=True,
            usuario_registro=str(
                session.get(
                    "usuario",
                    ""
                )
            )
        )

        try:
            db.session.add(proyecto)
            db.session.commit()

            flash(
                "Proyecto registrado correctamente.",
                "success"
            )

            return redirect(
                f"/crv/proyectos/{proyecto.id}"
            )

        except Exception as error:
            db.session.rollback()

            flash(
                f"No fue posible registrar el proyecto: {error}",
                "error"
            )

            return redirect("/crv/proyectos")

    buscar = request.args.get(
        "buscar",
        ""
    ).strip()

    estado = request.args.get(
        "estado",
        ""
    ).strip().upper()

    query = CRVProyecto.query.filter_by(
        empresa_id=empresa_id
    )

    if estado:
        query = query.filter_by(
            estado=estado
        )

    proyectos = query.order_by(
        CRVProyecto.activo.desc(),
        CRVProyecto.nombre.asc()
    ).all()

    if buscar:
        texto = buscar.lower()

        proyectos = [
            p for p in proyectos
            if (
                texto in (p.codigo or "").lower()
                or texto in (p.nombre or "").lower()
                or texto in (p.ubicacion or "").lower()
                or texto in (p.responsable or "").lower()
                or texto in (
                    (
                        p.cliente.razon_social
                        if p.cliente
                        else ""
                    ).lower()
                )
            )
        ]

    filas = []

    total_presupuesto = 0.0
    total_costos = 0.0
    total_ingresos = 0.0
    total_resultado = 0.0

    for proyecto in proyectos:
        resumen = _resumen_proyecto(
            proyecto
        )

        filas.append({
            "proyecto": proyecto,
            "resumen": resumen
        })

        total_presupuesto += float(
            proyecto.presupuesto or 0
        )

        total_costos += resumen[
            "costos"
        ]

        total_ingresos += resumen[
            "ingresos"
        ]

        total_resultado += resumen[
            "resultado"
        ]

    db.session.commit()

    return render_template(
        "contabilidad/crv_proyectos.html",
        empresa_nombre=empresa_nombre,
        filas=filas,
        clientes=_clientes_proyectos_crv(
            empresa_id
        ),
        buscar_actual=buscar,
        estado_actual=estado,
        total_proyectos=len(proyectos),
        total_presupuesto=total_presupuesto,
        total_costos=total_costos,
        total_ingresos=total_ingresos,
        total_resultado=total_resultado
    )


# ============================================================
# PROYECTOS / OBRAS - DETALLE
# ============================================================

@crv_bp.route(
    "/proyectos/<int:proyecto_id>"
)
def proyecto_detalle_crv(proyecto_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    proyecto = (
        CRVProyecto.query
        .get_or_404(
            proyecto_id
        )
    )

    if proyecto.empresa_id != empresa_id:
        return redirect(
            "/crv/proyectos"
        )

    resumen = _resumen_proyecto(
        proyecto
    )

    db.session.commit()

    return render_template(
        "contabilidad/crv_proyecto_detalle.html",
        empresa_nombre=session.get(
            "empresa_nombre",
            "CRV CONSTRUCTORA"
        ),
        proyecto=proyecto,
        resumen=resumen
    )


# ============================================================
# PROYECTOS / OBRAS - EDITAR
# ============================================================

@crv_bp.route(
    "/proyectos/<int:proyecto_id>/editar",
    methods=["POST"]
)
def proyecto_editar_crv(proyecto_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    proyecto = (
        CRVProyecto.query
        .get_or_404(
            proyecto_id
        )
    )

    if proyecto.empresa_id != empresa_id:
        return redirect(
            "/crv/proyectos"
        )

    nombre_anterior = proyecto.nombre

    codigo = request.form.get(
        "codigo",
        ""
    ).strip().upper()

    nombre = request.form.get(
        "nombre",
        ""
    ).strip().upper()

    if not codigo or not nombre:
        flash(
            "Código y nombre son obligatorios.",
            "error"
        )
        return redirect(
            f"/crv/proyectos/{proyecto.id}"
        )

    duplicado_codigo = (
        CRVProyecto.query
        .filter(
            CRVProyecto.empresa_id == empresa_id,
            CRVProyecto.codigo == codigo,
            CRVProyecto.id != proyecto.id
        )
        .first()
    )

    if duplicado_codigo:
        flash(
            "Ya existe otro proyecto con ese código.",
            "error"
        )
        return redirect(
            f"/crv/proyectos/{proyecto.id}"
        )

    cliente_texto = request.form.get(
        "cliente_id",
        ""
    ).strip()

    proyecto.codigo = codigo
    proyecto.nombre = nombre
    proyecto.tipo = request.form.get(
        "tipo",
        ""
    ).strip().upper()

    proyecto.cliente_id = (
        int(cliente_texto)
        if cliente_texto.isdigit()
        else None
    )

    proyecto.ubicacion = request.form.get(
        "ubicacion",
        ""
    ).strip()

    proyecto.fecha_inicio = convertir_fecha(
        request.form.get(
            "fecha_inicio"
        )
    )

    proyecto.fecha_fin_estimada = convertir_fecha(
        request.form.get(
            "fecha_fin_estimada"
        )
    )

    proyecto.presupuesto = convertir_numero(
        request.form.get(
            "presupuesto"
        )
    )

    proyecto.valor_contrato = convertir_numero(
        request.form.get(
            "valor_contrato"
        )
    )

    proyecto.responsable = request.form.get(
        "responsable",
        ""
    ).strip()

    proyecto.centro_costo = request.form.get(
        "centro_costo",
        ""
    ).strip()

    proyecto.estado = request.form.get(
        "estado",
        "ACTIVO"
    ).strip().upper() or "ACTIVO"

    proyecto.descripcion = request.form.get(
        "descripcion",
        ""
    ).strip()

    proyecto.observaciones = request.form.get(
        "observaciones",
        ""
    ).strip()

    # Si se cambia el nombre, trasladamos la clasificación histórica
    # de todos los registros vinculados al nombre anterior.
    if (
        nombre_anterior
        and nombre_anterior.strip().lower()
        != nombre.strip().lower()
    ):

        for movimiento in _movimientos_proyecto(
            empresa_id,
            nombre_anterior
        ):
            movimiento.proyecto = nombre

        for cuenta in _cuentas_pagar_proyecto(
            empresa_id,
            nombre_anterior
        ):
            cuenta.proyecto = nombre

        for cuenta in _cuentas_cobrar_proyecto(
            empresa_id,
            nombre_anterior
        ):
            cuenta.proyecto = nombre

    try:
        db.session.commit()

        flash(
            "Proyecto actualizado correctamente.",
            "success"
        )

    except Exception as error:
        db.session.rollback()

        flash(
            f"No fue posible actualizar el proyecto: {error}",
            "error"
        )

    return redirect(
        f"/crv/proyectos/{proyecto.id}"
    )


# ============================================================
# PROYECTOS / OBRAS - ARCHIVAR / REACTIVAR
# ============================================================

@crv_bp.route(
    "/proyectos/<int:proyecto_id>/estado",
    methods=["POST"]
)
def proyecto_estado_crv(proyecto_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    proyecto = (
        CRVProyecto.query
        .get_or_404(
            proyecto_id
        )
    )

    if proyecto.empresa_id != empresa_id:
        return redirect(
            "/crv/proyectos"
        )

    proyecto.activo = not bool(
        proyecto.activo
    )

    if not proyecto.activo:
        proyecto.estado = "ARCHIVADO"

    elif proyecto.estado == "ARCHIVADO":
        proyecto.estado = "ACTIVO"

    db.session.commit()

    flash(
        (
            "Proyecto reactivado correctamente."
            if proyecto.activo
            else "Proyecto archivado correctamente."
        ),
        "success"
    )

    return redirect(
        request.referrer
        or "/crv/proyectos"
    )


# ============================================================
# PROYECTOS / OBRAS - ELIMINAR SOLO SI NO TIENE HISTORIA
# ============================================================

@crv_bp.route(
    "/proyectos/<int:proyecto_id>/eliminar",
    methods=["POST"]
)
def proyecto_eliminar_crv(proyecto_id):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get("empresa_id")

    proyecto = (
        CRVProyecto.query
        .get_or_404(
            proyecto_id
        )
    )

    if proyecto.empresa_id != empresa_id:
        return redirect(
            "/crv/proyectos"
        )

    resumen = _resumen_proyecto(
        proyecto
    )

    if (
        resumen["movimientos"]
        or resumen["cuentas_pagar"]
        or resumen["cuentas_cobrar"]
    ):
        flash(
            "No se puede eliminar un proyecto con información contable. "
            "Puedes archivarlo para conservar la trazabilidad.",
            "warning"
        )

        return redirect(
            f"/crv/proyectos/{proyecto.id}"
        )

    try:
        db.session.delete(
            proyecto
        )

        db.session.commit()

        flash(
            "Proyecto eliminado correctamente.",
            "success"
        )

    except Exception as error:
        db.session.rollback()

        flash(
            f"No fue posible eliminar el proyecto: {error}",
            "error"
        )

    return redirect(
        "/crv/proyectos"
    )


# ============================================================
# PROYECTOS / OBRAS - EXCEL
# ============================================================

@crv_bp.route(
    "/proyectos/excel"
)
def proyectos_excel_crv():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    proyectos = (
        CRVProyecto.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVProyecto.nombre.asc()
        )
        .all()
    )

    wb = Workbook()

    ws = wb.active
    ws.title = "PROYECTOS"

    headers = [
        "Código",
        "Proyecto",
        "Estado",
        "Tipo",
        "Cliente",
        "Ubicación",
        "Responsable",
        "Centro costo",
        "Presupuesto",
        "Valor contrato",
        "Ingresos",
        "Compras",
        "Gastos",
        "Costos totales",
        "Resultado",
        "CxP pendiente",
        "CxC pendiente",
        "% presupuesto usado"
    ]

    _excel_titulo(
        ws,
        "CRV CONSTRUCTORA - PROYECTOS Y OBRAS",
        (
            "Control financiero por proyecto "
            f"· Generado {date.today().strftime('%d/%m/%Y')}"
        ),
        len(headers)
    )

    _excel_encabezado(
        ws,
        4,
        headers
    )

    for fila, proyecto in enumerate(
        proyectos,
        5
    ):
        resumen = _resumen_proyecto(
            proyecto
        )

        valores = [
            proyecto.codigo,
            proyecto.nombre,
            proyecto.estado,
            proyecto.tipo,
            (
                proyecto.cliente.razon_social
                if proyecto.cliente
                else ""
            ),
            proyecto.ubicacion,
            proyecto.responsable,
            proyecto.centro_costo,
            float(
                proyecto.presupuesto
                or 0
            ),
            float(
                proyecto.valor_contrato
                or 0
            ),
            resumen["ingresos"],
            resumen["compras"],
            resumen["gastos"],
            resumen["costos"],
            resumen["resultado"],
            resumen["cxp_pendiente"],
            resumen["cxc_pendiente"],
            resumen["porcentaje_presupuesto"] / 100
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            ws.cell(
                fila,
                col,
                valor
            )

        for col in range(
            9,
            18
        ):
            _excel_moneda(
                ws.cell(
                    fila,
                    col
                )
            )

        ws.cell(
            fila,
            18
        ).number_format = "0.00%"

    ws.freeze_panes = "A5"

    ws.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers))}"
        f"{max(4, ws.max_row)}"
    )

    _excel_auto_ancho(
        ws
    )

    detalle = wb.create_sheet(
        "MOVIMIENTOS_PROYECTO"
    )

    headers_detalle = [
        "Proyecto",
        "Fecha",
        "Tipo",
        "Tercero",
        "Documento",
        "Concepto",
        "Cuenta",
        "Centro costo",
        "Valor"
    ]

    _excel_titulo(
        detalle,
        "DETALLE DE MOVIMIENTOS POR PROYECTO",
        "Trazabilidad contable de proyectos y obras",
        len(headers_detalle)
    )

    _excel_encabezado(
        detalle,
        4,
        headers_detalle
    )

    fila = 5

    for proyecto in proyectos:

        for movimiento in _movimientos_proyecto(
            empresa_id,
            proyecto.nombre
        ):

            valores = [
                proyecto.nombre,
                movimiento.fecha,
                movimiento.tipo,
                movimiento.nombre_tercero,
                (
                    movimiento.numero_documento
                    or movimiento.numero_soporte
                ),
                movimiento.concepto,
                movimiento.cuenta,
                movimiento.centro_costo,
                float(
                    movimiento.valor_neto
                    or movimiento.valor
                    or 0
                )
            ]

            for col, valor in enumerate(
                valores,
                1
            ):
                detalle.cell(
                    fila,
                    col,
                    valor
                )

            _excel_moneda(
                detalle.cell(
                    fila,
                    9
                )
            )

            fila += 1

    detalle.freeze_panes = "A5"

    detalle.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers_detalle))}"
        f"{max(4, detalle.max_row)}"
    )

    _excel_auto_ancho(
        detalle
    )

    db.session.commit()

    return _enviar_excel(
        wb,
        (
            f"CRV_proyectos_obras_"
            f"{date.today().isoformat()}.xlsx"
        )
    )


# ============================================================
# PROYECTOS / OBRAS - PDF
# ============================================================

@crv_bp.route(
    "/proyectos/pdf"
)
def proyectos_pdf_crv():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    proyectos = (
        CRVProyecto.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVProyecto.nombre.asc()
        )
        .all()
    )

    story = []

    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - PROYECTOS Y OBRAS",
        (
            "Informe financiero consolidado por proyecto "
            f"· Generado {date.today().strftime('%d/%m/%Y')}"
        )
    )

    filas = []

    for proyecto in proyectos:
        resumen = _resumen_proyecto(
            proyecto
        )

        filas.append([
            proyecto.codigo,
            proyecto.nombre,
            proyecto.estado,
            _money_export(
                proyecto.presupuesto
            ),
            _money_export(
                resumen["ingresos"]
            ),
            _money_export(
                resumen["costos"]
            ),
            _money_export(
                resumen["resultado"]
            ),
            _money_export(
                resumen["cxp_pendiente"]
            ),
            _money_export(
                resumen["cxc_pendiente"]
            )
        ])

    story.append(
        _pdf_tabla(
            [
                "Código",
                "Proyecto",
                "Estado",
                "Presupuesto",
                "Ingresos",
                "Costos",
                "Resultado",
                "CxP pendiente",
                "CxC pendiente"
            ],
            filas,
            [
                20 * mm,
                48 * mm,
                22 * mm,
                28 * mm,
                28 * mm,
                28 * mm,
                28 * mm,
                28 * mm,
                28 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            f"CRV_proyectos_obras_"
            f"{date.today().isoformat()}.pdf"
        )
    )

# ============================================================
# CIERRE MENSUAL - FUNCIONES AUXILIARES
# ============================================================

def _rango_mes(anio, mes):
    ultimo_dia = calendar.monthrange(
        anio,
        mes
    )[1]

    return (
        date(anio, mes, 1),
        date(anio, mes, ultimo_dia)
    )


def _cierre_periodo(empresa_id, anio, mes):
    return (
        CRVCierreMensual.query
        .filter_by(
            empresa_id=empresa_id,
            anio=anio,
            mes=mes
        )
        .first()
    )


def _periodo_crv_cerrado(fecha, empresa_id=None):
    if not fecha:
        return False

    empresa_id = (
        empresa_id
        or session.get("empresa_id")
    )

    cierre = _cierre_periodo(
        empresa_id,
        fecha.year,
        fecha.month
    )

    return bool(
        cierre
        and cierre.estado == "CERRADO"
    )


def _movimientos_mes_crv(empresa_id, anio, mes):
    return (
        CRVMovimiento.query
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            db.extract(
                "year",
                CRVMovimiento.fecha
            ) == anio,
            db.extract(
                "month",
                CRVMovimiento.fecha
            ) == mes
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )


def _saldo_financiero_hasta(empresa_id, fecha_fin):
    movimientos = (
        CRVMovimiento.query
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            CRVMovimiento.fecha <= fecha_fin
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )

    cuentas = {}

    def asegurar(nombre):
        nombre = _nombre_cuenta_valido(
            nombre
        )

        if not nombre:
            return ""

        if nombre not in cuentas:
            cuentas[nombre] = 0.0

        return nombre

    for movimiento in movimientos:
        tipo = str(
            movimiento.tipo or ""
        ).strip().upper()

        origen = asegurar(
            movimiento.cuenta
        )

        destino = ""

        if tipo == "TRANSFERENCIA":
            destino = asegurar(
                movimiento.cuenta_destino
            )

        valor = _valor_movimiento_bancario(
            movimiento
        )

        if tipo == "SALDO INICIAL CUENTA":
            if origen:
                cuentas[origen] += valor
            continue

        if tipo == "INGRESO":
            if origen:
                cuentas[origen] += valor
            continue

        if tipo == "COBRO CUENTA POR COBRAR":
            if origen:
                cuentas[origen] += valor
            continue

        if tipo == "TRANSFERENCIA":
            if origen:
                cuentas[origen] -= valor
            if destino:
                cuentas[destino] += valor
            continue

        if tipo == "COMPRA":
            forma_pago = str(
                movimiento.forma_pago or ""
            ).strip().upper()

            if forma_pago in (
                "CREDITO",
                "CRÉDITO"
            ):
                continue

        if origen:
            cuentas[origen] -= valor

    return {
        "cuentas": cuentas,
        "total": sum(
            cuentas.values()
        )
    }


def _resumen_cierre_mes(
    empresa_id,
    anio,
    mes
):
    fecha_inicio, fecha_fin = _rango_mes(
        anio,
        mes
    )

    movimientos = _movimientos_mes_crv(
        empresa_id,
        anio,
        mes
    )

    ingresos = 0.0
    compras = 0.0
    gastos = 0.0

    tipos_gasto = {
        "GASTO",
        "NOMINA AFILIADOS",
        "NÓMINA AFILIADOS",
        "NOMINA OBRA",
        "NÓMINA OBRA",
        "ANTICIPO",
        "PRESTAMO EMPLEADO",
        "PRÉSTAMO EMPLEADO",
        "EGRESO",
        "OTRO EGRESO",
        "OTRO",
    }

    tipos_sin_resultado = {
        "TRANSFERENCIA",
        "PAGO CUENTA POR PAGAR",
        "COBRO CUENTA POR COBRAR",
        "SALDO INICIAL CUENTA",
    }

    for m in movimientos:
        tipo = str(
            m.tipo or ""
        ).strip().upper()

        valor = float(
            m.valor_neto
            or m.valor
            or 0
        )

        if tipo == "INGRESO":
            ingresos += valor

        elif tipo == "COMPRA":
            compras += valor

        elif tipo in tipos_gasto:
            gastos += valor

        elif tipo in tipos_sin_resultado:
            pass

        else:
            gastos += valor

    resultado = (
        ingresos
        - compras
        - gastos
    )

    cuentas_pagar = (
        CRVCuentaPagar.query
        .filter(
            CRVCuentaPagar.empresa_id == empresa_id,
            CRVCuentaPagar.fecha_documento >= fecha_inicio,
            CRVCuentaPagar.fecha_documento <= fecha_fin
        )
        .all()
    )

    for c in cuentas_pagar:
        c.recalcular_saldos()

    cuentas_cobrar = (
        CRVCuentaCobrar.query
        .filter(
            CRVCuentaCobrar.empresa_id == empresa_id,
            CRVCuentaCobrar.fecha_documento >= fecha_inicio,
            CRVCuentaCobrar.fecha_documento <= fecha_fin
        )
        .all()
    )

    for c in cuentas_cobrar:
        c.recalcular_saldos()

    cxp_facturado = sum(
        float(c.valor_neto or 0)
        for c in cuentas_pagar
    )

    cxp_pagado = sum(
        float(c.total_pagado or 0)
        for c in cuentas_pagar
    )

    cxp_pendiente = sum(
        float(c.saldo_pendiente or 0)
        for c in cuentas_pagar
    )

    cxc_facturado = sum(
        float(c.valor_neto or 0)
        for c in cuentas_cobrar
    )

    cxc_cobrado = sum(
        float(c.total_cobrado or 0)
        for c in cuentas_cobrar
    )

    cxc_pendiente = sum(
        float(c.saldo_pendiente or 0)
        for c in cuentas_cobrar
    )

    bancarios = [
        m
        for m in movimientos
        if _movimiento_es_bancario(m)
    ]

    sin_conciliar = [
        m
        for m in bancarios
        if not bool(m.revisado)
    ]

    saldos = _saldo_financiero_hasta(
        empresa_id,
        fecha_fin
    )

    # --------------------------------------------------------
    # VALIDACIÓN DE CONCILIACIÓN POR CUENTA
    # --------------------------------------------------------

    cuentas_periodo = sorted({
        _nombre_cuenta_valido(m.cuenta)
        for m in bancarios
        if _nombre_cuenta_valido(m.cuenta)
    })

    # También incluir la cuenta destino si se trata de una transferencia
    # entre cuentas propias de CRV.
    for m in bancarios:
        if str(m.tipo or "").strip().upper() == "TRANSFERENCIA":
            destino = _nombre_cuenta_valido(
                m.cuenta_destino
            )

            if destino and destino not in cuentas_periodo:
                cuentas_periodo.append(
                    destino
                )

    conciliaciones_pendientes = []

    for nombre_cuenta in cuentas_periodo:
        conciliacion_cuenta = _conciliacion_periodo(
            empresa_id,
            nombre_cuenta,
            anio,
            mes
        )

        if not conciliacion_cuenta:
            conciliaciones_pendientes.append({
                "cuenta": nombre_cuenta,
                "motivo": "SIN EXTRACTO REGISTRADO"
            })

            continue

        _, fecha_fin_periodo = _rango_mes(
            anio,
            mes
        )

        saldo_sistema_cuenta = _saldo_cuenta_hasta(
            empresa_id,
            nombre_cuenta,
            fecha_fin_periodo
        )

        diferencia_cuenta = (
            float(
                conciliacion_cuenta.saldo_extracto
                or 0
            )
            - saldo_sistema_cuenta
        )

        if (
            abs(diferencia_cuenta) >= 0.01
            or conciliacion_cuenta.estado != "CONCILIADA"
        ):
            conciliaciones_pendientes.append({
                "cuenta": nombre_cuenta,
                "motivo": (
                    f"DIFERENCIA ${diferencia_cuenta:,.0f}"
                    if abs(diferencia_cuenta) >= 0.01
                    else "CONCILIACIÓN PENDIENTE"
                )
            })

    alertas = []

    if sin_conciliar:
        alertas.append(
            (
                f"Hay {len(sin_conciliar)} movimiento(s) "
                "bancario(s) del mes pendientes de conciliación."
            )
        )

    movimientos_sin_tercero = [
        m
        for m in movimientos
        if (
            str(m.tipo or "").strip().upper()
            not in (
                "TRANSFERENCIA",
                "SALDO INICIAL CUENTA"
            )
            and not (
                m.tercero_id
                or str(
                    m.tercero_nombre
                    or ""
                ).strip()
            )
        )
    ]

    if movimientos_sin_tercero:
        alertas.append(
            (
                f"Hay {len(movimientos_sin_tercero)} movimiento(s) "
                "sin tercero identificado."
            )
        )

    if conciliaciones_pendientes:
        for item in conciliaciones_pendientes:
            alertas.append(
                (
                    f"Cuenta {item['cuenta']}: "
                    f"{item['motivo']}."
                )
            )

    return {
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "movimientos": movimientos,
        "ingresos": ingresos,
        "compras": compras,
        "gastos": gastos,
        "resultado": resultado,
        "cxp_facturado": cxp_facturado,
        "cxp_pagado": cxp_pagado,
        "cxp_pendiente": cxp_pendiente,
        "cxc_facturado": cxc_facturado,
        "cxc_cobrado": cxc_cobrado,
        "cxc_pendiente": cxc_pendiente,
        "saldo_bancos_caja": saldos["total"],
        "saldos_cuentas": saldos["cuentas"],
        "movimientos_mes": len(movimientos),
        "sin_conciliar": sin_conciliar,
        "movimientos_sin_conciliar": len(sin_conciliar),
        "alertas": alertas,
        "conciliaciones_pendientes": conciliaciones_pendientes,
        "puede_cerrar": (
            len(sin_conciliar) == 0
            and len(conciliaciones_pendientes) == 0
        ),
    }


def _fecha_formulario_cierre(valor):
    try:
        return datetime.strptime(
            str(valor or "").strip(),
            "%Y-%m-%d"
        ).date()
    except (ValueError, TypeError):
        return None


# ============================================================
# BLOQUEO DE MODIFICACIONES EN PERIODOS CERRADOS
# ============================================================

@crv_bp.before_request
def bloquear_periodos_cerrados_crv():
    if request.method != "POST":
        return None

    if not usuario_valido():
        return None

    empresa_id = session.get(
        "empresa_id"
    )

    if not empresa_id:
        return None

    endpoint = str(
        request.endpoint
        or ""
    )

    if endpoint in (
        "crv.cierre_mensual_crv",
        "crv.cierre_mensual_reabrir",
    ):
        return None

    fechas = []

    # Fecha enviada por formulario.
    for campo in (
        "fecha",
        "fecha_documento",
        "fecha_apertura",
    ):
        valor = request.form.get(
            campo
        )

        fecha = _fecha_formulario_cierre(
            valor
        )

        if fecha:
            fechas.append(
                fecha
            )

    # Movimiento existente.
    movimiento_id = (
        request.view_args.get(
            "movimiento_id"
        )
        if request.view_args
        else None
    )

    if movimiento_id:
        movimiento = db.session.get(
            CRVMovimiento,
            movimiento_id
        )

        if movimiento and movimiento.fecha:
            fechas.append(
                movimiento.fecha
            )

    # Cuenta por pagar existente.
    cuenta_id = (
        request.view_args.get(
            "cuenta_id"
        )
        if request.view_args
        else None
    )

    if cuenta_id and (
        "cuentas-por-pagar"
        in request.path
    ):
        cuenta = db.session.get(
            CRVCuentaPagar,
            cuenta_id
        )

        if (
            cuenta
            and cuenta.fecha_documento
        ):
            fechas.append(
                cuenta.fecha_documento
            )

    # Cuenta por cobrar existente.
    if cuenta_id and (
        "cuentas-por-cobrar"
        in request.path
    ):
        cuenta = db.session.get(
            CRVCuentaCobrar,
            cuenta_id
        )

        if (
            cuenta
            and cuenta.fecha_documento
        ):
            fechas.append(
                cuenta.fecha_documento
            )

    # Pago CxP existente.
    pago_id = (
        request.view_args.get(
            "pago_id"
        )
        if request.view_args
        else None
    )

    if pago_id:
        pago = db.session.get(
            CRVPagoCuentaPagar,
            pago_id
        )

        if pago and pago.fecha:
            fechas.append(
                pago.fecha
            )

    # Cobro CxC existente.
    cobro_id = (
        request.view_args.get(
            "cobro_id"
        )
        if request.view_args
        else None
    )

    if cobro_id:
        cobro = db.session.get(
            CRVCobroCuentaCobrar,
            cobro_id
        )

        if cobro and cobro.fecha:
            fechas.append(
                cobro.fecha
            )

    for fecha in fechas:
        if _periodo_crv_cerrado(
            fecha,
            empresa_id
        ):
            flash(
                (
                    f"El periodo {fecha.year}-{fecha.month:02d} "
                    "está CERRADO. Debes reabrir el mes antes "
                    "de modificar, eliminar o registrar información "
                    "con esa fecha."
                ),
                "warning"
            )

            return redirect(
                request.referrer
                or "/crv/cierre-mensual"
            )

    return None


# ============================================================
# CIERRE MENSUAL - PANEL
# ============================================================

@crv_bp.route(
    "/cierre-mensual",
    methods=["GET", "POST"]
)
def cierre_mensual_crv():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    periodo = request.values.get(
        "periodo",
        ""
    ).strip()

    if not periodo:
        hoy = date.today()

        periodo = (
            f"{hoy.year}-"
            f"{hoy.month:02d}"
        )

    try:
        anio_texto, mes_texto = periodo.split(
            "-"
        )

        anio = int(
            anio_texto
        )

        mes = int(
            mes_texto
        )

        if mes < 1 or mes > 12:
            raise ValueError

    except (ValueError, TypeError):
        hoy = date.today()
        anio = hoy.year
        mes = hoy.month
        periodo = (
            f"{anio}-"
            f"{mes:02d}"
        )

    resumen = _resumen_cierre_mes(
        empresa_id,
        anio,
        mes
    )

    cierre = _cierre_periodo(
        empresa_id,
        anio,
        mes
    )

    if request.method == "POST":

        accion = request.form.get(
            "accion",
            ""
        ).strip().lower()

        if accion == "cerrar":

            if (
                cierre
                and cierre.estado
                == "CERRADO"
            ):
                flash(
                    "Este periodo ya se encuentra cerrado.",
                    "warning"
                )

                return redirect(
                    f"/crv/cierre-mensual?periodo={periodo}"
                )

            if not resumen[
                "puede_cerrar"
            ]:
                flash(
                    (
                        "No se puede cerrar el periodo mientras "
                        "existan movimientos bancarios sin conciliar."
                    ),
                    "error"
                )

                return redirect(
                    f"/crv/cierre-mensual?periodo={periodo}"
                )

            if not cierre:
                cierre = CRVCierreMensual(
                    empresa_id=empresa_id,
                    anio=anio,
                    mes=mes
                )

                db.session.add(
                    cierre
                )

            cierre.estado = "CERRADO"
            cierre.ingresos = resumen[
                "ingresos"
            ]
            cierre.compras = resumen[
                "compras"
            ]
            cierre.gastos = resumen[
                "gastos"
            ]
            cierre.resultado = resumen[
                "resultado"
            ]

            cierre.cxp_facturado = resumen[
                "cxp_facturado"
            ]
            cierre.cxp_pagado = resumen[
                "cxp_pagado"
            ]
            cierre.cxp_pendiente = resumen[
                "cxp_pendiente"
            ]

            cierre.cxc_facturado = resumen[
                "cxc_facturado"
            ]
            cierre.cxc_cobrado = resumen[
                "cxc_cobrado"
            ]
            cierre.cxc_pendiente = resumen[
                "cxc_pendiente"
            ]

            cierre.saldo_bancos_caja = resumen[
                "saldo_bancos_caja"
            ]

            cierre.movimientos_mes = resumen[
                "movimientos_mes"
            ]

            cierre.movimientos_sin_conciliar = resumen[
                "movimientos_sin_conciliar"
            ]

            cierre.observaciones = request.form.get(
                "observaciones",
                ""
            ).strip()

            cierre.usuario_cierre = str(
                session.get(
                    "usuario",
                    ""
                )
            )

            cierre.fecha_cierre = datetime.utcnow()

            try:
                db.session.commit()

                flash(
                    (
                        f"Periodo {periodo} cerrado correctamente. "
                        "Los registros contables del mes quedaron bloqueados."
                    ),
                    "success"
                )

            except Exception as error:
                db.session.rollback()

                flash(
                    f"No fue posible cerrar el periodo: {error}",
                    "error"
                )

            return redirect(
                f"/crv/cierre-mensual?periodo={periodo}"
            )

    cierres = (
        CRVCierreMensual.query
        .filter_by(
            empresa_id=empresa_id
        )
        .order_by(
            CRVCierreMensual.anio.desc(),
            CRVCierreMensual.mes.desc()
        )
        .all()
    )

    return render_template(
        "contabilidad/crv_cierre_mensual.html",
        empresa_nombre=empresa_nombre,
        periodo=periodo,
        anio=anio,
        mes=mes,
        resumen=resumen,
        cierre=cierre,
        cierres=cierres
    )


# ============================================================
# CIERRE MENSUAL - REABRIR
# ============================================================

@crv_bp.route(
    "/cierre-mensual/<int:cierre_id>/reabrir",
    methods=["POST"]
)
def cierre_mensual_reabrir(
    cierre_id
):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    cierre = (
        CRVCierreMensual.query
        .get_or_404(
            cierre_id
        )
    )

    if cierre.empresa_id != empresa_id:
        return redirect(
            "/crv/cierre-mensual"
        )

    motivo = request.form.get(
        "motivo_reapertura",
        ""
    ).strip()

    if not motivo:
        flash(
            (
                "Para reabrir un periodo debes indicar "
                "el motivo de la reapertura."
            ),
            "error"
        )

        return redirect(
            (
                "/crv/cierre-mensual?"
                f"periodo={cierre.anio}-{cierre.mes:02d}"
            )
        )

    cierre.estado = "ABIERTO"

    cierre.usuario_reapertura = str(
        session.get(
            "usuario",
            ""
        )
    )

    cierre.fecha_reapertura = datetime.utcnow()

    cierre.motivo_reapertura = motivo

    db.session.commit()

    flash(
        (
            f"Periodo {cierre.anio}-{cierre.mes:02d} "
            "reabierto. Ya se pueden realizar correcciones."
        ),
        "success"
    )

    return redirect(
        (
            "/crv/cierre-mensual?"
            f"periodo={cierre.anio}-{cierre.mes:02d}"
        )
    )


# ============================================================
# CIERRE MENSUAL - EXCEL
# ============================================================

@crv_bp.route(
    "/cierre-mensual/excel"
)
def cierre_mensual_excel():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    periodo = request.args.get(
        "periodo",
        ""
    ).strip()

    try:
        anio, mes = [
            int(x)
            for x in periodo.split("-")
        ]
    except (ValueError, TypeError):
        hoy = date.today()
        anio = hoy.year
        mes = hoy.month
        periodo = (
            f"{anio}-"
            f"{mes:02d}"
        )

    resumen = _resumen_cierre_mes(
        empresa_id,
        anio,
        mes
    )

    cierre = _cierre_periodo(
        empresa_id,
        anio,
        mes
    )

    wb = Workbook()

    ws = wb.active
    ws.title = "CIERRE"

    _excel_titulo(
        ws,
        (
            "CRV CONSTRUCTORA - "
            f"CIERRE {periodo}"
        ),
        (
            "Informe mensual de control contable "
            f"· Estado: "
            f"{cierre.estado if cierre else 'ABIERTO'}"
        ),
        4
    )

    _excel_encabezado(
        ws,
        4,
        [
            "Concepto",
            "Valor",
            "Estado",
            "Observación"
        ]
    )

    filas = [
        (
            "Ingresos",
            resumen["ingresos"],
            "",
            ""
        ),
        (
            "Compras",
            resumen["compras"],
            "",
            ""
        ),
        (
            "Gastos",
            resumen["gastos"],
            "",
            ""
        ),
        (
            "Resultado",
            resumen["resultado"],
            "",
            ""
        ),
        (
            "CxP facturado",
            resumen["cxp_facturado"],
            "",
            ""
        ),
        (
            "CxP pagado",
            resumen["cxp_pagado"],
            "",
            ""
        ),
        (
            "CxP pendiente",
            resumen["cxp_pendiente"],
            "",
            ""
        ),
        (
            "CxC facturado",
            resumen["cxc_facturado"],
            "",
            ""
        ),
        (
            "CxC cobrado",
            resumen["cxc_cobrado"],
            "",
            ""
        ),
        (
            "CxC pendiente",
            resumen["cxc_pendiente"],
            "",
            ""
        ),
        (
            "Saldo bancos y caja al cierre",
            resumen["saldo_bancos_caja"],
            "",
            ""
        ),
        (
            "Movimientos del mes",
            resumen["movimientos_mes"],
            "",
            ""
        ),
        (
            "Movimientos sin conciliar",
            resumen[
                "movimientos_sin_conciliar"
            ],
            (
                "OK"
                if resumen[
                    "movimientos_sin_conciliar"
                ] == 0
                else "PENDIENTE"
            ),
            ""
        ),
    ]

    fila_excel = 5

    for concepto, valor, estado, nota in filas:
        ws.cell(
            fila_excel,
            1,
            concepto
        )

        ws.cell(
            fila_excel,
            2,
            valor
        )

        ws.cell(
            fila_excel,
            3,
            estado
        )

        ws.cell(
            fila_excel,
            4,
            nota
        )

        if concepto not in (
            "Movimientos del mes",
            "Movimientos sin conciliar",
        ):
            _excel_moneda(
                ws.cell(
                    fila_excel,
                    2
                )
            )

        fila_excel += 1

    _excel_auto_ancho(
        ws
    )

    mov_ws = wb.create_sheet(
        "MOVIMIENTOS"
    )

    headers = [
        "Fecha",
        "Tipo",
        "Tercero",
        "Documento",
        "Concepto",
        "Cuenta",
        "Cuenta destino",
        "Proyecto",
        "Centro costo",
        "Valor",
        "Conciliado"
    ]

    _excel_titulo(
        mov_ws,
        f"MOVIMIENTOS DEL PERIODO {periodo}",
        "Detalle utilizado para el cierre mensual",
        len(headers)
    )

    _excel_encabezado(
        mov_ws,
        4,
        headers
    )

    for fila, m in enumerate(
        resumen["movimientos"],
        5
    ):
        valores = [
            m.fecha,
            m.tipo,
            m.nombre_tercero,
            (
                m.numero_documento
                or m.numero_soporte
            ),
            m.concepto,
            m.cuenta,
            m.cuenta_destino,
            m.proyecto,
            m.centro_costo,
            float(
                m.valor_neto
                or m.valor
                or 0
            ),
            (
                "SI"
                if m.revisado
                else "NO"
            )
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            mov_ws.cell(
                fila,
                col,
                valor
            )

        _excel_moneda(
            mov_ws.cell(
                fila,
                10
            )
        )

    mov_ws.freeze_panes = "A5"

    mov_ws.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers))}"
        f"{max(4, mov_ws.max_row)}"
    )

    _excel_auto_ancho(
        mov_ws
    )

    bank_ws = wb.create_sheet(
        "SALDOS_BANCOS_CAJA"
    )

    _excel_titulo(
        bank_ws,
        f"SALDOS FINANCIEROS AL {resumen['fecha_fin'].strftime('%d/%m/%Y')}",
        "Saldos calculados de cuentas propias de CRV",
        2
    )

    _excel_encabezado(
        bank_ws,
        4,
        [
            "Cuenta CRV",
            "Saldo"
        ]
    )

    for fila, (
        nombre,
        saldo
    ) in enumerate(
        sorted(
            resumen[
                "saldos_cuentas"
            ].items()
        ),
        5
    ):
        bank_ws.cell(
            fila,
            1,
            nombre
        )

        bank_ws.cell(
            fila,
            2,
            saldo
        )

        _excel_moneda(
            bank_ws.cell(
                fila,
                2
            )
        )

    _excel_auto_ancho(
        bank_ws
    )

    return _enviar_excel(
        wb,
        (
            f"CRV_cierre_mensual_"
            f"{periodo}.xlsx"
        )
    )


# ============================================================
# CIERRE MENSUAL - PDF
# ============================================================

@crv_bp.route(
    "/cierre-mensual/pdf"
)
def cierre_mensual_pdf():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    periodo = request.args.get(
        "periodo",
        ""
    ).strip()

    try:
        anio, mes = [
            int(x)
            for x in periodo.split("-")
        ]
    except (ValueError, TypeError):
        hoy = date.today()
        anio = hoy.year
        mes = hoy.month
        periodo = (
            f"{anio}-"
            f"{mes:02d}"
        )

    resumen = _resumen_cierre_mes(
        empresa_id,
        anio,
        mes
    )

    cierre = _cierre_periodo(
        empresa_id,
        anio,
        mes
    )

    story = []

    _pdf_encabezado_historia(
        story,
        (
            "CRV CONSTRUCTORA - "
            f"CIERRE MENSUAL {periodo}"
        ),
        (
            f"Estado: "
            f"{cierre.estado if cierre else 'ABIERTO'} "
            "· Informe para administración, gerencia y contador"
        )
    )

    story.append(
        Paragraph(
            "Resumen financiero",
            _pdf_estilos_crv()[2]
        )
    )

    story.append(
        _pdf_tabla(
            [
                "Concepto",
                "Valor"
            ],
            [
                [
                    "Ingresos",
                    _money_export(
                        resumen["ingresos"]
                    )
                ],
                [
                    "Compras",
                    _money_export(
                        resumen["compras"]
                    )
                ],
                [
                    "Gastos",
                    _money_export(
                        resumen["gastos"]
                    )
                ],
                [
                    "Resultado",
                    _money_export(
                        resumen["resultado"]
                    )
                ],
                [
                    "CxP pendiente",
                    _money_export(
                        resumen["cxp_pendiente"]
                    )
                ],
                [
                    "CxC pendiente",
                    _money_export(
                        resumen["cxc_pendiente"]
                    )
                ],
                [
                    "Saldo bancos y caja",
                    _money_export(
                        resumen[
                            "saldo_bancos_caja"
                        ]
                    )
                ],
                [
                    "Movimientos sin conciliar",
                    resumen[
                        "movimientos_sin_conciliar"
                    ]
                ],
            ],
            [
                85 * mm,
                45 * mm
            ]
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )

    story.append(
        Paragraph(
            "Movimientos del periodo",
            _pdf_estilos_crv()[2]
        )
    )

    filas_mov = []

    for m in resumen["movimientos"]:
        filas_mov.append([
            (
                m.fecha.strftime(
                    "%d/%m/%Y"
                )
                if m.fecha
                else ""
            ),
            m.tipo or "",
            m.nombre_tercero or "",
            m.concepto or "",
            m.cuenta or "",
            m.proyecto or "",
            _money_export(
                m.valor_neto
                or m.valor
                or 0
            ),
            (
                "SI"
                if m.revisado
                else "NO"
            )
        ])

    story.append(
        _pdf_tabla(
            [
                "Fecha",
                "Tipo",
                "Tercero",
                "Concepto",
                "Cuenta",
                "Proyecto",
                "Valor",
                "Conciliado"
            ],
            filas_mov,
            [
                20 * mm,
                30 * mm,
                43 * mm,
                52 * mm,
                34 * mm,
                34 * mm,
                27 * mm,
                24 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            f"CRV_cierre_mensual_"
            f"{periodo}.pdf"
        )
    )

# ============================================================
# CONCILIACIÓN BANCARIA - FUNCIONES AUXILIARES
# ============================================================

def _impacto_bancario_movimiento(movimiento, cuenta):
    tipo = str(
        movimiento.tipo or ""
    ).strip().upper()

    origen = _nombre_cuenta_valido(
        movimiento.cuenta
    )

    destino = ""

    if tipo == "TRANSFERENCIA":
        destino = _nombre_cuenta_valido(
            movimiento.cuenta_destino
        )

    valor = _valor_movimiento_bancario(
        movimiento
    )

    impacto = 0.0

    if tipo == "SALDO INICIAL CUENTA":
        if origen == cuenta:
            impacto += valor
        return impacto

    if tipo == "INGRESO":
        if origen == cuenta:
            impacto += valor
        return impacto

    if tipo == "COBRO CUENTA POR COBRAR":
        if origen == cuenta:
            impacto += valor
        return impacto

    if tipo == "TRANSFERENCIA":
        if origen == cuenta:
            impacto -= valor

        if destino == cuenta:
            impacto += valor

        return impacto

    if tipo == "COMPRA":
        forma_pago = str(
            movimiento.forma_pago or ""
        ).strip().upper()

        if forma_pago in (
            "CREDITO",
            "CRÉDITO"
        ):
            return 0.0

    if origen == cuenta:
        impacto -= valor

    return impacto


def _saldo_cuenta_hasta(
    empresa_id,
    cuenta,
    fecha_fin
):
    movimientos = (
        CRVMovimiento.query
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            CRVMovimiento.fecha <= fecha_fin
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )

    return sum(
        _impacto_bancario_movimiento(
            movimiento,
            cuenta
        )
        for movimiento in movimientos
    )


def _movimientos_cuenta_periodo(
    empresa_id,
    cuenta,
    anio,
    mes
):
    movimientos = (
        CRVMovimiento.query
        .filter(
            CRVMovimiento.empresa_id == empresa_id,
            db.extract(
                "year",
                CRVMovimiento.fecha
            ) == anio,
            db.extract(
                "month",
                CRVMovimiento.fecha
            ) == mes
        )
        .order_by(
            CRVMovimiento.fecha.asc(),
            CRVMovimiento.id.asc()
        )
        .all()
    )

    return [
        movimiento
        for movimiento in movimientos
        if abs(
            _impacto_bancario_movimiento(
                movimiento,
                cuenta
            )
        ) > 0
    ]


def _conciliacion_periodo(
    empresa_id,
    cuenta,
    anio,
    mes
):
    return (
        CRVConciliacionBancaria.query
        .filter_by(
            empresa_id=empresa_id,
            cuenta=cuenta,
            anio=anio,
            mes=mes
        )
        .first()
    )


# ============================================================
# CONCILIACIÓN BANCARIA - PANEL PRINCIPAL
# ============================================================

@crv_bp.route(
    "/conciliacion-bancaria",
    methods=["GET", "POST"]
)
def conciliacion_bancaria_crv():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    empresa_nombre = session.get(
        "empresa_nombre",
        "CRV CONSTRUCTORA"
    )

    cuentas_data = _cuentas_financieras_crv(
        empresa_id
    )

    cuentas = [
        c["nombre"]
        for c in cuentas_data
    ]

    periodo = request.values.get(
        "periodo",
        ""
    ).strip()

    if not periodo:
        hoy = date.today()
        periodo = (
            f"{hoy.year}-"
            f"{hoy.month:02d}"
        )

    try:
        anio_texto, mes_texto = periodo.split("-")
        anio = int(anio_texto)
        mes = int(mes_texto)

        if mes < 1 or mes > 12:
            raise ValueError

    except (ValueError, TypeError):
        hoy = date.today()
        anio = hoy.year
        mes = hoy.month
        periodo = (
            f"{anio}-"
            f"{mes:02d}"
        )

    cuenta = request.values.get(
        "cuenta",
        ""
    ).strip()

    if not cuenta and cuentas:
        cuenta = cuentas[0]

    if cuenta and cuenta not in cuentas:
        cuenta = (
            cuentas[0]
            if cuentas
            else ""
        )

    fecha_inicio, fecha_fin = _rango_mes(
        anio,
        mes
    )

    movimientos = []

    saldo_sistema = 0.0
    saldo_extracto = None
    diferencia = None
    conciliacion = None

    if cuenta:
        movimientos = _movimientos_cuenta_periodo(
            empresa_id,
            cuenta,
            anio,
            mes
        )

        saldo_sistema = _saldo_cuenta_hasta(
            empresa_id,
            cuenta,
            fecha_fin
        )

        conciliacion = _conciliacion_periodo(
            empresa_id,
            cuenta,
            anio,
            mes
        )

        if conciliacion:
            saldo_extracto = float(
                conciliacion.saldo_extracto
                or 0
            )

            diferencia = (
                saldo_extracto
                - saldo_sistema
            )

    if request.method == "POST":

        accion = request.form.get(
            "accion",
            ""
        ).strip().lower()

        if accion == "guardar_extracto":

            if not cuenta:
                flash(
                    "Debes seleccionar una cuenta propia de CRV.",
                    "error"
                )

                return redirect(
                    f"/crv/conciliacion-bancaria?periodo={periodo}"
                )

            saldo_extracto = convertir_numero(
                request.form.get(
                    "saldo_extracto"
                )
            )

            observaciones = request.form.get(
                "observaciones",
                ""
            ).strip()

            if not conciliacion:
                conciliacion = CRVConciliacionBancaria(
                    empresa_id=empresa_id,
                    cuenta=cuenta,
                    anio=anio,
                    mes=mes
                )

                db.session.add(
                    conciliacion
                )

            conciliacion.saldo_extracto = saldo_extracto
            conciliacion.saldo_sistema = saldo_sistema
            conciliacion.diferencia = (
                saldo_extracto
                - saldo_sistema
            )

            conciliacion.observaciones = observaciones
            conciliacion.usuario_registro = str(
                session.get(
                    "usuario",
                    ""
                )
            )

            pendientes = [
                m
                for m in movimientos
                if not bool(m.revisado)
            ]

            if (
                abs(
                    conciliacion.diferencia
                ) < 0.01
                and not pendientes
            ):
                conciliacion.estado = "CONCILIADA"
                conciliacion.fecha_conciliacion = datetime.utcnow()

            else:
                conciliacion.estado = "PENDIENTE"
                conciliacion.fecha_conciliacion = None

            db.session.commit()

            flash(
                "Saldo del extracto guardado correctamente.",
                "success"
            )

            return redirect(
                (
                    "/crv/conciliacion-bancaria?"
                    f"periodo={periodo}"
                    f"&cuenta={cuenta}"
                )
            )

    pendientes = [
        m
        for m in movimientos
        if not bool(m.revisado)
    ]

    conciliados = [
        m
        for m in movimientos
        if bool(m.revisado)
    ]

    entradas = sum(
        max(
            0,
            _impacto_bancario_movimiento(
                m,
                cuenta
            )
        )
        for m in movimientos
    ) if cuenta else 0

    salidas = sum(
        abs(
            min(
                0,
                _impacto_bancario_movimiento(
                    m,
                    cuenta
                )
            )
        )
        for m in movimientos
    ) if cuenta else 0

    return render_template(
        "contabilidad/crv_conciliacion_bancaria.html",
        empresa_nombre=empresa_nombre,
        cuentas=cuentas,
        cuenta_actual=cuenta,
        periodo=periodo,
        anio=anio,
        mes=mes,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        movimientos=movimientos,
        pendientes=pendientes,
        conciliados=conciliados,
        entradas=entradas,
        salidas=salidas,
        saldo_sistema=saldo_sistema,
        saldo_extracto=saldo_extracto,
        diferencia=diferencia,
        extracto_registrado=bool(conciliacion),
        conciliacion=conciliacion
    )


# ============================================================
# CONCILIACIÓN BANCARIA - MARCAR MOVIMIENTO
# ============================================================

@crv_bp.route(
    "/conciliacion-bancaria/movimiento/<int:movimiento_id>",
    methods=["POST"]
)
def conciliacion_bancaria_movimiento(
    movimiento_id
):

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    movimiento = (
        CRVMovimiento.query
        .get_or_404(
            movimiento_id
        )
    )

    if movimiento.empresa_id != empresa_id:
        return redirect(
            "/crv/conciliacion-bancaria"
        )

    periodo = request.form.get(
        "periodo",
        ""
    ).strip()

    cuenta = request.form.get(
        "cuenta",
        ""
    ).strip()

    movimiento.revisado = not bool(
        movimiento.revisado
    )

    db.session.commit()

    flash(
        (
            "Movimiento marcado como conciliado."
            if movimiento.revisado
            else "Movimiento devuelto a pendiente."
        ),
        "success"
    )

    return redirect(
        (
            "/crv/conciliacion-bancaria?"
            f"periodo={periodo}"
            f"&cuenta={cuenta}"
        )
    )


# ============================================================
# CONCILIACIÓN BANCARIA - CONCILIAR TODO EL PERIODO
# ============================================================

@crv_bp.route(
    "/conciliacion-bancaria/conciliar-todo",
    methods=["POST"]
)
def conciliacion_bancaria_todo():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    periodo = request.form.get(
        "periodo",
        ""
    ).strip()

    cuenta = request.form.get(
        "cuenta",
        ""
    ).strip()

    try:
        anio, mes = [
            int(x)
            for x in periodo.split("-")
        ]
    except (ValueError, TypeError):
        flash(
            "Periodo inválido.",
            "error"
        )

        return redirect(
            "/crv/conciliacion-bancaria"
        )

    movimientos = _movimientos_cuenta_periodo(
        empresa_id,
        cuenta,
        anio,
        mes
    )

    for movimiento in movimientos:
        movimiento.revisado = True

    _, fecha_fin = _rango_mes(
        anio,
        mes
    )

    saldo_sistema = _saldo_cuenta_hasta(
        empresa_id,
        cuenta,
        fecha_fin
    )

    conciliacion = _conciliacion_periodo(
        empresa_id,
        cuenta,
        anio,
        mes
    )

    if conciliacion:
        conciliacion.saldo_sistema = saldo_sistema
        conciliacion.diferencia = (
            float(
                conciliacion.saldo_extracto
                or 0
            )
            - saldo_sistema
        )

        if abs(
            conciliacion.diferencia
        ) < 0.01:
            conciliacion.estado = "CONCILIADA"
            conciliacion.fecha_conciliacion = datetime.utcnow()

    db.session.commit()

    flash(
        "Todos los movimientos de la cuenta y periodo fueron marcados como conciliados.",
        "success"
    )

    return redirect(
        (
            "/crv/conciliacion-bancaria?"
            f"periodo={periodo}"
            f"&cuenta={cuenta}"
        )
    )


# ============================================================
# CONCILIACIÓN BANCARIA - EXCEL
# ============================================================

@crv_bp.route(
    "/conciliacion-bancaria/excel"
)
def conciliacion_bancaria_excel():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    periodo = request.args.get(
        "periodo",
        ""
    ).strip()

    cuenta = request.args.get(
        "cuenta",
        ""
    ).strip()

    try:
        anio, mes = [
            int(x)
            for x in periodo.split("-")
        ]
    except (ValueError, TypeError):
        hoy = date.today()
        anio = hoy.year
        mes = hoy.month
        periodo = (
            f"{anio}-"
            f"{mes:02d}"
        )

    _, fecha_fin = _rango_mes(
        anio,
        mes
    )

    movimientos = _movimientos_cuenta_periodo(
        empresa_id,
        cuenta,
        anio,
        mes
    )

    saldo_sistema = _saldo_cuenta_hasta(
        empresa_id,
        cuenta,
        fecha_fin
    )

    conciliacion = _conciliacion_periodo(
        empresa_id,
        cuenta,
        anio,
        mes
    )

    saldo_extracto = (
        float(
            conciliacion.saldo_extracto
            or 0
        )
        if conciliacion
        else 0
    )

    wb = Workbook()

    ws = wb.active
    ws.title = "CONCILIACION"

    _excel_titulo(
        ws,
        "CRV CONSTRUCTORA - CONCILIACIÓN BANCARIA",
        (
            f"Cuenta: {cuenta} · Periodo: {periodo}"
        ),
        4
    )

    _excel_encabezado(
        ws,
        4,
        [
            "Concepto",
            "Sistema",
            "Extracto",
            "Diferencia"
        ]
    )

    ws.cell(5, 1, "Saldo final")
    ws.cell(5, 2, saldo_sistema)
    ws.cell(5, 3, saldo_extracto)
    ws.cell(
        5,
        4,
        saldo_extracto - saldo_sistema
    )

    for col in range(
        2,
        5
    ):
        _excel_moneda(
            ws.cell(
                5,
                col
            )
        )

    _excel_auto_ancho(
        ws
    )

    mov = wb.create_sheet(
        "MOVIMIENTOS"
    )

    headers = [
        "Fecha",
        "Tipo",
        "Tercero",
        "Concepto",
        "Documento",
        "Cuenta",
        "Cuenta destino",
        "Entrada",
        "Salida",
        "Valor",
        "Conciliado"
    ]

    _excel_titulo(
        mov,
        "MOVIMIENTOS DE CONCILIACIÓN",
        f"{cuenta} · {periodo}",
        len(headers)
    )

    _excel_encabezado(
        mov,
        4,
        headers
    )

    for fila, m in enumerate(
        movimientos,
        5
    ):
        impacto = _impacto_bancario_movimiento(
            m,
            cuenta
        )

        valores = [
            m.fecha,
            m.tipo,
            m.nombre_tercero,
            m.concepto,
            (
                m.numero_documento
                or m.numero_soporte
            ),
            m.cuenta,
            m.cuenta_destino,
            max(0, impacto),
            abs(min(0, impacto)),
            abs(impacto),
            (
                "SI"
                if m.revisado
                else "NO"
            )
        ]

        for col, valor in enumerate(
            valores,
            1
        ):
            mov.cell(
                fila,
                col,
                valor
            )

        for col in (
            8,
            9,
            10
        ):
            _excel_moneda(
                mov.cell(
                    fila,
                    col
                )
            )

    mov.freeze_panes = "A5"
    mov.auto_filter.ref = (
        f"A4:"
        f"{get_column_letter(len(headers))}"
        f"{max(4, mov.max_row)}"
    )

    _excel_auto_ancho(
        mov
    )

    return _enviar_excel(
        wb,
        (
            "CRV_conciliacion_"
            f"{cuenta.replace(' ', '_')}_"
            f"{periodo}.xlsx"
        )
    )


# ============================================================
# CONCILIACIÓN BANCARIA - PDF
# ============================================================

@crv_bp.route(
    "/conciliacion-bancaria/pdf"
)
def conciliacion_bancaria_pdf():

    bloqueo = proteger_crv()
    if bloqueo:
        return bloqueo

    empresa_id = session.get(
        "empresa_id"
    )

    periodo = request.args.get(
        "periodo",
        ""
    ).strip()

    cuenta = request.args.get(
        "cuenta",
        ""
    ).strip()

    try:
        anio, mes = [
            int(x)
            for x in periodo.split("-")
        ]
    except (ValueError, TypeError):
        hoy = date.today()
        anio = hoy.year
        mes = hoy.month
        periodo = (
            f"{anio}-"
            f"{mes:02d}"
        )

    _, fecha_fin = _rango_mes(
        anio,
        mes
    )

    movimientos = _movimientos_cuenta_periodo(
        empresa_id,
        cuenta,
        anio,
        mes
    )

    saldo_sistema = _saldo_cuenta_hasta(
        empresa_id,
        cuenta,
        fecha_fin
    )

    conciliacion = _conciliacion_periodo(
        empresa_id,
        cuenta,
        anio,
        mes
    )

    saldo_extracto = (
        float(
            conciliacion.saldo_extracto
            or 0
        )
        if conciliacion
        else 0
    )

    story = []

    _pdf_encabezado_historia(
        story,
        "CRV CONSTRUCTORA - CONCILIACIÓN BANCARIA",
        (
            f"Cuenta: {cuenta} · Periodo: {periodo}"
        )
    )

    story.append(
        _pdf_tabla(
            [
                "Saldo sistema",
                "Saldo extracto",
                "Diferencia",
                "Estado"
            ],
            [[
                _money_export(
                    saldo_sistema
                ),
                _money_export(
                    saldo_extracto
                ),
                _money_export(
                    saldo_extracto
                    - saldo_sistema
                ),
                (
                    conciliacion.estado
                    if conciliacion
                    else "PENDIENTE"
                )
            ]],
            [
                45 * mm,
                45 * mm,
                45 * mm,
                35 * mm
            ]
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )

    filas = []

    for m in movimientos:
        impacto = _impacto_bancario_movimiento(
            m,
            cuenta
        )

        filas.append([
            (
                m.fecha.strftime(
                    "%d/%m/%Y"
                )
                if m.fecha
                else ""
            ),
            m.tipo or "",
            m.nombre_tercero or "",
            m.concepto or "",
            _money_export(
                max(
                    0,
                    impacto
                )
            ),
            _money_export(
                abs(
                    min(
                        0,
                        impacto
                    )
                )
            ),
            (
                "SI"
                if m.revisado
                else "NO"
            )
        ])

    story.append(
        _pdf_tabla(
            [
                "Fecha",
                "Tipo",
                "Tercero",
                "Concepto",
                "Entrada",
                "Salida",
                "Conciliado"
            ],
            filas,
            [
                20 * mm,
                30 * mm,
                45 * mm,
                65 * mm,
                28 * mm,
                28 * mm,
                24 * mm
            ]
        )
    )

    return _enviar_pdf(
        story,
        (
            "CRV_conciliacion_"
            f"{cuenta.replace(' ', '_')}_"
            f"{periodo}.pdf"
        )
    )